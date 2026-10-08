"""Tests de cohérence du carnet guidé des 13 points de contrôle rapides."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).parent.parent
WORKBOOK_PATH = PROJECT_ROOT / "03-easy-checks" / "grille-audit-easy-checks.xlsx"
DOCS_WORKBOOK_PATH = PROJECT_ROOT / "docs" / "grille-audit-easy-checks.xlsx"
DOWNLOAD_WORKBOOK_PATH = (
    PROJECT_ROOT / "docs" / "assets" / "downloads" / "grille-audit-easy-checks.xlsx"
)
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from generate_grille_audit import (  # noqa: E402
    BASE_URL,
    CHECKS,
    EXERCICE_PAGES,
    FORMATION,
    WORKBOOK_TITLE,
)
from slides import discover_slides  # noqa: E402


def iter_text(workbook):
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    yield worksheet.title, cell.coordinate, cell.value


def test_le_classeur_contient_demarrer_les_13_points_et_la_synthese():
    workbook = load_workbook(WORKBOOK_PATH, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Démarrer",
            *(check["onglet"] for check in CHECKS),
            "Synthèse",
        ]
        assert workbook.properties.title.endswith(WORKBOOK_TITLE)
        assert workbook.properties.language == "fr-FR"
    finally:
        workbook.close()


def test_chaque_fiche_est_verticale_vide_et_reliee_a_sa_page():
    workbook = load_workbook(WORKBOOK_PATH, read_only=False, data_only=False)
    try:
        for index, (check, page) in enumerate(zip(CHECKS, EXERCICE_PAGES, strict=True)):
            worksheet = workbook[check["onglet"]]
            assert worksheet["A1"].value == f"#{check['id']:02d}"
            assert worksheet["B1"].value == check["titre"]
            assert worksheet["B3"].value == page["title"]
            assert worksheet["B3"].hyperlink.target == (
                f"{BASE_URL}/site-inaccessible/{page['id']}.html"
            )
            assert worksheet["B4"].value == page["slides"]
            assert worksheet["B5"].value == check["question"]
            assert worksheet["B6"].value == check["wcag"]
            assert worksheet["B7"].value == check["rgaa"]
            assert worksheet["B11"].value in (None, "")
            assert worksheet["B12"].value in (None, "")
            assert worksheet["B13"].value in (None, "")

            previous_sheet = "Démarrer" if index == 0 else CHECKS[index - 1]["onglet"]
            next_sheet = "Synthèse" if index == 12 else CHECKS[index + 1]["onglet"]
            assert worksheet["A15"].hyperlink.location == "'Démarrer'!A1"
            assert worksheet["B15"].hyperlink.location == f"'{previous_sheet}'!A1"
            assert worksheet["A16"].hyperlink.location == f"'{next_sheet}'!A1"
            assert worksheet["B16"].hyperlink.location == "'Synthèse'!A1"

            targets = [
                cell.hyperlink.target
                for row in worksheet.iter_rows()
                for cell in row
                if cell.hyperlink is not None and cell.hyperlink.target is not None
            ]
            assert not any("site-aide-correction" in target for target in targets)
            assert not any("site-accessible" in target for target in targets)
    finally:
        workbook.close()


def test_demarrer_regroupe_les_liens_aide_et_correction():
    workbook = load_workbook(WORKBOOK_PATH, read_only=False, data_only=False)
    try:
        worksheet = workbook["Démarrer"]
        targets = [
            cell.hyperlink.target
            for row in worksheet.iter_rows()
            for cell in row
            if cell.hyperlink is not None and cell.hyperlink.target is not None
        ]
        assert sum("site-inaccessible" in target for target in targets) == 13
        assert sum("site-aide-correction" in target for target in targets) == 13
        assert sum("site-accessible" in target for target in targets) == 13
        assert worksheet["A26"].value == "Après la recherche autonome"
        assert workbook["#06 - Clavier"]["B7"].value == "7.3, 10.7, 12.8, 12.9"
        assert workbook["#07 - Langue"]["B7"].value == "8.3, 8.4, 8.7, 8.8"
        assert workbook["#12 - Étiquettes"]["B7"].value.endswith("11.5, 11.6, 11.7")
        assert workbook["#01 - Images"]["B8"].value == (
            "Afficher les attributs alt avec Web Developer "
            "(Images > Display Alt Attributes) ou utiliser le bookmarklet « Check images », "
            "puis vérifier leur pertinence selon le rôle de l'image et le contexte du lien."
        )
        assert workbook["Démarrer"].row_dimensions[27].height == 42
        assert workbook["Synthèse"].row_dimensions[1].height == 34
        assert [worksheet[cell].value for cell in ("A5", "C5", "A6", "C6")] == [
            "Auditeur",
            "Date",
            "Navigateur",
            "Outils",
        ]
    finally:
        workbook.close()


def test_la_synthese_relie_les_trois_champs_des_13_fiches():
    workbook = load_workbook(WORKBOOK_PATH, read_only=False, data_only=False)
    try:
        worksheet = workbook["Synthèse"]
        assert [worksheet.cell(row=4, column=column).value for column in range(1, 7)] == [
            "#",
            "Page",
            "Point testé",
            "Constat",
            "Mise en conformité à réaliser",
            "Preuve de l'écart",
        ]
        for row, check in enumerate(CHECKS, start=5):
            assert "À renseigner" in worksheet.cell(row=row, column=4).value
            assert "!B12" in worksheet.cell(row=row, column=5).value
            assert '\"\",\"\"' in worksheet.cell(row=row, column=5).value
            assert "!B13" in worksheet.cell(row=row, column=6).value
            assert '\"\",\"\"' in worksheet.cell(row=row, column=6).value
            assert worksheet.cell(row=row, column=3).hyperlink.location == f"'{check['onglet']}'!A1"
    finally:
        workbook.close()


def test_les_anciens_mecanismes_de_notation_ont_disparu():
    workbook = load_workbook(WORKBOOK_PATH, read_only=False, data_only=False)
    try:
        text = "\n".join(value for _, _, value in iter_text(workbook))
        for forbidden in (
            "Mon engagement pour demain 9 h",
            "Conventions de verdict",
            "Conventions de sévérité",
            "Top 3 des actions prioritaires",
            "Taux de conformité",
        ):
            assert forbidden not in text
        assert "Mode d'emploi" not in workbook.sheetnames
        assert "Exemple" not in workbook.sheetnames
        assert "Échantillon RGAA" not in workbook.sheetnames
    finally:
        workbook.close()


def test_les_trois_copies_du_classeur_sont_strictement_identiques():
    expected = WORKBOOK_PATH.read_bytes()
    assert DOCS_WORKBOOK_PATH.read_bytes() == expected
    assert DOWNLOAD_WORKBOOK_PATH.read_bytes() == expected


def test_l_url_du_site_est_issue_de_la_configuration():
    assert BASE_URL == FORMATION["site_url"].rstrip("/")
    source = (PROJECT_ROOT / "scripts" / "generate_grille_audit.py").read_text(
        encoding="utf-8"
    )
    assert 'BASE_URL = FORMATION["site_url"].rstrip("/")' in source


def test_le_poids_affiche_sur_l_accueil_correspond_au_classeur():
    size_kib = round(DOWNLOAD_WORKBOOK_PATH.stat().st_size / 1024)
    home = (PROJECT_ROOT / "docs" / "index.html").read_text(encoding="utf-8")
    assert f"XLSX - {size_kib} Ko" in home


def test_les_sources_pedagogiques_utilisent_le_diagnostic_a_trois_champs():
    paths = [
        PROJECT_ROOT / "03-easy-checks" / "evaluation_contract.yml",
        PROJECT_ROOT / "03-easy-checks" / "grille-audit-easy-checks.md",
        PROJECT_ROOT / "03-easy-checks" / "exercice-site-easy-checks-spec.md",
        PROJECT_ROOT / "docs" / "corrige-easy-checks.md",
        PROJECT_ROOT / "scripts" / "slides" / "51a_grille-remontee.py",
        PROJECT_ROOT / "scripts" / "slides" / "52_mission-13-checks.py",
    ]
    text = "\n".join(path.read_text(encoding="utf-8-sig") for path in paths)
    for forbidden in (
        "Conventions de verdict",
        "Conventions de sévérité",
        "Sévérité indicative",
        "grid.verdict",
        "grid.severity",
        "Taux de conformité points de contrôle rapides",
        "Une seule NC prouvée",
    ):
        assert forbidden not in text


def test_les_index_des_sites_d_exercice_utilisent_des_titres_de_niveau_2():
    for variant in ("site-inaccessible", "site-aide-correction"):
        source = (PROJECT_ROOT / "docs" / variant / "index.html").read_text(
            encoding="utf-8"
        )
        assert source.count('<h2 class="fr-card__title">') == 13
        assert '<h3 class="fr-card__title">' not in source


def test_les_pages_a_auditer_annoncent_le_controle_et_sa_reference_wcag():
    for check, page in zip(CHECKS, EXERCICE_PAGES, strict=True):
        source = (
            PROJECT_ROOT / "docs" / "site-inaccessible" / f"{page['id']}.html"
        ).read_text(encoding="utf-8")
        exercise_number = f"#{check['id']:02d}"
        assert source.count(
            f'aria-label="Consigne de l\'exercice {exercise_number}"'
        ) == 1
        assert f"<strong>WCAG 2.2 :</strong> {check['wcag']}." in source
        assert (
            f"Consignez votre diagnostic dans l'onglet « {check['onglet']} » de la grille."
            in source
        )


def test_le_point_01_utilise_le_titre_et_le_sous_titre_valides():
    title = "Contrôler les images avant publication"
    card_description = "Rôle des images et alternatives textuelles"
    assert EXERCICE_PAGES[0]["title"] == title

    index_paths = [
        PROJECT_ROOT / "docs" / "index.html",
        PROJECT_ROOT / "docs" / "site-inaccessible" / "index.html",
        PROJECT_ROOT / "docs" / "site-aide-correction" / "index.html",
        PROJECT_ROOT / "docs" / "site-accessible" / "index.html",
    ]
    for path in index_paths:
        source = path.read_text(encoding="utf-8")
        assert title in source
        assert card_description in source

    for variant in ("site-inaccessible", "site-aide-correction", "site-accessible"):
        source = (
            PROJECT_ROOT / "docs" / variant / "ec01-images.html"
        ).read_text(encoding="utf-8")
        assert title in source
        assert "Actualité illustrée" not in source


def test_les_pages_medias_corrigees_restent_sobres():
    expected_absent = {
        "ec09-captions.html": ("Contrôle de la correction", 'role="group"'),
        "ec10-transcript.html": (
            "Contrôle de la correction",
            "Choisir le bon niveau de transcription",
            'role="group"',
        ),
        "ec11-audio-description.html": (
            "Contrôle de la correction",
            "Correction principale",
            "Compléments de vérification",
            'role="group"',
        ),
    }
    for filename, forbidden_values in expected_absent.items():
        source = (
            PROJECT_ROOT / "docs" / "site-accessible" / filename
        ).read_text(encoding="utf-8")
        for forbidden in forbidden_values:
            assert forbidden not in source


def test_la_version_corrigee_ne_montre_pas_le_dispositif_pedagogique():
    forbidden_values = (
        "Mission d'audit pédagogique",
        "Site à auditer",
        "Aide à la correction",
        "Site corrigé",
        "Point de contrôle rapide",
        "Index généré depuis le contrat d'évaluation",
    )
    for path in sorted((PROJECT_ROOT / "docs" / "site-accessible").glob("*.html")):
        source = path.read_text(encoding="utf-8")
        for forbidden in forbidden_values:
            assert forbidden not in source, f"{path.name}: marqueur pédagogique visible {forbidden!r}"


def test_ec06_corrige_synchronise_aria_expanded_aux_evenements_dsfr():
    source = (
        PROJECT_ROOT / "docs" / "site-accessible" / "ec06-keyboard-focus.html"
    ).read_text(encoding="utf-8")
    assert 'modal.addEventListener("dsfr.disclose", function ()' in source
    assert 'modal.addEventListener("dsfr.conceal", function ()' in source
    assert 'opener.setAttribute("aria-expanded", "false")' in source


def test_les_pages_d_exercice_pointent_vers_les_slides_du_deck_courant():
    positions = {path.name: index for index, path in enumerate(discover_slides(), start=1)}
    module_bounds = {
        "ec01-images": ("29_check01_alt-types.py", "31_check01_alt-exemples.py"),
        "ec02-page-title": ("32_check02_titre-page.py",),
        "ec03-headings": ("33_check03_titres-hierarchie.py", "34_check03_titres-outils.py"),
        "ec04-contrast": ("35_check04_contraste-principe.py", "36_check04_contraste-outils.py"),
        "ec05-skiplinks": ("37_check05_lien-evitement.py",),
        "ec06-keyboard-focus": ("38_navigation-clavier-ouverture.py", "41_mission-clavier.py"),
        "ec07-language": ("42_check07_langue.py",),
        "ec08-zoom": ("43_check08_zoom.py",),
        "ec09-captions": ("44_check09_sous-titres-principe.py", "45_check09_sous-titres-auto.py"),
        "ec10-transcript": ("46_check10_transcriptions.py",),
        "ec11-audio-description": ("47_check11_audiodescription.py",),
        "ec12-form-labels": ("48_check12_etiquettes-principe.py", "50_check12_etiquettes-groupes.py"),
        "ec13-required-errors": ("51_check13_champs-obligatoires.py",),
    }
    expected = {}
    for page_id, bounds in module_bounds.items():
        first = positions[bounds[0]]
        last = positions[bounds[-1]]
        expected[page_id] = str(first) if first == last else f"{first}-{last}"

    assert {page["id"]: page["slides"] for page in EXERCICE_PAGES} == expected
