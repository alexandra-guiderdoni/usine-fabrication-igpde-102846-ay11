"""Cohérence finale du pack avec le TP Sami organisé en étapes."""

import sys
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from config import load_formation_config

PACK = ROOT / load_formation_config()["livrables"]
ADMIN = PACK / "Livrables-Formateur" / "documents-administratifs-igpde"
TP_WORD = PACK / "Livrables-Stagiaires" / "tp-word-igpde"
PROJECTIONS = PACK / "Livrables-Stagiaires" / "supports-projections"

LEGACY_PATTERNS = (
    "21 " + "critères",
    "21 bonnes " + "pratiques",
    "21 " + "criteres",
    "25 " + "min",
    "25 " + "minutes",
    "sans " + "checklist",
    "sans " + "filet",
    "rapport " + "trimestriel",
    "quiz " + "final",
    "138 " + "slides",
    "102" + "638",
)


def _docx_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join(paragraphs + cells)


def _active_text_files() -> list[Path]:
    files = [
        ROOT / "AGENTS.md",
        ROOT / "README.md",
        PACK / "README.md",
        ROOT / "fiche-pratique" / "README.md",
        ROOT / "fiche-pratique" / "memo-word.md",
        ROOT / "fiche-pratique" / "memo-libreoffice-writer.md",
        ROOT / "_source" / "exercice-sami-spec.md",
        ROOT / "_source" / "exercice-sami-diff.md",
        ROOT / "_source" / "exercice-sami-matrice.yml",
        ROOT / "_source" / "checklist-accessibilite-bureautique.md",
        ROOT / "scripts" / "generate_exercice_sami.py",
        ROOT / "scripts" / "exercice_sami_matrice.py",
        ROOT / "scripts" / "sami_slide_data.py",
        ROOT / "scripts" / "slides" / "README.md",
        PACK
        / "Livrables-Formateur"
        / "_alex"
        / "formation-102846-octobre-2026-bureautique.md",
    ]
    files.extend(sorted((ROOT / "scripts" / "slides").glob("*.py")))
    files.extend(sorted((ROOT / "tests").glob("*.py")))
    return files


def test_anciens_contrats_absents_des_sources_actives():
    findings = []
    for path in _active_text_files():
        content = path.read_text(encoding="utf-8").casefold()
        if path == ROOT / "AGENTS.md":
            content = content.replace("ex-" + "102" + "638", "")
        for pattern in LEGACY_PATTERNS:
            if pattern.casefold() in content:
                findings.append(f"{path.relative_to(ROOT)}: {pattern}")

    for path in sorted(ADMIN.glob("*.docx")):
        content = _docx_text(path).casefold()
        for pattern in LEGACY_PATTERNS:
            if pattern.casefold() in content:
                findings.append(f"{path.relative_to(ROOT)}: {pattern}")

    assert findings == []


def test_pack_inventorie_et_contient_les_livrables_du_tp_word():
    expected = {
        "tp-doc-inaccessible.docx",
        "tp-doc-aide-correction.docx",
        "tp-doc-accessible.docx",
        "checklist-accessibilite-bureautique.docx",
        "checklist-accessibilite-bureautique.pdf",
        "memo-word-accessibilite.pdf",
        "memo-libreoffice-writer-accessibilite.pdf",
    }
    readme = (PACK / "README.md").read_text(encoding="utf-8")
    assert expected <= {path.name for path in TP_WORD.iterdir()}
    assert all(filename in readme for filename in expected)
    assert "Au début du TP" in readme
    assert "À la fin seulement" in readme


def test_anciennes_valeurs_de_contraste_ne_sont_plus_un_echec_actif():
    narrative_files = [
        ROOT / "README.md",
        PACK / "README.md",
        ROOT / "fiche-pratique" / "README.md",
        ROOT / "fiche-pratique" / "memo-word.md",
        ROOT / "fiche-pratique" / "memo-libreoffice-writer.md",
        PACK
        / "Livrables-Formateur"
        / "_alex"
        / "formation-102846-octobre-2026-bureautique.md",
    ]
    narrative = "\n".join(path.read_text(encoding="utf-8") for path in narrative_files)
    narrative += "\n" + "\n".join(_docx_text(path) for path in ADMIN.glob("*.docx"))
    assert "#767" + "676" not in narrative
    assert "4," + "48" not in narrative

    matrix = (ROOT / "_source" / "exercice-sami-matrice.yml").read_text(
        encoding="utf-8"
    )
    assert ("#767" + "676 sur blanc n'est pas présentée comme un échec") in matrix


def test_deroule_et_fiche_technique_sont_alignes():
    deroule = _docx_text(ADMIN / "Derped-deroule-pedagogique-102846-v2.docx")
    technique = _docx_text(ADMIN / "102846FiTechn-v2.docx")

    assert "10h30 - 12h00" in deroule
    assert "12h00 - 12h15" in deroule
    assert "slides 54 à 79" in deroule
    assert "slide 80" in deroule
    assert "méthode en cinq étapes" in deroule
    assert "TP guidé en cinq étapes" in deroule
    assert "checklist renseignée après chaque étape" in deroule
    assert "finaliser, vérifier, exporter et contrôler" in deroule
    for projection in (
        "00-introduction-et-idees-recues.pptx, slides 1 à 14",
        "01-accessibilite-numerique-et-cadre-legal.pptx, slides 1 à 39",
        "02-documents-bureautiques-accessibles-tp.pptx, slides 1 à 26",
        "02-documents-bureautiques-accessibles-tp.pptx, slide 27",
        "03-web-accessible-tp.pptx, slides 1 à 30",
        "04-reseaux-sociaux-accessibles.pptx, slides 1 à 55",
    ):
        assert projection in deroule
    assert "PAC 24.4.4.0" in technique
    assert "disponible sur les postes" in technique
    assert "Acrobat Pro" in technique


def test_projection_word_utilise_les_etapes():
    path = PROJECTIONS / "02-documents-bureautiques-accessibles-tp.pptx"
    presentation = Presentation(path)
    image_titles = "\n".join(
        shape._element.nvPicPr.cNvPr.get("title", "")
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    presenter_notes = "\n".join(
        slide.notes_slide.notes_text_frame.text for slide in presentation.slides
    )

    assert len(presentation.slides) == 27
    assert all(len(slide.shapes) == 1 for slide in presentation.slides)
    assert "Étape 1" in image_titles
    assert "Étape 5" in image_titles
    assert "Lire la transcription" in presenter_notes
    assert "Lire le discours oral" in presenter_notes
    assert "Station" not in image_titles
    assert "Station" not in presenter_notes


def test_index_des_slides_couvre_la_fin_de_la_partie_2():
    index = (ROOT / "scripts" / "slides" / "README.md").read_text(encoding="utf-8")
    for module in (
        "05a_quiz-flash-reponse.py",
        "26d_checklist-station-5-signales.py",
        "26e_synthese-matinee.py",
    ):
        assert module in index
