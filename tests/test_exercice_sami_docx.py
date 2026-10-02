"""Tests d'intégration XML des trois DOCX de l'exercice Sami."""

# PDG-LARGE-FILE-JUSTIFICATION: suite d'intégration unique qui compare les
# trois variantes DOCX et leurs contrats OOXML station par station.

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree
from PIL import Image

from exercice_sami_matrice import load_sami_matrix
from generate_exercice_sami import build_accessible, build_inaccessible


PROJECT_ROOT = Path(__file__).parent.parent
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _archive_text(path: Path, member: str) -> str:
    with ZipFile(path) as archive:
        return archive.read(member).decode("utf-8")


def _archive_members(path: Path) -> set[str]:
    with ZipFile(path) as archive:
        return set(archive.namelist())


def _comment_anchor(path: Path, marker: str) -> dict[str, object]:
    """Retourne le paragraphe réellement couvert par un commentaire."""
    namespaces = {"w": WORD_NS, "r": OFFICE_REL_NS}
    with ZipFile(path) as archive:
        comments_root = etree.fromstring(archive.read("word/comments.xml"))
        document_root = etree.fromstring(archive.read("word/document.xml"))

    comment = next(
        item
        for item in comments_root.xpath("//w:comment", namespaces=namespaces)
        if marker in "".join(item.xpath(".//w:t/text()", namespaces=namespaces))
    )
    comment_id = comment.get(qn("w:id"))
    paragraph = document_root.xpath(
        f'//w:p[.//w:commentRangeStart[@w:id="{comment_id}"]]',
        namespaces=namespaces,
    )[0]
    return {
        "text": "".join(paragraph.xpath(".//w:t/text()", namespaces=namespaces)),
        "drawing": bool(paragraph.xpath(".//w:drawing", namespaces=namespaces)),
        "hyperlink": bool(paragraph.xpath(".//w:hyperlink", namespaces=namespaces)),
    }


def _comment_text(path: Path, marker: str) -> str:
    namespaces = {"w": WORD_NS}
    with ZipFile(path) as archive:
        comments_root = etree.fromstring(archive.read("word/comments.xml"))
    return next(
        "".join(item.xpath(".//w:t/text()", namespaces=namespaces))
        for item in comments_root.xpath("//w:comment", namespaces=namespaces)
        if marker in "".join(item.xpath(".//w:t/text()", namespaces=namespaces))
    )


def _embedded_asset_count(path: Path, asset_path: Path) -> int:
    """Compte les médias embarqués identiques à un fichier source."""
    expected = asset_path.read_bytes()
    with ZipFile(path) as archive:
        return sum(
            archive.read(member) == expected
            for member in archive.namelist()
            if member.startswith("word/media/")
        )


def _external_hyperlinks(path: Path) -> list[dict[str, object]]:
    """Retourne les liens et leur mise en forme depuis leur nœud OOXML."""
    document_namespaces = {"w": WORD_NS, "r": OFFICE_REL_NS}
    relationship_namespaces = {"pr": PACKAGE_REL_NS}
    with ZipFile(path) as archive:
        document_root = etree.fromstring(archive.read("word/document.xml"))
        relationships_root = etree.fromstring(
            archive.read("word/_rels/document.xml.rels")
        )
    targets = {
        item.get("Id"): item.get("Target")
        for item in relationships_root.xpath(
            "//pr:Relationship", namespaces=relationship_namespaces
        )
        if item.get("Type", "").endswith("/hyperlink")
    }
    return [
        {
            "target": targets.get(link.get(qn("r:id"))),
            "text": "".join(
                link.xpath(".//w:t/text()", namespaces=document_namespaces)
            ),
            "colors": link.xpath(
                ".//w:rPr/w:color/@w:val", namespaces=document_namespaces
            ),
            "underlines": link.xpath(
                ".//w:rPr/w:u/@w:val", namespaces=document_namespaces
            ),
        }
        for link in document_root.xpath("//w:hyperlink", namespaces=document_namespaces)
    ]


def _linear_channel(value: int) -> float:
    channel = value / 255
    return channel / 12.92 if channel <= 0.04045 else ((channel + 0.055) / 1.055) ** 2.4


def _contrast_ratio(foreground: str, background: str = "FFFFFF") -> float:
    """Calcule le contraste WCAG de deux couleurs hexadécimales."""
    colors = []
    for value in (foreground, background):
        rgb = tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))
        colors.append(
            0.2126 * _linear_channel(rgb[0])
            + 0.7152 * _linear_channel(rgb[1])
            + 0.0722 * _linear_channel(rgb[2])
        )
    light, dark = sorted(colors, reverse=True)
    return (light + 0.05) / (dark + 0.05)


def _table_containing(document: Document, expected_text: str):
    return next(
        table
        for table in document.tables
        if any(expected_text in cell.text for row in table.rows for cell in row.cells)
    )


def _table_tokens(table) -> list[str]:
    tokens = []
    seen_cells = set()
    for row in table.rows:
        for cell in row.cells:
            if cell._tc in seen_cells:
                continue
            seen_cells.add(cell._tc)
            tokens.extend(token for token in cell.text.splitlines() if token)
    return tokens


def _direct_run_language(run):
    run_properties = run._r.rPr
    if run_properties is None:
        return None
    language = run_properties.find(qn("w:lang"))
    return None if language is None else language.get(qn("w:val"))


def _style_language(document, style_name):
    run_properties = document.styles[style_name].element.rPr
    if run_properties is None:
        return None
    language = run_properties.find(qn("w:lang"))
    return None if language is None else language.get(qn("w:val"))


def _default_document_language(document):
    styles = document.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    run_defaults = defaults.find(qn("w:rPrDefault")) if defaults is not None else None
    run_properties = (
        run_defaults.find(qn("w:rPr")) if run_defaults is not None else None
    )
    language = run_properties.find(qn("w:lang")) if run_properties is not None else None
    return None if language is None else language.get(qn("w:val"))


def _body_runs(document):
    for paragraph in document.paragraphs:
        yield from paragraph.runs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield from paragraph.runs


def _useful_document_runs(document):
    """Parcourt le corps et les informations utiles d'en-tête et de pied."""
    yield from _body_runs(document)
    for section in document.sections:
        for story in (section.header, section.footer):
            for paragraph in story.paragraphs:
                yield from paragraph.runs


def _body_signature(document):
    return (
        [paragraph.text for paragraph in document.paragraphs],
        [_table_tokens(table) for table in document.tables],
    )


def _assert_detectable_station_families_are_corrected(path):
    """Oracle indépendant couvrant un signal automatisable par station."""
    document = Document(path)
    introduction = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "Introduction"
    )
    assert introduction.style.name == "Heading 1", "station-1"

    links = _external_hyperlinks(path)
    assert links and all(
        "cliquez ici" not in link["text"].casefold() for link in links
    ), "station-2"

    contrast_run = next(
        paragraph.runs[0]
        for paragraph in document.paragraphs
        if paragraph.text == "Information complémentaire : résultats provisoires."
    )
    assert _contrast_ratio(str(contrast_run.font.color.rgb)) >= 4.5, "station-3"

    normal = document.styles["Normal"]
    assert normal.font.size.pt >= 12, "station-4"
    assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT, "station-4"

    assert document.core_properties.title == "Rendre un document Word accessible", (
        "station-5"
    )
    assert document.core_properties.author == "Sami Dupont", "station-5"
    assert document.core_properties.language == "fr-FR", "station-5"


def _paragraph_region(path: Path, start: str, end: str):
    paragraphs = Document(path).paragraphs
    start_index = next(
        index for index, paragraph in enumerate(paragraphs) if paragraph.text == start
    )
    end_index = next(
        index
        for index, paragraph in enumerate(paragraphs[start_index:], start_index)
        if paragraph.text == end
    )
    return paragraphs[start_index : end_index + 1]


def test_le_docx_guide_consomme_directement_le_libelle_de_la_matrice(tmp_path):
    matrix = deepcopy(load_sami_matrix())
    control = next(item for item in matrix["controles"] if item["id"] == "P-01")
    control["intitule"] = "Libellé injecté depuis la matrice"

    output = build_inaccessible(
        PROJECT_ROOT / "_assets" / "graphique-inaccessible.png",
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )

    comments = _archive_text(output, "word/comments.xml")
    assert "Libellé injecté depuis la matrice" in comments


def test_le_corrige_porte_un_titre_principal_et_le_controle_p01(tmp_path):
    matrix = deepcopy(load_sami_matrix())
    control = next(item for item in matrix["controles"] if item["id"] == "P-01")
    control["intitule"] = "Titre principal piloté par la matrice"

    output = build_accessible(
        PROJECT_ROOT / "_assets" / "graphique-accessible.png",
        matrix=matrix,
        output_dir=tmp_path,
    )

    document = Document(output)
    title = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "Rendre un document Word accessible"
    )
    assert title.style.name == "Title"
    assert "Titre principal piloté par la matrice" in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )


def test_la_hierarchie_p02_est_fautive_au_depart_et_corrigee(tmp_path):
    matrix = load_sami_matrix()
    control = next(item for item in matrix["controles"] if item["id"] == "P-02")
    heading_text = f"{control['id']} - {control['intitule']}"
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"

    inaccessible = build_inaccessible(
        chart_bad,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(chart_good, matrix=matrix, output_dir=tmp_path)

    bad_doc = Document(inaccessible)
    good_doc = Document(corrected)

    def style_for(document, text):
        return next(
            paragraph.style.name
            for paragraph in document.paragraphs
            if paragraph.text == text
        )

    station_title = next(
        block["titre"] for block in matrix["sequence"] if block["id"] == "station-1"
    )
    assert style_for(bad_doc, "Rendre un document Word accessible") == "Normal"
    assert style_for(bad_doc, station_title) == "Heading 1"
    assert style_for(bad_doc, heading_text) == "Heading 4"
    assert style_for(good_doc, heading_text) == "Heading 2"

    bad_heading = next(
        paragraph for paragraph in bad_doc.paragraphs if paragraph.text == heading_text
    )
    good_heading = next(
        paragraph for paragraph in good_doc.paragraphs if paragraph.text == heading_text
    )
    bad_numbering = bad_heading._p.pPr.numPr
    good_numbering = good_heading._p.pPr.numPr
    assert bad_numbering is not None
    assert good_numbering is not None
    assert bad_numbering.ilvl.val == 3
    assert good_numbering.ilvl.val == 1
    for sibling in ("P-01", "P-03", "P-04", "P-05"):
        sibling_control = next(
            item for item in matrix["controles"] if item["id"] == sibling
        )
        sibling_text = f"{sibling} - {sibling_control['intitule']}"
        assert style_for(bad_doc, sibling_text) == "Heading 2"
        assert style_for(good_doc, sibling_text) == "Heading 2"


def test_les_titres_hors_station_un_restant_ne_creent_pas_d_occurrence_cachee(
    tmp_path,
):
    matrix = load_sami_matrix()
    output = build_inaccessible(
        PROJECT_ROOT / "_assets" / "graphique-inaccessible.png",
        icon_path=PROJECT_ROOT / "_assets" / "icone-enveloppe.png",
        organigramme_path=PROJECT_ROOT / "_assets" / "organigramme.png",
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    document = Document(output)
    expected_styles = {
        "Introduction": "Heading 1",
        "Résultats du trimestre": "Heading 2",
        "Détail par canal": "Heading 3",
        "Répartition par service": "Heading 3",
        "ANNEXES - ACCESSIBILITE": "Heading 3",
    }

    for text, style_name in expected_styles.items():
        paragraph = next(item for item in document.paragraphs if item.text == text)
        assert paragraph.style.name == style_name


def test_p03_guide_le_sommaire_manuel_et_le_corrige_le_rend_actualisable(
    tmp_path,
):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    inaccessible = build_inaccessible(
        chart_bad,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(chart_good, matrix=matrix, output_dir=tmp_path)

    inaccessible_xml = _archive_text(inaccessible, "word/document.xml")
    guided_xml = _archive_text(guided, "word/document.xml")
    corrected_xml = _archive_text(corrected, "word/document.xml")
    comments = _archive_text(guided, "word/comments.xml")

    assert "TOC \\o" not in inaccessible_xml
    assert "TOC \\o" not in guided_xml
    assert "TOC \\o" in corrected_xml
    assert comments.count("Document — P-03") == 1


def test_p04_remplace_les_marqueurs_saisis_par_des_listes_natives(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    guided = build_inaccessible(
        chart_bad,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(chart_good, matrix=matrix, output_dir=tmp_path)

    guided_doc = Document(guided)
    corrected_doc = Document(corrected)
    bad_bullet = next(
        paragraph
        for paragraph in guided_doc.paragraphs
        if "Augmenter le trafic de 10 %" in paragraph.text
    )
    good_bullet = next(
        paragraph
        for paragraph in corrected_doc.paragraphs
        if paragraph.text == "Augmenter le trafic de 10 %"
    )
    bad_number = next(
        paragraph
        for paragraph in guided_doc.paragraphs
        if "Refonte de la page d'accueil" in paragraph.text
    )
    good_number = next(
        paragraph
        for paragraph in corrected_doc.paragraphs
        if paragraph.text == "Refonte de la page d'accueil"
    )
    comments = _archive_text(guided, "word/comments.xml")

    assert bad_bullet.style.name == "Normal"
    assert bad_number.style.name == "Normal"
    assert good_bullet.style.name == "List Bullet"
    assert good_number.style.name == "List Number"
    assert comments.count("P-04 -") == 1


def test_p05_remplace_les_artifices_par_des_fonctions_de_mise_en_page(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    guided = build_inaccessible(
        chart_bad,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(chart_good, matrix=matrix, output_dir=tmp_path)

    start = "Mise en page robuste"
    end = "Fin de la zone Mise en page robuste."
    bad_region = _paragraph_region(guided, start, end)
    good_region = _paragraph_region(corrected, start, end)
    bad_xml = "".join(paragraph._p.xml for paragraph in bad_region)
    good_xml = "".join(paragraph._p.xml for paragraph in good_region)
    corrected_xml = _archive_text(corrected, "word/document.xml")
    comments = _archive_text(guided, "word/comments.xml")

    assert sum(not paragraph.text for paragraph in bad_region) >= 4
    assert any("  " in paragraph.text for paragraph in bad_region)
    assert "<w:tab" in bad_xml
    assert "<w:br" in bad_xml
    assert "w:pageBreakBefore" not in bad_xml
    assert 'w:num="2"' not in _archive_text(guided, "word/document.xml")

    assert not any("  " in paragraph.text for paragraph in good_region)
    assert "<w:tab" not in good_xml
    assert "<w:br" not in good_xml
    assert "w:pageBreakBefore" in good_xml
    assert "<w:spacing" in good_xml
    assert "<w:ind" in good_xml
    assert 'w:num="2"' in corrected_xml
    assert comments.count("P-05 -") == 1

    bad_text = re.sub(r"\s+", " ", " ".join(p.text for p in bad_region)).strip()
    good_text = re.sub(r"\s+", " ", " ".join(p.text for p in good_region)).strip()
    assert bad_text == good_text


def test_le_titre_de_station_est_lu_dans_la_matrice(tmp_path):
    matrix = deepcopy(load_sami_matrix())
    station = next(item for item in matrix["sequence"] if item["id"] == "station-1")
    station["titre"] = "Titre de station injecté depuis la matrice"

    output = build_accessible(
        PROJECT_ROOT / "_assets" / "graphique-accessible.png",
        matrix=matrix,
        output_dir=tmp_path,
    )

    heading = next(
        paragraph
        for paragraph in Document(output).paragraphs
        if paragraph.text == "Titre de station injecté depuis la matrice"
    )
    assert heading.style.name == "Heading 1"


def test_les_builders_lisent_les_noms_des_trois_versions_dans_la_matrice(tmp_path):
    matrix = deepcopy(load_sami_matrix())
    matrix["identite_editoriale"]["versions"] = {
        "inaccessible": "depart.docx",
        "avec_pistes": "guide.docx",
        "corrigee": "corrige.docx",
    }
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"

    inaccessible = build_inaccessible(chart_bad, matrix=matrix, output_dir=tmp_path)
    guided = build_inaccessible(
        chart_bad, with_guidance=True, matrix=matrix, output_dir=tmp_path
    )
    corrected = build_accessible(chart_good, matrix=matrix, output_dir=tmp_path)

    assert [inaccessible.name, guided.name, corrected.name] == [
        "depart.docx",
        "guide.docx",
        "corrige.docx",
    ]


def test_le_guide_contient_une_piste_par_occurrence_de_la_station_un(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    inaccessible = build_inaccessible(
        chart_bad,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(
        PROJECT_ROOT / "_assets" / "graphique-accessible.png",
        matrix=matrix,
        output_dir=tmp_path,
    )
    comments = _archive_text(guided, "word/comments.xml")
    controls = [
        control for control in matrix["controles"] if control["station"] == "station-1"
    ]

    for control in controls:
        assert comments.count(f"{control['id']} -") == control["occurrences_attendues"]
        assert control["regle"] in comments
        assert control["action_attendue"] in comments
    assert sum(comments.count(f"{control['id']} -") for control in controls) == 5
    assert "word/comments.xml" not in _archive_members(inaccessible)
    assert "word/comments.xml" not in _archive_members(corrected)


def test_les_trois_versions_partagent_le_contenu_pedagogique_de_la_station_un(
    tmp_path,
):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    paths = [
        build_inaccessible(
            chart_bad,
            output_name="inaccessible.docx",
            matrix=matrix,
            output_dir=tmp_path,
        ),
        build_inaccessible(
            chart_bad,
            with_guidance=True,
            output_name="guide.docx",
            matrix=matrix,
            output_dir=tmp_path,
        ),
        build_accessible(
            PROJECT_ROOT / "_assets" / "graphique-accessible.png",
            matrix=matrix,
            output_dir=tmp_path,
        ),
    ]
    documents = [Document(path) for path in paths]
    controls = [
        control for control in matrix["controles"] if control["station"] == "station-1"
    ]
    expected_paragraphs = {"Rendre un document Word accessible"}
    for control in controls:
        expected_paragraphs.update(
            {
                f"{control['id']} - {control['intitule']}",
                f"Problème : {control['defaut']}",
                f"Pourquoi : {control['impact']}",
                f"Règle : {control['regle']}",
                f"Dans Word : {control['procedure_word']}",
                f"Dans LibreOffice Writer : {control['procedure_writer']}",
                f"À faire : {control['action_attendue']}",
                f"Preuve : {control['preuve']['attendu']}",
            }
        )

    for document in documents:
        paragraphs = {paragraph.text for paragraph in document.paragraphs}
        assert expected_paragraphs <= paragraphs
        normalized = re.sub(
            r"\s+",
            " ",
            " ".join(paragraph.text for paragraph in document.paragraphs),
        )
        for content in (
            "Augmenter le trafic de 10 %",
            "Refonte de la page d'accueil",
            "Retrait du paragraphe aligné avec",
            "Ligne principale Suite sur un paragraphe distinct",
            "Début de la page suivante",
            "Colonne gauche : structure Colonne droite : navigation",
        ):
            assert content in normalized


def test_les_pistes_documentaires_sont_ancrees_dans_le_corps(tmp_path):
    matrix = load_sami_matrix()
    output = build_inaccessible(
        PROJECT_ROOT / "_assets" / "graphique-inaccessible.png",
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )

    members = _archive_members(output)
    for member in members:
        if member.startswith("word/header") or member.startswith("word/footer"):
            xml = _archive_text(output, member)
            assert "commentRangeStart" not in xml
            assert "commentReference" not in xml

    comments = _archive_text(output, "word/comments.xml")
    assert "Document — P-03" in comments
    assert "Document — P-11" in comments


def test_p06_demande_une_alternative_redigee_puis_l_applique(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    icon = PROJECT_ROOT / "_assets" / "icone-enveloppe.png"
    inaccessible = build_inaccessible(
        chart_bad,
        icon_path=icon,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        icon_path=icon,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(
        chart_good,
        icon_path=icon,
        matrix=matrix,
        output_dir=tmp_path,
    )

    control = next(item for item in matrix["controles"] if item["id"] == "P-06")
    station_title = next(
        item["titre"] for item in matrix["sequence"] if item["id"] == "station-2"
    )
    heading_text = f"{control['id']} - {control['intitule']}"
    for path in (inaccessible, guided, corrected):
        document = Document(path)
        station = next(p for p in document.paragraphs if p.text == station_title)
        heading = next(p for p in document.paragraphs if p.text == heading_text)
        assert station.style.name == "Heading 1"
        assert heading.style.name == "Heading 2"
        assert f"Problème : {control['defaut']}" in {
            p.text for p in document.paragraphs
        }

    bad_doc_pr = Document(inaccessible).inline_shapes[0]._inline.find(qn("wp:docPr"))
    good_doc_pr = Document(corrected).inline_shapes[0]._inline.find(qn("wp:docPr"))
    assert not bad_doc_pr.get("descr")
    assert good_doc_pr.get("descr") == "Contact par courriel."

    comments = _archive_text(guided, "word/comments.xml")
    assert comments.count("P-06 -") == 1
    assert control["regle"] in comments
    assert control["action_attendue"] in comments


def test_p07_associe_l_image_complexe_a_une_description_detaillee(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    organigramme = PROJECT_ROOT / "_assets" / "organigramme.png"
    inaccessible = build_inaccessible(
        chart_bad,
        organigramme_path=organigramme,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        organigramme_path=organigramme,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(
        chart_good,
        organigramme_path=organigramme,
        matrix=matrix,
        output_dir=tmp_path,
    )

    control = next(item for item in matrix["controles"] if item["id"] == "P-07")
    assert "P-07" in matrix["identite_editoriale"]["transformations_autorisees"]
    assert control["transformations_editoriales"] == [
        "description_image_complexe_vers_texte_adjacent"
    ]
    heading_text = f"{control['id']} - {control['intitule']}"
    for path in (inaccessible, guided, corrected):
        document = Document(path)
        heading = next(p for p in document.paragraphs if p.text == heading_text)
        assert heading.style.name == "Heading 2"

    bad_alts = [
        shape._inline.find(qn("wp:docPr")).get("descr")
        for shape in Document(inaccessible).inline_shapes
    ]
    good_alts = [
        shape._inline.find(qn("wp:docPr")).get("descr")
        for shape in Document(corrected).inline_shapes
    ]
    assert bad_alts.count("image.png") == 1
    assert (
        good_alts.count(
            "Organigramme de la Direction des affaires juridiques "
            "(description ci-dessous)."
        )
        == 1
    )

    description = (
        "La Direction des affaires juridiques comprend 4 bureaux : "
        "le Bureau du droit public, le Bureau du droit social, "
        "le Bureau de la communication et le Bureau des affaires "
        "internationales. Chaque bureau est rattaché directement à la direction."
    )
    assert description not in {p.text for p in Document(inaccessible).paragraphs}
    assert description in {p.text for p in Document(corrected).paragraphs}

    comments = _archive_text(guided, "word/comments.xml")
    assert comments.count("P-07 -") == 1
    assert control["regle"] in comments


def test_p08_marque_l_image_redondante_comme_decorative(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    icon = PROJECT_ROOT / "_assets" / "icone-enveloppe.png"
    inaccessible = build_inaccessible(
        chart_bad,
        icon_path=icon,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        icon_path=icon,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(
        chart_good,
        icon_path=icon,
        matrix=matrix,
        output_dir=tmp_path,
    )

    control = next(item for item in matrix["controles"] if item["id"] == "P-08")
    heading_text = f"{control['id']} - {control['intitule']}"
    for path in (inaccessible, guided, corrected):
        document = Document(path)
        heading = next(p for p in document.paragraphs if p.text == heading_text)
        assert heading.style.name == "Heading 2"
        assert f"Dans LibreOffice Writer : {control['procedure_writer']}" in {
            p.text for p in document.paragraphs
        }

    bad_alts = [
        shape._inline.find(qn("wp:docPr")).get("descr")
        for shape in Document(inaccessible).inline_shapes
    ]
    assert bad_alts.count("E-mail") == 1
    decorative_marker = "C183D7F6-B498-43B3-948B-1728B52AA6E4"
    assert decorative_marker not in _archive_text(inaccessible, "word/document.xml")
    assert _archive_text(corrected, "word/document.xml").count(decorative_marker) == 1

    contact = (
        "Pour toute question, contactez-nous par  e-mail pour plus d'informations."
    )
    assert contact in {p.text for p in Document(inaccessible).paragraphs}
    assert contact in {p.text for p in Document(corrected).paragraphs}

    comments = _archive_text(guided, "word/comments.xml")
    assert comments.count("P-08 -") == 1
    assert control["regle"] in comments


def test_p09_remplace_l_image_de_texte_par_un_texte_selectionnable(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    text_image = PROJECT_ROOT / "_assets" / "texte-image.png"
    inaccessible = build_inaccessible(
        chart_bad,
        texte_image_path=text_image,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        texte_image_path=text_image,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(
        chart_good,
        texte_image_path=text_image,
        matrix=matrix,
        output_dir=tmp_path,
    )

    control = next(item for item in matrix["controles"] if item["id"] == "P-09")
    heading_text = f"{control['id']} - {control['intitule']}"
    for path in (inaccessible, guided, corrected):
        document = Document(path)
        heading = next(p for p in document.paragraphs if p.text == heading_text)
        assert heading.style.name == "Heading 2"

    real_text = (
        "Avis important : les indicateurs du T2 2025 "
        "seront transmis avant le 15 septembre 2025."
    )
    assert real_text not in {p.text for p in Document(inaccessible).paragraphs}
    assert real_text not in {p.text for p in Document(guided).paragraphs}
    assert real_text in {p.text for p in Document(corrected).paragraphs}
    assert _embedded_asset_count(inaccessible, text_image) == 1
    assert _embedded_asset_count(guided, text_image) == 1
    assert _embedded_asset_count(corrected, text_image) == 0

    comments = _archive_text(guided, "word/comments.xml")
    assert comments.count("P-09 -") == 1
    assert control["action_attendue"] in comments


def test_p10_conserve_la_destination_et_rend_le_lien_autonome(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    inaccessible = build_inaccessible(
        chart_bad,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(chart_good, matrix=matrix, output_dir=tmp_path)

    control = next(item for item in matrix["controles"] if item["id"] == "P-10")
    expected_proof = (
        "Destination inchangée, libellé explicite et identification visuelle vérifiées."
    )
    assert control["preuve"]["attendu"] == expected_proof
    heading_text = f"{control['id']} - {control['intitule']}"
    for path in (inaccessible, guided, corrected):
        document = Document(path)
        heading = next(p for p in document.paragraphs if p.text == heading_text)
        assert heading.style.name == "Heading 2"
        assert f"Preuve : {expected_proof}" in {
            paragraph.text for paragraph in document.paragraphs
        }
        links = _external_hyperlinks(path)
        assert [link["target"] for link in links].count(
            "https://example.org/annexes-rapport-t1-2025.pdf"
        ) == 1

    bad_xml = _archive_text(inaccessible, "word/document.xml")
    assert "cliquez ici" in bad_xml
    corrected_link = next(
        link
        for link in _external_hyperlinks(corrected)
        if link["target"] == "https://example.org/annexes-rapport-t1-2025.pdf"
    )
    assert corrected_link["text"] == (
        "Consulter les annexes du rapport T1 2025 (PDF, 1,2 Mo, français)"
    )
    assert corrected_link["colors"] == ["0000FF"]
    assert corrected_link["underlines"] == ["single"]

    comments = _archive_text(guided, "word/comments.xml")
    assert comments.count("P-10 -") == 1
    assert control["regle"] in comments


def test_p11_est_explique_sans_filigrane_illisible(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    inaccessible = build_inaccessible(
        chart_bad,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(chart_good, matrix=matrix, output_dir=tmp_path)

    control = next(item for item in matrix["controles"] if item["id"] == "P-11")
    heading_text = f"{control['id']} - {control['intitule']}"
    for path in (inaccessible, guided, corrected):
        document = Document(path)
        heading = next(p for p in document.paragraphs if p.text == heading_text)
        assert heading.style.name == "Heading 2"

    assert "CONFIDENTIEL" not in _archive_text(inaccessible, "word/header1.xml")
    assert "CONFIDENTIEL" not in _archive_text(guided, "word/header1.xml")
    assert "CONFIDENTIEL" not in _archive_text(corrected, "word/header1.xml")
    assert "Document confidentiel" not in {
        p.text for p in Document(inaccessible).paragraphs
    }
    assert "Document confidentiel" not in {p.text for p in Document(guided).paragraphs}
    assert "Document confidentiel" in {p.text for p in Document(corrected).paragraphs}

    comments = _archive_text(guided, "word/comments.xml")
    assert comments.count("Document — P-11 -") == 1
    assert "Document — Critère 18" not in comments
    assert "commentReference" not in _archive_text(guided, "word/header1.xml")


def test_la_station_deux_est_ordonnee_et_pilotee_par_la_matrice(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    assets = {
        "icon_path": PROJECT_ROOT / "_assets" / "icone-enveloppe.png",
        "organigramme_path": PROJECT_ROOT / "_assets" / "organigramme.png",
        "texte_image_path": PROJECT_ROOT / "_assets" / "texte-image.png",
    }
    paths = [
        build_inaccessible(
            chart_bad,
            output_name="inaccessible.docx",
            matrix=matrix,
            output_dir=tmp_path,
            **assets,
        ),
        build_inaccessible(
            chart_bad,
            with_guidance=True,
            output_name="guide.docx",
            matrix=matrix,
            output_dir=tmp_path,
            **assets,
        ),
        build_accessible(
            PROJECT_ROOT / "_assets" / "graphique-accessible.png",
            matrix=matrix,
            output_dir=tmp_path,
            **assets,
        ),
    ]
    controls = [item for item in matrix["controles"] if item["station"] == "station-2"]
    station_title = next(
        item["titre"] for item in matrix["sequence"] if item["id"] == "station-2"
    )

    for path in paths:
        paragraphs = Document(path).paragraphs
        texts = [paragraph.text for paragraph in paragraphs]
        assert texts.count(station_title) == 1
        heading_indexes = []
        for control in controls:
            heading_text = f"{control['id']} - {control['intitule']}"
            heading_indexes.append(texts.index(heading_text))
            for expected in (
                f"Problème : {control['defaut']}",
                f"Pourquoi : {control['impact']}",
                f"Règle : {control['regle']}",
                f"Dans Word : {control['procedure_word']}",
                f"Dans LibreOffice Writer : {control['procedure_writer']}",
                f"À faire : {control['action_attendue']}",
                f"Preuve : {control['preuve']['attendu']}",
            ):
                assert expected in texts
        assert heading_indexes == sorted(heading_indexes)

    comments = _archive_text(paths[1], "word/comments.xml")
    for control in controls:
        assert comments.count(f"{control['id']} -") == control["occurrences_attendues"]
        assert control["regle"] in comments
        assert control["action_attendue"] in comments
    assert sum(comments.count(f"{control['id']} -") for control in controls) == 6
    for legacy_label in (
        "Critère 8 -",
        "Critère 9 -",
        "Critère 10 -",
        "Critère 18 -",
        "Critère 20 -",
    ):
        assert legacy_label not in comments


def test_la_station_deux_consomme_un_libelle_injecte_depuis_la_matrice(tmp_path):
    matrix = deepcopy(load_sami_matrix())
    control = next(item for item in matrix["controles"] if item["id"] == "P-06")
    control["intitule"] = "Libellé P-06 injecté depuis la matrice"
    output = build_inaccessible(
        PROJECT_ROOT / "_assets" / "graphique-inaccessible.png",
        icon_path=PROJECT_ROOT / "_assets" / "icone-enveloppe.png",
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )

    texts = {paragraph.text for paragraph in Document(output).paragraphs}
    assert "P-06 - Libellé P-06 injecté depuis la matrice" in texts
    assert "Libellé P-06 injecté depuis la matrice" in _archive_text(
        output, "word/comments.xml"
    )


def test_les_pistes_de_station_deux_sont_ancrees_sur_les_occurrences(tmp_path):
    matrix = load_sami_matrix()
    guided = build_inaccessible(
        PROJECT_ROOT / "_assets" / "graphique-inaccessible.png",
        icon_path=PROJECT_ROOT / "_assets" / "icone-enveloppe.png",
        organigramme_path=PROJECT_ROOT / "_assets" / "organigramme.png",
        texte_image_path=PROJECT_ROOT / "_assets" / "texte-image.png",
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )

    p06 = _comment_anchor(guided, "P-06 -")
    p07 = _comment_anchor(guided, "P-07 -")
    p08 = _comment_anchor(guided, "P-08 -")
    p09 = _comment_anchor(guided, "P-09 -")
    p10 = _comment_anchor(guided, "P-10 -")
    p11 = _comment_anchor(guided, "Document — P-11 -")

    assert p06["drawing"] and not p06["text"]
    assert p07["drawing"] and not p07["text"]
    assert p08["drawing"] and "contactez-nous" in p08["text"]
    assert p09["drawing"] and not p09["text"]
    assert p10["hyperlink"] and "cliquez ici" in p10["text"]
    assert p11["text"].startswith("Ce guide pratique présente")


def test_le_faux_sommaire_ne_reference_plus_les_titres_supprimes(tmp_path):
    matrix = load_sami_matrix()
    guided = build_inaccessible(
        PROJECT_ROOT / "_assets" / "graphique-inaccessible.png",
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    texts = {paragraph.text for paragraph in Document(guided).paragraphs}
    station_title = next(
        block["titre"] for block in matrix["sequence"] if block["id"] == "station-2"
    )

    assert not any(text.startswith("Organisation du service ") for text in texts)
    assert not any(text.startswith("Contact ") for text in texts)
    assert any(text.startswith(f"{station_title} ") for text in texts)


def test_la_station_trois_est_ordonnee_et_pilotee_par_la_matrice(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    paths = [
        build_inaccessible(
            chart_bad,
            output_name="inaccessible.docx",
            matrix=matrix,
            output_dir=tmp_path,
        ),
        build_inaccessible(
            chart_bad,
            with_guidance=True,
            output_name="guide.docx",
            matrix=matrix,
            output_dir=tmp_path,
        ),
        build_accessible(
            chart_good,
            matrix=matrix,
            output_dir=tmp_path,
        ),
    ]
    controls = [item for item in matrix["controles"] if item["station"] == "station-3"]
    station_title = next(
        item["titre"] for item in matrix["sequence"] if item["id"] == "station-3"
    )

    for path in paths:
        texts = [paragraph.text for paragraph in Document(path).paragraphs]
        assert texts.count(station_title) == 1
        heading_indexes = []
        for control in controls:
            heading_text = f"{control['id']} - {control['intitule']}"
            heading_indexes.append(texts.index(heading_text))
            for expected in (
                f"Problème : {control['defaut']}",
                f"Pourquoi : {control['impact']}",
                f"Règle : {control['regle']}",
                f"Dans Word : {control['procedure_word']}",
                f"Dans LibreOffice Writer : {control['procedure_writer']}",
                f"À faire : {control['action_attendue']}",
                f"Preuve : {control['preuve']['attendu']}",
            ):
                assert expected in texts
        assert heading_indexes == sorted(heading_indexes)

    comments = _archive_text(paths[1], "word/comments.xml")
    for control in controls:
        assert comments.count(f"{control['id']} -") == control["occurrences_attendues"]
        assert control["regle"] in comments
    for legacy_label in (
        "Critère 4 -",
        "Critère 5 -",
        "Critère 6 -",
        "Critère 7 -",
        "Critère 21 -",
    ):
        assert legacy_label not in comments

    p12 = _comment_anchor(paths[1], "P-12 -")
    p13 = _comment_anchor(paths[1], "P-13 -")
    p14 = _comment_anchor(paths[1], "P-14 -")
    assert p12["text"].startswith("Information complémentaire")
    assert p13["drawing"]
    assert "Effectif" in p14["text"]


def test_p12_calcule_un_vrai_defaut_de_contraste_et_sa_correction(tmp_path):
    matrix = load_sami_matrix()
    text = "Information complémentaire : résultats provisoires."
    inaccessible = build_inaccessible(
        PROJECT_ROOT / "_assets" / "graphique-inaccessible.png",
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        PROJECT_ROOT / "_assets" / "graphique-inaccessible.png",
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(
        PROJECT_ROOT / "_assets" / "graphique-accessible.png",
        matrix=matrix,
        output_dir=tmp_path,
    )

    bad_run = next(
        paragraph.runs[0]
        for paragraph in Document(inaccessible).paragraphs
        if paragraph.text == text
    )
    good_run = next(
        paragraph.runs[0]
        for paragraph in Document(corrected).paragraphs
        if paragraph.text == text
    )
    bad_color = str(bad_run.font.color.rgb)
    good_color = str(good_run.font.color.rgb)
    assert bad_run.font.size.pt == 11
    assert _contrast_ratio(bad_color) < 4.5
    assert _contrast_ratio(good_color) >= 4.5

    comments = _archive_text(guided, "word/comments.xml")
    document_xml = _archive_text(guided, "word/document.xml")
    assert "#767676" not in comments
    assert "4,48" not in comments
    assert "#767676" not in document_xml
    assert "4,48" not in document_xml
    assert "l'urgence repose surtout sur le rouge" not in comments

    with Image.open(PROJECT_ROOT / "_assets" / "graphique-inaccessible.png") as image:
        pixels = set(image.convert("RGB").get_flattened_data())
    for color in ((208, 0, 0), (24, 117, 60)):
        assert color in pixels
        assert _contrast_ratio("".join(f"{channel:02X}" for channel in color)) >= 3


def test_p13_fournit_dans_word_les_donnees_pour_corriger_le_graphique(tmp_path):
    matrix = load_sami_matrix()
    control = next(item for item in matrix["controles"] if item["id"] == "P-13")
    assert control["ancrage"] == "Image du graphique de la station 3."
    assert control["procedure_word"].startswith("Insertion > Graphique")
    assert "valeurs affichées" in control["procedure_word"]
    assert control["transformations_editoriales"] == ["alternative_graphique_complete"]
    assert "P-13" in matrix["identite_editoriale"]["transformations_autorisees"]
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    inaccessible = build_inaccessible(
        chart_bad,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(
        chart_good,
        matrix=matrix,
        output_dir=tmp_path,
    )
    for path in (inaccessible, guided):
        document = Document(path)
        alt_values = [
            shape._inline.find(qn("wp:docPr")).get("descr")
            for shape in document.inline_shapes
        ]
        assert (
            alt_values.count(
                "Graphique : évolution du trafic web entre T4 2024 et T1 2025."
            )
            == 1
        )
        document_text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
        )
        for value in (
            "18 200",
            "20 400",
            "21 000",
            "23 400",
            "6 000",
            "6 800",
        ):
            assert value not in document_text
            assert value not in "\n".join(alt_values)

    corrected_document = Document(corrected)
    corrected_alt_values = [
        shape._inline.find(qn("wp:docPr")).get("descr")
        for shape in corrected_document.inline_shapes
    ]
    assert (
        corrected_alt_values.count(
            "T4 2024 puis T1 2025 : accès directs, 18 200 puis 20 400 "
            "(+12 %) ; moteurs de recherche, 21 000 puis 23 400 (+11 %) ; "
            "sites référents, 6 000 puis 6 800 (+13 %)."
        )
        == 1
    )

    generator_source = (
        PROJECT_ROOT / "scripts" / "generate_exercice_sami.py"
    ).read_text(encoding="utf-8")
    assert 'ax.set_title("Evolution du trafic web"' not in generator_source
    assert generator_source.count("ax.set_title(CHART_TITLE") == 2
    assert 'label="T4 2024"' in generator_source
    assert 'label="T1 2025"' in generator_source

    assert _embedded_asset_count(inaccessible, chart_bad) == 1
    assert _embedded_asset_count(guided, chart_bad) == 1
    assert _embedded_asset_count(corrected, chart_bad) == 0
    assert _embedded_asset_count(corrected, chart_good) == 1
    assert _comment_anchor(guided, "P-13 -")["drawing"]


def test_p14_corrige_la_structure_du_tableau_de_donnees(tmp_path):
    matrix = load_sami_matrix()
    inaccessible = build_inaccessible(
        PROJECT_ROOT / "_assets" / "graphique-inaccessible.png",
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        PROJECT_ROOT / "_assets" / "graphique-inaccessible.png",
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(
        PROJECT_ROOT / "_assets" / "graphique-accessible.png",
        matrix=matrix,
        output_dir=tmp_path,
    )
    bad_table = _table_containing(Document(inaccessible), "120 000")
    guided_table = _table_containing(Document(guided), "120 000")
    good_document = Document(corrected)
    good_table = _table_containing(good_document, "120 000")

    assert _table_tokens(bad_table) == _table_tokens(good_table)
    assert bad_table._tbl.xpath(".//w:gridSpan")
    assert guided_table._tbl.xpath(".//w:gridSpan")
    assert not good_table._tbl.xpath(".//w:gridSpan")
    assert not good_table._tbl.xpath(".//w:vMerge")
    assert not bad_table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert good_table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert any(not row._tr.xpath("./w:trPr/w:cantSplit") for row in bad_table.rows)
    assert all(row._tr.xpath("./w:trPr/w:cantSplit") for row in good_table.rows)
    assert not any(cell.tables for row in good_table.rows for cell in row.cells)
    assert any(
        paragraph.text == "Répartition par service"
        and paragraph.style.name == "Heading 3"
        for paragraph in good_document.paragraphs
    )
    assert _comment_anchor(guided, "P-14 -")["text"].startswith("Effectif")


def test_p15_distingue_la_langue_principale_et_le_passage_anglais(tmp_path):
    matrix = load_sami_matrix()
    control = next(item for item in matrix["controles"] if item["id"] == "P-15")
    station_title = next(
        block["titre"] for block in matrix["sequence"] if block["id"] == "station-4"
    )
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    inaccessible = build_inaccessible(
        chart_bad,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(chart_good, matrix=matrix, output_dir=tmp_path)
    documents = [Document(path) for path in (inaccessible, guided, corrected)]
    english_text = (
        "The quarterly report is available upon request. "
        "Please contact the communication department for further details."
    )
    heading_text = f"{control['id']} - {control['intitule']}"

    for document in documents:
        texts = [paragraph.text for paragraph in document.paragraphs]
        assert texts.count(station_title) == 1
        assert texts.count(heading_text) == 1

    bad_document, guided_document, good_document = documents
    for document in (bad_document, guided_document):
        assert document.core_properties.language == "de-DE"
        assert _default_document_language(document) == "de-DE"
        for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
            assert _style_language(document, style_name) == "de-DE"

    assert good_document.core_properties.language == "fr-FR"
    assert _default_document_language(good_document) == "fr-FR"
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        assert _style_language(good_document, style_name) == "fr-FR"

    bad_english = next(p for p in bad_document.paragraphs if p.text == english_text)
    guided_english = next(
        p for p in guided_document.paragraphs if p.text == english_text
    )
    good_english = next(p for p in good_document.paragraphs if p.text == english_text)
    assert all(_direct_run_language(run) is None for run in bad_english.runs)
    assert all(_direct_run_language(run) is None for run in guided_english.runs)
    assert {_direct_run_language(run) for run in good_english.runs} == {"en-US"}

    comments = _archive_text(guided, "word/comments.xml")
    assert comments.count("P-15 -") == 1
    assert control["regle"] in comments
    assert _comment_anchor(guided, "P-15 -")["text"] == english_text


def test_p16_corrige_la_lisibilite_par_le_style_normal(tmp_path):
    matrix = load_sami_matrix()
    control = next(item for item in matrix["controles"] if item["id"] == "P-16")
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    inaccessible = build_inaccessible(
        chart_bad,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(chart_good, matrix=matrix, output_dir=tmp_path)
    bad_document = Document(inaccessible)
    guided_document = Document(guided)
    good_document = Document(corrected)
    sample_text = (
        "La mise en forme du corps du document est pilotée par le style Normal."
    )

    for document in (bad_document, guided_document, good_document):
        heading_text = f"{control['id']} - {control['intitule']}"
        assert sum(p.text == heading_text for p in document.paragraphs) == 1
        sample = next(p for p in document.paragraphs if p.text == sample_text)
        assert sample.style.name == "Normal"
        assert all(run.font.size is None for run in sample.runs)

    for document in (bad_document, guided_document):
        normal = document.styles["Normal"]
        assert normal.font.name == "Arial"
        assert normal.font.size.pt == 11
        assert normal.paragraph_format.line_spacing == 1.0
        assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    normal = good_document.styles["Normal"]
    assert normal.font.name == "Arial"
    assert normal.font.size.pt == 12
    assert normal.paragraph_format.line_spacing >= 1.15
    assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
    explicit_sizes = [
        run.font.size.pt
        for run in _useful_document_runs(good_document)
        if run.font.size is not None and run.text.strip()
    ]
    assert explicit_sizes
    assert min(explicit_sizes) >= 12

    comments = _archive_text(guided, "word/comments.xml")
    assert comments.count("P-16 -") == 1
    assert control["procedure_word"] in _comment_text(guided, "P-16 -")
    assert _comment_anchor(guided, "P-16 -")["text"] == sample_text


def test_p17_conserve_les_accents_et_applique_la_casse_par_la_forme(tmp_path):
    matrix = load_sami_matrix()
    control = next(item for item in matrix["controles"] if item["id"] == "P-17")
    assert control["transformations_editoriales"] == [
        "retablissement_casse_accents_source"
    ]
    assert "P-17" in matrix["identite_editoriale"]["transformations_autorisees"]
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    inaccessible = build_inaccessible(
        chart_bad,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(chart_good, matrix=matrix, output_dir=tmp_path)
    bad_document = Document(inaccessible)
    guided_document = Document(guided)
    good_document = Document(corrected)
    heading_text = f"{control['id']} - {control['intitule']}"

    for document in (bad_document, guided_document, good_document):
        assert sum(p.text == heading_text for p in document.paragraphs) == 1

    for document in (bad_document, guided_document):
        target = next(
            p for p in document.paragraphs if p.text == "ANNEXES - ACCESSIBILITE"
        )
        assert target.style.name == "Heading 3"
        assert "w:caps" not in target._p.xml

    corrected_target = next(
        p for p in good_document.paragraphs if p.text == "Annexes - accessibilité"
    )
    assert corrected_target.style.name == "Heading 3"
    assert corrected_target.runs
    assert all(run.font.all_caps for run in corrected_target.runs)
    assert "ACCESSIBILITE" not in corrected_target._p.xml
    assert "accessibilité" in corrected_target._p.xml

    comments = _archive_text(guided, "word/comments.xml")
    assert comments.count("P-17 -") == 1
    assert "Critère 17 -" not in comments
    assert control["regle"] in _comment_text(guided, "P-17 -")
    assert _comment_anchor(guided, "P-17 -")["text"] == ("ANNEXES - ACCESSIBILITE")


def test_p18_developpe_le_rgaa_et_documente_le_reglage_du_correcteur(tmp_path):
    matrix = load_sami_matrix()
    control = next(item for item in matrix["controles"] if item["id"] == "P-18")
    assert control["transformations_editoriales"] == [
        "developpement_acronyme_premiere_occurrence"
    ]
    assert "P-18" in matrix["identite_editoriale"]["transformations_autorisees"]
    assert "Ignorer les mots en MAJUSCULES" in control["procedure_word"]
    assert "Paramètres linguistiques" in control["procedure_writer"]

    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    inaccessible = build_inaccessible(
        chart_bad,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(chart_good, matrix=matrix, output_dir=tmp_path)
    bad_document = Document(inaccessible)
    guided_document = Document(guided)
    good_document = Document(corrected)
    bad_text = "Le RGAA structure le contrôle des documents numériques."
    good_text = (
        "Le Référentiel général d’amélioration de l’accessibilité (RGAA) "
        "structure le contrôle des documents numériques."
    )
    heading_text = f"{control['id']} - {control['intitule']}"

    for document in (bad_document, guided_document, good_document):
        assert sum(p.text == heading_text for p in document.paragraphs) == 1
        body_text = {paragraph.text for paragraph in document.paragraphs}
        assert f"Dans Word : {control['procedure_word']}" in body_text
        assert f"Dans LibreOffice Writer : {control['procedure_writer']}" in body_text

    for document in (bad_document, guided_document):
        assert sum(p.text == bad_text for p in document.paragraphs) == 1
        assert good_text not in {p.text for p in document.paragraphs}

    assert sum(p.text == good_text for p in good_document.paragraphs) == 1
    assert bad_text not in {p.text for p in good_document.paragraphs}
    first_rgaa = next(p.text for p in good_document.paragraphs if "RGAA" in p.text)
    assert first_rgaa == good_text

    comments = _archive_text(guided, "word/comments.xml")
    assert comments.count("P-18 -") == 1
    assert control["procedure_word"] in _comment_text(guided, "P-18 -")
    assert _comment_anchor(guided, "P-18 -")["text"] == bad_text


def test_la_station_quatre_est_ordonnee_et_pilotee_par_la_matrice(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    paths = [
        build_inaccessible(
            chart_bad,
            output_name="inaccessible.docx",
            matrix=matrix,
            output_dir=tmp_path,
        ),
        build_inaccessible(
            chart_bad,
            with_guidance=True,
            output_name="guide.docx",
            matrix=matrix,
            output_dir=tmp_path,
        ),
        build_accessible(chart_good, matrix=matrix, output_dir=tmp_path),
    ]
    controls = [item for item in matrix["controles"] if item["station"] == "station-4"]
    station_title = next(
        item["titre"] for item in matrix["sequence"] if item["id"] == "station-4"
    )

    for path in paths:
        texts = [paragraph.text for paragraph in Document(path).paragraphs]
        assert texts.count(station_title) == 1
        heading_indexes = []
        for control in controls:
            heading_text = f"{control['id']} - {control['intitule']}"
            heading_indexes.append(texts.index(heading_text))
            for expected in (
                f"Problème : {control['defaut']}",
                f"Pourquoi : {control['impact']}",
                f"Règle : {control['regle']}",
                f"Dans Word : {control['procedure_word']}",
                f"Dans LibreOffice Writer : {control['procedure_writer']}",
                f"À faire : {control['action_attendue']}",
                f"Preuve : {control['preuve']['attendu']}",
            ):
                assert expected in texts
        assert heading_indexes == sorted(heading_indexes)

    comments = _archive_text(paths[1], "word/comments.xml")
    for control in controls:
        assert comments.count(f"{control['id']} -") == 1
        assert control["regle"] in _comment_text(paths[1], f"{control['id']} -")
    for legacy_label in ("Critère 13 -", "Critère 15 -", "Critère 17 -"):
        assert legacy_label not in comments


def test_p19_renseigne_les_proprietes_et_preserve_les_noms_des_versions(tmp_path):
    matrix = load_sami_matrix()
    control = next(item for item in matrix["controles"] if item["id"] == "P-19")
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    paths = [
        build_inaccessible(chart_bad, matrix=matrix, output_dir=tmp_path),
        build_inaccessible(
            chart_bad,
            with_guidance=True,
            matrix=matrix,
            output_dir=tmp_path,
        ),
        build_accessible(chart_good, matrix=matrix, output_dir=tmp_path),
    ]
    documents = [Document(path) for path in paths]
    anchor_text = (
        "Document — propriétés : vérifiez le titre, l’auteur, la langue et le "
        "nom de votre copie de travail."
    )

    assert [path.name for path in paths] == [
        "tp-doc-inaccessible.docx",
        "tp-doc-aide-correction.docx",
        "tp-doc-accessible.docx",
    ]
    for document in documents:
        texts = [paragraph.text for paragraph in document.paragraphs]
        assert f"{control['id']} - {control['intitule']}" in texts
        assert texts.count(anchor_text) == 1

    for document in documents[:2]:
        assert document.core_properties.title == ""
        assert document.core_properties.author == ""
        assert document.core_properties.language == "de-DE"

    corrected = documents[2]
    assert corrected.core_properties.title == "Rendre un document Word accessible"
    assert corrected.core_properties.author == "Sami Dupont"
    assert corrected.core_properties.language == "fr-FR"

    comments = _archive_text(paths[1], "word/comments.xml")
    assert comments.count("P-19 -") == 1
    assert "Critère 14 -" not in comments
    assert control["regle"] in _comment_text(paths[1], "P-19 -")
    assert _comment_anchor(paths[1], "P-19 -")["text"] == anchor_text


def test_la_station_cinq_est_complete_ordonnee_et_identique_dans_les_trois_docx(
    tmp_path,
):
    matrix = load_sami_matrix()
    controls = [item for item in matrix["controles"] if item["station"] == "station-5"]
    assert [control["id"] for control in controls] == [
        "P-19",
        "P-20",
        "C-01",
        "C-02",
    ]
    station_title = next(
        item["titre"] for item in matrix["sequence"] if item["id"] == "station-5"
    )
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    chart_good = PROJECT_ROOT / "_assets" / "graphique-accessible.png"
    paths = [
        build_inaccessible(
            chart_bad,
            output_name="inaccessible.docx",
            matrix=matrix,
            output_dir=tmp_path,
        ),
        build_inaccessible(
            chart_bad,
            with_guidance=True,
            output_name="guide.docx",
            matrix=matrix,
            output_dir=tmp_path,
        ),
        build_accessible(chart_good, matrix=matrix, output_dir=tmp_path),
    ]
    station_bodies = []
    no_injected_defect = (
        "Cette étape de finalisation ne peut pas être prouvée par le DOCX seul."
    )

    for path in paths:
        texts = [paragraph.text for paragraph in Document(path).paragraphs]
        assert texts.count(station_title) == 1
        heading_indexes = []
        for control in controls:
            heading_text = f"{control['id']} - {control['intitule']}"
            heading_indexes.append(texts.index(heading_text))
            problem = control["defaut"] or no_injected_defect
            for expected in (
                f"Problème : {problem}",
                f"Pourquoi : {control['impact']}",
                f"Règle : {control['regle']}",
                f"Dans Word : {control['procedure_word']}",
                f"Dans LibreOffice Writer : {control['procedure_writer']}",
                f"À faire : {control['action_attendue']}",
                f"Preuve : {control['preuve']['attendu']}",
            ):
                assert expected in texts
        assert heading_indexes == sorted(heading_indexes)
        station_bodies.append(texts[texts.index(station_title) :])

    assert station_bodies[0] == station_bodies[1] == station_bodies[2]

    comments = _archive_text(paths[1], "word/comments.xml")
    assert comments.count("P-19 -") == 1
    for control_id in ("P-20", "C-01", "C-02"):
        assert comments.count(f"{control_id} -") == 0

    c01 = next(control for control in controls if control["id"] == "C-01")
    assert "zéro erreur" in c01["regle"]
    assert "ne remplace pas" in c01["impact"]
    p20 = next(control for control in controls if control["id"] == "P-20")
    assert all(
        term in p20["procedure_word"] for term in ("propriétés", "balises", "signets")
    )
    c02 = next(control for control in controls if control["id"] == "C-02")
    assert "PAC" in c02["regle"]
    assert "Acrobat Pro" in c02["regle"]
    assert "checklist humaine" in c02["regle"]


def test_le_guide_couvre_exactement_les_occurrences_et_preserve_le_corps(tmp_path):
    matrix = load_sami_matrix()
    chart_bad = PROJECT_ROOT / "_assets" / "graphique-inaccessible.png"
    icon = PROJECT_ROOT / "_assets" / "icone-enveloppe.png"
    organigramme = PROJECT_ROOT / "_assets" / "organigramme.png"
    texte_image = PROJECT_ROOT / "_assets" / "texte-image.png"
    inaccessible = build_inaccessible(
        chart_bad,
        icon_path=icon,
        organigramme_path=organigramme,
        texte_image_path=texte_image,
        output_name="inaccessible.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    guided = build_inaccessible(
        chart_bad,
        icon_path=icon,
        organigramme_path=organigramme,
        texte_image_path=texte_image,
        with_guidance=True,
        output_name="guide.docx",
        matrix=matrix,
        output_dir=tmp_path,
    )
    corrected = build_accessible(
        PROJECT_ROOT / "_assets" / "graphique-accessible.png",
        icon_path=icon,
        organigramme_path=organigramme,
        texte_image_path=texte_image,
        matrix=matrix,
        output_dir=tmp_path,
    )

    assert _body_signature(Document(inaccessible)) == _body_signature(Document(guided))
    assert "word/comments.xml" not in _archive_members(inaccessible)
    assert "word/comments.xml" not in _archive_members(corrected)

    comments_xml = _archive_text(guided, "word/comments.xml")
    comments_root = etree.fromstring(comments_xml.encode("utf-8"))
    expected_count = sum(
        control["occurrences_attendues"] for control in matrix["controles"]
    )
    assert len(comments_root.xpath("//w:comment", namespaces={"w": WORD_NS})) == (
        expected_count
    )
    for control in matrix["controles"]:
        assert (
            comments_xml.count(f"{control['id']} -") == control["occurrences_attendues"]
        )
        if control["occurrences_attendues"]:
            comment = _comment_text(guided, f"{control['id']} -")
            for expected in (
                f"Problème : {control['defaut']}",
                f"Impact : {control['impact']}",
                f"Règle : {control['regle']}",
                f"Première action : {control['action_attendue']}",
            ):
                assert expected in comment


def test_le_controle_negatif_detecte_un_defaut_reinjecte_par_station(tmp_path):
    corrected = build_accessible(
        PROJECT_ROOT / "_assets" / "graphique-accessible.png",
        matrix=load_sami_matrix(),
        output_dir=tmp_path,
    )
    _assert_detectable_station_families_are_corrected(corrected)

    def reintroduce_structure(document):
        next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text == "Introduction"
        ).style = document.styles["Normal"]

    def reintroduce_link(document):
        document.element.xpath("//w:hyperlink//w:t")[0].text = "cliquez ici"

    def reintroduce_contrast(document):
        run = next(
            paragraph.runs[0]
            for paragraph in document.paragraphs
            if paragraph.text == "Information complémentaire : résultats provisoires."
        )
        run.font.color.rgb = RGBColor(0xB5, 0xB5, 0xB5)

    def reintroduce_typography(document):
        document.styles["Normal"].font.size = Pt(11)

    def reintroduce_metadata(document):
        document.core_properties.title = ""

    for family, mutate in (
        ("station-1", reintroduce_structure),
        ("station-2", reintroduce_link),
        ("station-3", reintroduce_contrast),
        ("station-4", reintroduce_typography),
        ("station-5", reintroduce_metadata),
    ):
        document = Document(corrected)
        mutate(document)
        candidate = tmp_path / f"defaut-{family}.docx"
        document.save(candidate)
        with pytest.raises(AssertionError, match=family):
            _assert_detectable_station_families_are_corrected(candidate)


def test_le_corrige_decoupe_le_guide_en_unites_imprimables_sans_titre_orphelin(
    tmp_path,
):
    matrix = load_sami_matrix()
    corrected = build_accessible(
        PROJECT_ROOT / "_assets" / "graphique-accessible.png",
        matrix=matrix,
        output_dir=tmp_path,
    )
    document = Document(corrected)
    paragraphs = {paragraph.text: paragraph for paragraph in document.paragraphs}

    for station in (
        block for block in matrix["sequence"] if block["id"].startswith("station-")
    ):
        station_heading = paragraphs[station["titre"]]
        assert station_heading.paragraph_format.page_break_before is True
        assert station_heading.paragraph_format.keep_with_next is True

        controls = [
            control
            for control in matrix["controles"]
            if control["station"] == station["id"]
        ]
        for index, control in enumerate(controls):
            heading = paragraphs[f"{control['id']} - {control['intitule']}"]
            assert heading.paragraph_format.keep_with_next is True
            if index:
                assert heading.paragraph_format.page_break_before is True
            else:
                assert heading.paragraph_format.page_break_before is not True


def test_commentaires_de_correction_en_francais_sans_corriger_le_defaut_de_langue(
    tmp_path,
):
    guided = build_inaccessible(
        PROJECT_ROOT / "_assets" / "graphique-inaccessible.png",
        with_guidance=True,
        output_name="guide.docx",
        matrix=load_sami_matrix(),
        output_dir=tmp_path,
    )
    namespaces = {"w": WORD_NS}
    comments_root = etree.fromstring(
        _archive_text(guided, "word/comments.xml").encode("utf-8")
    )
    runs = comments_root.xpath("//w:comment//w:r[w:t]", namespaces=namespaces)

    assert runs
    for run in runs:
        assert run.xpath("w:rPr/w:lang/@w:val", namespaces=namespaces) == ["fr-FR"]
    # Le défaut pédagogique de la station 4 reste dans le document lui-même.
    styles = _archive_text(guided, "word/styles.xml")
    assert re.search(r'<w:docDefaults>.*w:val="de-DE"', styles, re.S)
