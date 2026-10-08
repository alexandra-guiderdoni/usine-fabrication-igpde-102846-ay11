"""Tests d'intégration des checklists générées depuis la matrice Sami."""

from __future__ import annotations

import re
import subprocess
from pathlib import Path
from zipfile import ZipFile

import fabriquer_pack
from docx import Document
from exercice_sami_matrice import load_sami_matrix
from generate_exercice_sami import build_checklists

from config import load_formation_config

PROJECT_ROOT = Path(__file__).parent.parent
PACK = PROJECT_ROOT / load_formation_config()["livrables"]
CHECKLIST_DOCX = "checklist-accessibilite-bureautique.docx"
CHECKLIST_PDF = "checklist-accessibilite-bureautique.pdf"
CHECKLIST_MARKDOWN = "_source/checklist-accessibilite-bureautique.md"


def test_les_deux_checklists_reprennent_la_matrice_dans_l_ordre(tmp_path):
    matrix = load_sami_matrix()
    markdown = tmp_path / "checklist.md"
    docx = tmp_path / "checklist.docx"

    build_checklists(matrix=matrix, markdown_output=markdown, docx_output=docx)

    assert markdown.is_file()
    assert docx.is_file()

    markdown_text = markdown.read_text(encoding="utf-8")
    markdown_positions = []
    expected_rows = []
    for control in matrix["controles"]:
        expected = (
            f"**{control['id']} · {control['niveau']} - {control['checklist']}**"
        )
        markdown_positions.append(markdown_text.index(expected))
        expected_rows.append(
            (
                f"{control['id']} · {control['niveau']}",
                control["checklist"],
            )
        )
    assert markdown_positions == sorted(markdown_positions)

    document = Document(docx)
    actual_rows = [
        (row.cells[0].text, row.cells[1].text)
        for table in document.tables
        for row in table.rows[1:]
    ]
    assert actual_rows == expected_rows


def test_la_checklist_docx_met_en_gras_les_intitules_des_points(tmp_path):
    markdown = tmp_path / "checklist.md"
    docx = tmp_path / CHECKLIST_DOCX
    build_checklists(
        matrix=load_sami_matrix(), markdown_output=markdown, docx_output=docx
    )

    for table in Document(docx).tables:
        for row in table.rows[1:]:
            for cell in row.cells[:2]:
                runs = [run for paragraph in cell.paragraphs for run in paragraph.runs]
                assert runs and all(run.bold is True for run in runs)
            suivi_runs = [
                run for paragraph in row.cells[2].paragraphs for run in paragraph.runs
            ]
            assert suivi_runs and not any(run.bold for run in suivi_runs)


def test_make_expose_la_cible_checklist():
    result = subprocess.run(
        ["make", "-n", "checklist"],
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert "generate_exercice_sami.py --checklist" in result.stdout


def test_make_pack_regenere_les_pdf_avant_le_controle_de_fraicheur():
    result = subprocess.run(
        ["make", "-n", "pack"],
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert result.stdout.index("fabriquer_pack.py pdf") < result.stdout.index(
        "verifier_fraicheur_pack.py"
    )

    makefile = (PROJECT_ROOT / "Makefile").read_text(encoding="utf-8")
    recette_pack = makefile.split("\npack:\n", 1)[1].split("\n\napercu:", 1)[0]
    assert recette_pack.index("$(MAKE) pdf") < recette_pack.index(
        "$(MAKE) fraicheur-pack"
    )
    assert recette_pack.index("$(MAKE) fraicheur-pack") < recette_pack.index(
        "$(MAKE) deck"
    )


def test_make_pdf_reutilise_la_source_markdown_de_la_checklist():
    checklist = next(
        entry
        for entry in fabriquer_pack.PDFS
        if entry[0] == "_source/checklist-accessibilite-bureautique.md"
    )

    assert checklist[1].endswith(
        "/Livrables-Stagiaires/tp-word-igpde/checklist-accessibilite-bureautique.pdf"
    )
    assert checklist[3][-2:] == [
        "--subtitle",
        "Suivi progressif du TP Word accessible",
    ]
    assert checklist[4] is None


def test_checklists_sont_suivables_sans_formulaire_interactif(tmp_path):
    matrix = load_sami_matrix()
    markdown = tmp_path / "checklist.md"
    docx = tmp_path / "checklist.docx"

    build_checklists(matrix=matrix, markdown_output=markdown, docx_output=docx)

    markdown_text = markdown.read_text(encoding="utf-8")
    assert "Généré depuis `_source/exercice-sami-matrice.yml`" in markdown_text
    assert markdown_text.count("Suivi : ☐ À vérifier · ☐ Fait · ☐ À reprendre") == len(
        matrix["controles"]
    )
    for control in matrix["controles"]:
        if control["niveau"] == "S":
            assert control["checklist"] in markdown_text
    assert "sans manipulation obligatoire pendant le TP" in markdown_text

    document = Document(docx)
    assert document.core_properties.title == (
        "Checklist accessibilité des documents bureautiques"
    )
    assert document.core_properties.language == "fr-FR"
    checklist_rows = [row for table in document.tables for row in table.rows[1:]]
    assert len(checklist_rows) == len(matrix["controles"])
    assert all(
        row.cells[2].text == "☐ À vérifier\n☐ Fait\n☐ À reprendre\nNotes :"
        for row in checklist_rows
    )
    with ZipFile(docx) as archive:
        word_xml = b"".join(
            archive.read(member)
            for member in archive.namelist()
            if member.startswith("word/") and member.endswith(".xml")
        )
    assert b"<w:sdt" not in word_xml
    assert b"<w:ffData" not in word_xml
    assert b"<w:checkBox" not in word_xml
    assert word_xml.count(b"<w:tblHeader") == len(document.tables)


def test_checklists_ecartent_les_formulations_interdites(tmp_path):
    markdown = tmp_path / "checklist.md"
    docx = tmp_path / "checklist.docx"
    build_checklists(markdown_output=markdown, docx_output=docx)

    document = Document(docx)
    docx_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    outputs = markdown.read_text(encoding="utf-8") + "\n" + docx_text
    for forbidden in (
        "H1/H2/H3",
        "CSS",
        "sans erreur résiduelle",
        "Aucun Tableaux de mise en page avec habillage Aucun",
    ):
        assert forbidden.casefold() not in outputs.casefold()


def test_markdown_a_un_sous_titre_explicite_sans_case_redondante(tmp_path):
    markdown = tmp_path / "checklist.md"
    docx = tmp_path / "checklist.docx"
    build_checklists(markdown_output=markdown, docx_output=docx)

    text = markdown.read_text(encoding="utf-8")
    assert 'subtitle: "Suivi progressif du TP Word accessible"' in text
    assert "break-inside: avoid" in text
    assert "- ☐ **" not in text


def test_markdown_utilise_le_code_de_formation_configure(tmp_path):
    markdown = tmp_path / "checklist.md"
    docx = tmp_path / "checklist.docx"
    build_checklists(
        markdown_output=markdown,
        docx_output=docx,
        formation_code="999999",
    )

    assert 'author: "IGPDE - Formation 999999"' in markdown.read_text(encoding="utf-8")


def test_checklists_sont_declarees_sans_liste_normative_parallele():
    agents = (PROJECT_ROOT / "AGENTS.md").read_text(encoding="utf-8")
    pack_readme = (PACK / "README.md").read_text(encoding="utf-8")
    legacy = (
        PACK / "Livrables-Formateur" / "_alex" / "checklist-bureautique.md"
    ).read_text(encoding="utf-8")

    for filename in (CHECKLIST_DOCX, CHECKLIST_PDF, CHECKLIST_MARKDOWN):
        assert filename in agents
        assert filename in pack_readme
    assert "make checklist" in agents
    assert "- [ ]" not in legacy
    assert CHECKLIST_DOCX in legacy
    assert CHECKLIST_PDF in legacy
    assert "_source/exercice-sami-matrice.yml" in legacy


def test_la_checklist_docx_est_paginee_page_x_sur_y(tmp_path):
    markdown = tmp_path / "checklist.md"
    docx = tmp_path / CHECKLIST_DOCX
    build_checklists(
        matrix=load_sami_matrix(), markdown_output=markdown, docx_output=docx
    )
    footer_xml = Document(docx).sections[0].footer._element.xml

    instructions = re.findall(r"<w:instrText[^>]*>([^<]*)</w:instrText>", footer_xml)
    assert [instruction.strip() for instruction in instructions] == ["PAGE", "NUMPAGES"]
    texts = re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", footer_xml)
    assert "".join(texts) == "Page 1 sur 1"
    assert set(re.findall(r'<w:sz w:val="(\d+)"/>', footer_xml)) == {"24"}


def test_la_checklist_docx_met_en_gras_les_niveaux_p_c_s(tmp_path):
    markdown = tmp_path / "checklist.md"
    docx = tmp_path / CHECKLIST_DOCX
    build_checklists(
        matrix=load_sami_matrix(), markdown_output=markdown, docx_output=docx
    )
    definitions = {
        paragraph.text.split(" : ", 1)[0]: paragraph
        for paragraph in Document(docx).paragraphs
        if paragraph.style.name == "List Bullet"
    }

    for label in ("P - pratiqué", "C - contrôlé", "S - signalé"):
        first, *rest = definitions[label].runs
        assert first.text == f"{label} :"
        assert first.bold is True
        assert rest and not any(run.bold for run in rest)


def test_les_tableaux_de_la_checklist_suivent_les_regles_d_accessibilite(tmp_path):
    markdown = tmp_path / "checklist.md"
    docx = tmp_path / CHECKLIST_DOCX
    matrix = load_sami_matrix()
    build_checklists(matrix=matrix, markdown_output=markdown, docx_output=docx)
    document = Document(docx)
    titles = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "Heading 1"
    ]

    assert len(document.tables) == len(titles)
    for table, title in zip(document.tables, titles):
        table_xml = table._tbl.xml
        header = table.rows[0]
        assert [cell.text for cell in header.cells] == [
            "Identifiant et niveau",
            "Point à vérifier",
            "Suivi",
        ]
        assert all(
            run.bold is True
            for cell in header.cells
            for paragraph in cell.paragraphs
            for run in paragraph.runs
        )
        # Titre et description du tableau (Texte de remplacement dans Word).
        caption = re.findall(r'<w:tblCaption w:val="([^"]*)"/>', table_xml)
        description = re.findall(r'<w:tblDescription w:val="([^"]*)"/>', table_xml)
        assert caption == [title]
        assert len(description) == 1 and description[0]
        # Règle P-14 : en-tête répété, lignes insécables, ni fusion ni flottant.
        assert "w:tblHeader" in table_xml
        assert table_xml.count("<w:cantSplit") == len(table.rows)
        assert "w:gridSpan" not in table_xml and "w:vMerge" not in table_xml
        assert "w:tblpPr" not in table_xml
        assert not table._tbl.xpath(".//w:tbl")


def test_la_checklist_docx_a_des_titres_arial_et_une_colonne_de_points_large(
    tmp_path,
):
    markdown = tmp_path / "checklist.md"
    docx = tmp_path / CHECKLIST_DOCX
    build_checklists(
        matrix=load_sami_matrix(), markdown_output=markdown, docx_output=docx
    )
    document = Document(docx)

    for style_name in ("Title", "Heading 1"):
        fonts = re.findall(r"<w:rFonts [^>]*>", document.styles[style_name].element.xml)
        assert len(fonts) == 1
        assert 'w:ascii="Arial"' in fonts[0] and 'w:hAnsi="Arial"' in fonts[0]
        assert "Theme" not in fonts[0] and "theme" not in fonts[0]

    # 1,3 + 3,2 + 1,5 pouces : la largeur utile de 6 pouces, en twips.
    expected = ["1872", "4608", "2160"]
    for table in document.tables:
        table_xml = table._tbl.xml
        assert re.findall(r'<w:gridCol w:w="(\d+)"/>', table_xml) == expected
        for row in table.rows:
            widths = [
                re.findall(r'<w:tcW [^>]*w:w="(\d+)"', cell._tc.xml)[0]
                for cell in row.cells
            ]
            assert widths == expected


def test_la_checklist_pdf_reprend_la_mise_en_forme_de_la_docx(tmp_path):
    markdown = tmp_path / "checklist.md"
    docx = tmp_path / CHECKLIST_DOCX
    build_checklists(
        matrix=load_sami_matrix(), markdown_output=markdown, docx_output=docx
    )
    text = markdown.read_text(encoding="utf-8")
    style = re.search(r"<style>(.*?)</style>", text, re.DOTALL).group(1)

    for label in ("P - pratiqué", "C - contrôlé", "S - signalé"):
        assert f"- **{label} :** " in text
    assert "break-inside: avoid" in style
    assert "font-family: Arial" in style
    # Sans césure : pas de mot coupé, comme dans la checklist DOCX.
    assert "hyphens: manual" in style
    assert 'content: "Page " counter(page) " sur " counter(pages)' in style
