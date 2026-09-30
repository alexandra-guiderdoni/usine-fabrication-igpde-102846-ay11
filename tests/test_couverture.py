"""Tests du contenu propre a la slide 1 de couverture."""

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation as PrsLoad
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from igpde_dsfr_components import NSMAP_A, create_presentation, finalize_pptx
from slides import SlideContext, load_slide_module


SLIDE_PATH = Path(__file__).parent.parent / "scripts" / "slides" / "01_couverture.py"
TITRE_COMPLET = "L'accessibilité numérique pour la bureautique et le web"
NOTE_ORALE = "Slide de couverture - accueil des stagiaires et installation."


def test_couverture_titre_semantique_complet_et_note_orale_a_jour(tmp_path):
    prs, layouts = create_presentation()
    ctx = SlideContext(
        page_num=1,
        date="9 octobre 2026",
        footer_base="Formation 102846",
        formation_code="102846",
    )

    slide = load_slide_module(SLIDE_PATH).build(prs, layouts, ctx)

    assert slide.shapes.title is not None
    assert slide.shapes.title.text == TITRE_COMPLET
    assert slide.notes_slide.notes_text_frame.text.strip() == NOTE_ORALE

    sous_titre = next(
        shape for shape in slide.shapes if shape.name == "DSFR-couverture-soustitre"
    )
    assert (
        sous_titre.top
        >= slide.shapes.title.top + slide.shapes.title.height + Inches(0.12)
    )
    assert sous_titre.top + sous_titre.height <= Inches(6.80)

    output = tmp_path / "couverture.pptx"
    finalize_pptx(prs, str(output))
    slide = PrsLoad(str(output)).slides[0]

    titres = [
        ph
        for ph in slide.placeholders
        if ph.placeholder_format.type
        in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    ]
    assert len(titres) == 1
    assert titres[0].text == TITRE_COMPLET

    textes = []
    for child in slide.shapes._spTree:
        if etree.QName(child.tag).localname not in (
            "sp",
            "pic",
            "graphicFrame",
            "cxnSp",
        ):
            continue
        texte = "".join(
            node.text or "" for node in child.findall(f".//{{{NSMAP_A}}}t")
        ).strip()
        if texte:
            textes.append(texte)
    assert textes[0] == TITRE_COMPLET
