"""Tests des slides 2 et 3 relues avec la formatrice."""

import sys
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from igpde_dsfr_components import create_presentation
from slides import SlideContext, load_slide_module


SLIDES_DIR = Path(__file__).parent.parent / "scripts" / "slides"


def _build_slide(filename, page_num):
    prs, layouts = create_presentation()
    ctx = SlideContext(
        page_num=page_num,
        date="9 octobre 2026",
        footer_base="Formation 102846",
        formation_code="102846",
    )
    return load_slide_module(SLIDES_DIR / filename).build(prs, layouts, ctx)


def test_slide_2_presente_les_objectifs_sans_chevauchement():
    slide = _build_slide("02_objectifs.py", 2)

    assert slide.shapes.title is not None
    assert slide.shapes.title.text == "Objectifs pédagogiques"
    assert not any(shape.name == "DSFR-callout-titre" for shape in slide.shapes)

    carte = next(shape for shape in slide.shapes if shape.name == "DSFR-box")
    qrcode = next(shape for shape in slide.shapes if shape.name == "DSFR-qrcode")
    assert qrcode.top - (carte.top + carte.height) >= Inches(0.35)

    url = next(
        shape for shape in slide.shapes if shape.name == "DSFR-qrcode-url-visible"
    )
    taille_url = url.text_frame.paragraphs[0].runs[0].font.size
    assert taille_url is not None
    assert taille_url.pt >= 10

    affiche = next(
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name != "DSFR-qrcode"
    )
    c_nv_pr = affiche._element.find(
        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"
    )
    alternative = c_nv_pr.get("descr", "")
    assert "agents publics" in alternative.lower()
    assert "accessibilite.gouv.fr" in alternative.lower()


def test_slide_3_hierarchise_et_espace_les_quatre_modules():
    slide = _build_slide("02a_sommaire.py", 3)

    titres = [shape.text for shape in slide.shapes if shape.name == "DSFR-card-titre"]
    assert titres == [
        "Accessibilité et cadre légal",
        "Bureautique accessible",
        "Points de contrôle rapides W3C",
        "Réseaux sociaux",
    ]

    contenus = "\n".join(
        shape.text for shape in slide.shapes if shape.name == "DSFR-card-contenu"
    )
    assert "Documents Word (LibreOffice)" in contenus

    numeros = [shape for shape in slide.shapes if shape.name == "DSFR-card-numero"]
    assert [shape.text for shape in numeros] == ["1", "2", "3", "4"]
    assert all(shape.width >= Inches(0.60) for shape in numeros)
    assert all(
        shape.text_frame.paragraphs[0].runs[0].font.size.pt >= 20 for shape in numeros
    )

    cartes = [shape for shape in slide.shapes if shape.name == "DSFR-box"]
    rangees = sorted({shape.top for shape in cartes})
    assert len(rangees) == 2
    assert rangees[0] >= Inches(2.25)
    assert rangees[1] - (rangees[0] + cartes[0].height) >= Inches(0.30)
    assert max(shape.top + shape.height for shape in cartes) <= Inches(6.75)
