"""Tests des composants visuels (necessitent le template PPTX)."""

import sys
from pathlib import Path

import pytest
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from igpde_dsfr_components import (
    add_callout,
    add_alert,
    add_highlight,
    add_quote,
    add_card,
    add_pave_chiffre,
    add_stepper,
    add_tableau,
    add_texte_libre,
    add_notes,
    add_encadre,
    add_image,
    MARGIN_L,
    CONTENT_W,
    COL_W,
    TOP_CONTENT,
)


class TestAddCallout:
    def test_cree_des_shapes(self, slide):
        n_before = len(slide.shapes)
        add_callout(slide, "Titre", ["Bullet 1", "Bullet 2"], top=TOP_CONTENT)
        assert len(slide.shapes) > n_before

    def test_accepte_liste_vide(self, slide):
        add_callout(slide, "Titre", [], top=TOP_CONTENT)

    def test_hauteur_explicite(self, slide):
        add_callout(slide, "Titre", ["B1"], top=TOP_CONTENT, height=3.0)

    def test_mode_compact_reduit_la_hauteur(self, slide):
        add_callout(
            slide,
            "Titre",
            ["Une phrase suffisamment longue pour tester le calcul de hauteur"] * 4,
            top=TOP_CONTENT,
        )
        hauteur_normale = [
            shape.height for shape in slide.shapes if shape.name == "DSFR-box"
        ][-1]
        add_callout(
            slide,
            "Titre",
            ["Une phrase suffisamment longue pour tester le calcul de hauteur"] * 4,
            top=TOP_CONTENT,
            compact=True,
        )
        hauteur_compacte = [
            shape.height for shape in slide.shapes if shape.name == "DSFR-box"
        ][-1]
        assert hauteur_compacte < hauteur_normale

    def test_string_en_bullets_itere_caracteres(self, slide):
        # Piege documente dans CLAUDE.md : string au lieu de liste
        # Ne leve pas d'erreur mais produit un bullet par caractere (degrade)
        add_callout(slide, "Titre", "abc", top=TOP_CONTENT)
        # Le test verifie que ca ne crashe pas — le resultat est incorrect
        # mais le code ne valide pas le type de bullets


class TestAddAlert:
    def test_types_valides(self, slide):
        for t in ("warning", "error", "success", "info"):
            n_before = len(slide.shapes)
            add_alert(slide, f"Titre {t}", ["Item"], top=TOP_CONTENT, alert_type=t)
            assert len(slide.shapes) > n_before

    def test_type_invalide_default(self, slide):
        # Ne doit pas crasher, utilise un fallback
        add_alert(slide, "Titre", ["Item"], top=TOP_CONTENT, alert_type="inconnu")


class TestAddHighlight:
    def test_cree_shape(self, slide):
        n_before = len(slide.shapes)
        add_highlight(slide, "Message important", top=TOP_CONTENT)
        assert len(slide.shapes) > n_before

    def test_avec_url(self, slide):
        add_highlight(slide, "Lien", top=TOP_CONTENT, url="https://example.com")


class TestAddQuote:
    def test_avec_auteur(self, slide):
        n_before = len(slide.shapes)
        add_quote(slide, "Citation profonde", auteur="Voltaire", top=TOP_CONTENT)
        assert len(slide.shapes) > n_before

    def test_sans_auteur(self, slide):
        add_quote(slide, "Citation anonyme", top=TOP_CONTENT)


class TestAddCard:
    def test_cree_shape(self, slide):
        n_before = len(slide.shapes)
        add_card(slide, "Titre", ["Contenu"], top=TOP_CONTENT, left=MARGIN_L)
        assert len(slide.shapes) > n_before

    def test_avec_numero(self, slide):
        add_card(slide, "Titre", ["Contenu"], top=TOP_CONTENT, left=MARGIN_L,
                 numero="01")


class TestAddPaveChiffre:
    def test_cree_shape(self, slide):
        n_before = len(slide.shapes)
        add_pave_chiffre(slide, "42%", "des utilisateurs", top=TOP_CONTENT,
                         left=MARGIN_L)
        assert len(slide.shapes) > n_before


class TestAddStepper:
    def test_3_etapes(self, slide):
        n_before = len(slide.shapes)
        add_stepper(slide, [("E1", "Desc 1"), ("E2", "Desc 2"), ("E3", "Desc 3")],
                    top=TOP_CONTENT)
        assert len(slide.shapes) > n_before

    def test_etape_unique(self, slide):
        add_stepper(slide, [("Seule", "Description")], top=TOP_CONTENT)


class TestAddTableau:
    def test_cree_table(self, slide):
        n_before = len(slide.shapes)
        add_tableau(slide, ["Col A", "Col B"],
                    [["A1", "B1"], ["A2", "B2"]], top=TOP_CONTENT)
        assert len(slide.shapes) > n_before

    def test_retourne_hauteur(self, slide):
        h = add_tableau(slide, ["H1", "H2"], [["r1", "r2"]], top=TOP_CONTENT)
        assert isinstance(h, (int, float))
        assert h > 0


class TestAddTextLibre:
    def test_cree_shape(self, slide):
        n_before = len(slide.shapes)
        add_texte_libre(slide, "Texte positionne", top=TOP_CONTENT)
        assert len(slide.shapes) > n_before


class TestAddNotes:
    def test_ajoute_notes(self, slide):
        add_notes(slide, "Notes presentateur")
        assert slide.has_notes_slide
        notes_tf = slide.notes_slide.notes_text_frame
        assert "Notes presentateur" in notes_tf.text

    def test_multilignes(self, slide):
        add_notes(slide, "Ligne 1\nLigne 2\nLigne 3")
        assert "Ligne 2" in slide.notes_slide.notes_text_frame.text


class TestAddEncadre:
    def test_cree_shape(self, slide):
        n_before = len(slide.shapes)
        add_encadre(slide, top=TOP_CONTENT, left=MARGIN_L, width=CONTENT_W,
                    height=2.0, titre="Encadre", bullets=["Point 1"])
        assert len(slide.shapes) > n_before
