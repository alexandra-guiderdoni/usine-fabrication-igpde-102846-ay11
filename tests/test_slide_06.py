"""Tests de la slide 6 - tour de table."""

import sys
from pathlib import Path

from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from igpde_dsfr_components import create_presentation
from slides import SlideContext, load_slide_module


SLIDE_PATH = (
    Path(__file__).parent.parent / "scripts" / "slides" / "02d_tour-de-table.py"
)


def test_slide_6_cartes_plus_compactes_et_contenus_plus_aeres():
    prs, layouts = create_presentation()
    ctx = SlideContext(
        page_num=6,
        date="9 octobre 2026",
        footer_base="Formation 102846",
        formation_code="102846",
    )

    slide = load_slide_module(SLIDE_PATH).build(prs, layouts, ctx)

    cartes = [shape for shape in slide.shapes if shape.name == "DSFR-box"]
    assert len(cartes) == 3
    assert all(shape.height == Inches(3.45) for shape in cartes)
    assert all(shape.top + shape.height <= Inches(5.75) for shape in cartes)

    contenus = [shape for shape in slide.shapes if shape.name == "DSFR-card-contenu"]
    assert len(contenus) == 3
    assert all(shape.text_frame.paragraphs[0].line_spacing == 1.5 for shape in contenus)
