"""Tests du post-traitement a11y (finalize_pptx et helpers associes)."""

import sys
import tempfile
from pathlib import Path

import pytest
from lxml import etree

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from igpde_dsfr_components import (
    create_presentation,
    new_slide,
    add_callout,
    add_notes,
    finalize_pptx,
    _set_lang_on_runs,
    _mark_decoratives,
    _reorder_shapes,
    NSMAP_A,
    MARGIN_L,
    CONTENT_W,
    TOP_CONTENT,
)


@pytest.fixture
def prs_with_content():
    """Presentation avec une slide contenant du contenu."""
    prs, layouts = create_presentation()
    slide = new_slide(prs, layouts, layout_name="titre_contenu",
                      titre="Slide test a11y", page_num=1)
    add_callout(slide, "Rappel", ["Point important"], top=TOP_CONTENT)
    add_notes(slide, "Notes du presentateur")
    return prs


class TestSetLangOnRuns:
    def test_applique_lang_fr(self, slide):
        _set_lang_on_runs(slide._element, lang="fr-FR")
        r_tag = f"{{{NSMAP_A}}}r"
        rPr_tag = f"{{{NSMAP_A}}}rPr"
        for r in slide._element.iter(r_tag):
            rPr = r.find(rPr_tag)
            assert rPr is not None
            assert rPr.get("lang") == "fr-FR"

    def test_lang_personnalisee(self, slide):
        _set_lang_on_runs(slide._element, lang="en-US")
        r_tag = f"{{{NSMAP_A}}}r"
        rPr_tag = f"{{{NSMAP_A}}}rPr"
        for r in slide._element.iter(r_tag):
            rPr = r.find(rPr_tag)
            if rPr is not None:
                assert rPr.get("lang") == "en-US"


class TestMarkDecoratives:
    def test_shape_decoratif_reçoit_alt_vide(self, slide):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches

        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), Inches(1), Inches(1),
        )
        shape.name = "accent-decoratif-gauche"
        _mark_decoratives(slide)
        # Verifier que descr="" sur le cNvPr
        cNvPr = shape._element.find(f".//{{{NSMAP_A}}}cNvPr")
        if cNvPr is None:
            from igpde_dsfr_components import NSMAP_P
            cNvPr = shape._element.find(f".//{{{NSMAP_P}}}cNvPr")
        assert cNvPr is not None
        assert cNvPr.get("descr") == ""

    def test_shape_normal_pas_touche(self, slide):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches

        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), Inches(1), Inches(1),
        )
        shape.name = "contenu-principal"
        _mark_decoratives(slide)
        cNvPr = shape._element.find(f".//{{{NSMAP_A}}}cNvPr")
        if cNvPr is not None:
            assert cNvPr.get("descr") != ""  # pas modifie a vide


class TestReorderShapes:
    def test_titre_avant_contenu(self, slide):
        _reorder_shapes(slide)
        spTree = slide.shapes._spTree
        names = []
        for child in spTree:
            tag = etree.QName(child.tag).localname
            if tag in ("sp", "pic", "graphicFrame"):
                cNvPr = child.find(f".//{{{NSMAP_A}}}cNvPr")
                if cNvPr is not None:
                    names.append(cNvPr.get("name", ""))
        # Le titre doit apparaitre avant les contenus
        titre_idx = None
        for i, n in enumerate(names):
            if n.startswith("Titre") or "chapitre-titre" in n.lower():
                titre_idx = i
                break
        if titre_idx is not None:
            # Tous les contenus apres le titre
            for i, n in enumerate(names[titre_idx + 1:], titre_idx + 1):
                if "decoratif" in n.lower():
                    for j, n2 in enumerate(names[i + 1:], i + 1):
                        assert "decoratif" in n2.lower() or "pied" in n2.lower() \
                            or "numero" in n2.lower() or "date" in n2.lower(), \
                            f"Shape non-decoratif '{n2}' apres decoratif '{n}'"


class TestFinalizePptx:
    def test_sauvegarde_fichier(self, prs_with_content, tmp_path):
        output = tmp_path / "test-output.pptx"
        result = finalize_pptx(prs_with_content, str(output),
                               title="Test", author="Test")
        assert output.exists()
        assert output.stat().st_size > 0
        assert str(result) == str(output)

    def test_core_properties(self, prs_with_content, tmp_path):
        from pptx import Presentation as PrsLoad

        output = tmp_path / "test-props.pptx"
        finalize_pptx(prs_with_content, str(output),
                      title="Mon titre", author="Alex", subject="A11y")
        prs_reload = PrsLoad(str(output))
        assert prs_reload.core_properties.title == "Mon titre"
        assert prs_reload.core_properties.author == "Alex"
        assert prs_reload.core_properties.subject == "A11y"

    def test_lang_appliquee_sur_runs(self, prs_with_content, tmp_path):
        from pptx import Presentation as PrsLoad

        output = tmp_path / "test-lang.pptx"
        finalize_pptx(prs_with_content, str(output), lang="fr-FR")
        prs_reload = PrsLoad(str(output))
        slide = prs_reload.slides[0]
        r_tag = f"{{{NSMAP_A}}}r"
        rPr_tag = f"{{{NSMAP_A}}}rPr"
        langs_found = set()
        for r in slide._element.iter(r_tag):
            rPr = r.find(rPr_tag)
            if rPr is not None and rPr.get("lang"):
                langs_found.add(rPr.get("lang"))
        assert "fr-FR" in langs_found

    def test_notes_recoivent_lang(self, prs_with_content, tmp_path):
        from pptx import Presentation as PrsLoad

        output = tmp_path / "test-notes-lang.pptx"
        finalize_pptx(prs_with_content, str(output), lang="fr-FR")
        prs_reload = PrsLoad(str(output))
        slide = prs_reload.slides[0]
        assert slide.has_notes_slide
        r_tag = f"{{{NSMAP_A}}}r"
        rPr_tag = f"{{{NSMAP_A}}}rPr"
        for r in slide.notes_slide._element.iter(r_tag):
            rPr = r.find(rPr_tag)
            if rPr is not None:
                assert rPr.get("lang") == "fr-FR"
