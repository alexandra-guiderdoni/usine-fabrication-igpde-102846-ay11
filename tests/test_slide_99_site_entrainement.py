"""Tests du lien vers le site d'entraînement."""

import sys
from pathlib import Path

from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from config import load_formation_config
from igpde_dsfr_components import create_presentation
from slides import SlideContext, discover_slides, load_slide_module


SLIDE_FILE = "40_signaux-alerte.py"
SITE_ENTRAINEMENT = load_formation_config()["site_url"]


def test_slide_signaux_alerte_reste_dans_la_sequence_clavier():
    fichiers = [chemin.name for chemin in discover_slides()]

    assert fichiers.index("39_cinq-touches.py") < fichiers.index(SLIDE_FILE)
    assert fichiers.index(SLIDE_FILE) < fichiers.index("41_mission-clavier.py")


def test_slide_signaux_alerte_affiche_le_lien_et_son_qrcode_sous_les_cartes():
    prs, layouts = create_presentation()
    ctx = SlideContext(
        page_num=93,
        date="9 octobre 2026",
        footer_base="Formation 102846",
        formation_code="102846",
    )
    slide = load_slide_module(
        Path(__file__).parent.parent / "scripts" / "slides" / SLIDE_FILE
    ).build(prs, layouts, ctx)

    qrcode = next(shape for shape in slide.shapes if shape.name == "DSFR-qrcode")
    url = next(
        shape for shape in slide.shapes if shape.name == "DSFR-qrcode-url-visible"
    )
    cartes = [shape for shape in slide.shapes if shape.name == "DSFR-box"]

    assert url.text == SITE_ENTRAINEMENT
    assert qrcode.top - max(carte.top + carte.height for carte in cartes) >= Inches(
        0.20
    )
    assert qrcode.top + qrcode.height <= Inches(6.30)

    taille_url = url.text_frame.paragraphs[0].runs[0].font.size
    assert taille_url is not None
    assert taille_url.pt >= 10

    c_nv_pr = qrcode._element.find(
        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"
    )
    assert c_nv_pr.get("descr") == f"QR code : {SITE_ENTRAINEMENT}"
    assert slide.notes_slide.notes_text_frame.text.strip()
