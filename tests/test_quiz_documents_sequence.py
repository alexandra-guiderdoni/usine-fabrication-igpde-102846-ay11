"""Tests de la séquence question-réponse sur les documents accessibles."""

import sys
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from igpde_dsfr_components import create_presentation
from slides import SlideContext, discover_slides, load_slide_module


OPENING_FILE = "04_ouverture-lecteur-ecran.py"
QUESTION_FILE = "05_quiz-flash-a-vs-b.py"
ANSWER_FILE = "05a_quiz-flash-reponse.py"
NEXT_FILE = "06_pourquoi-concerne.py"


def _build_slide(filename, page_num):
    prs, layouts = create_presentation()
    ctx = SlideContext(
        page_num=page_num,
        date="9 octobre 2026",
        footer_base="Formation 102846",
        formation_code="102846",
    )
    return load_slide_module(
        Path(__file__).parent.parent / "scripts" / "slides" / filename
    ).build(prs, layouts, ctx)


def test_quiz_et_reponse_occupent_les_slides_56_et_57():
    fichiers = [chemin.name for chemin in discover_slides()]

    assert fichiers[55:58] == [QUESTION_FILE, ANSWER_FILE, NEXT_FILE]
    assert len(fichiers) == 131


def test_slide_55_presente_deux_cartes_cote_a_cote():
    slide = _build_slide(OPENING_FILE, 55)

    cartes = sorted(
        (shape for shape in slide.shapes if shape.name == "DSFR-box"),
        key=lambda shape: shape.left,
    )
    assert len(cartes) == 2
    assert cartes[0].top == cartes[1].top
    assert cartes[0].height == cartes[1].height
    assert cartes[0].left < cartes[1].left

    textes = "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )
    assert "Ce qu’entend le lecteur d’écran" in textes
    assert "Texte, texte, texte" in textes
    assert "15 % de vos destinataires sont concernés" in textes
    assert "80 % de ces handicaps sont invisibles" in textes
    assert slide.notes_slide.notes_text_frame.text.strip()


def test_slide_question_montre_deux_documents_sans_reveler_la_reponse():
    slide = _build_slide(QUESTION_FILE, 56)

    assert slide.shapes.title.text == (
        "Quiz - lequel de ces deux documents est accessible ?"
    )
    images = [
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(images) == 2
    assert all(image.width >= Inches(2.10) for image in images)
    assert all(image.top + image.height <= Inches(6.60) for image in images)
    textes = "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )
    assert "Document A" in textes
    assert "Document B" in textes
    assert "Titre 1" not in textes
    assert "texte alternatif" not in textes
    assert "Votre réponse ?" not in textes
    assert slide.notes_slide.notes_text_frame.text.strip()


def test_slide_reponse_explique_pourquoi_le_document_b_est_accessible():
    slide = _build_slide(ANSWER_FILE, 57)

    assert slide.shapes.title.text == "Réponse - le document B est accessible"
    textes = "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )
    assert "Titres mis en gras, police Arial 16" in textes
    assert "Titres avec le style « Titre 1 »" in textes
    assert "Image avec texte alternatif" in textes
    assert "rapport-bilan-2024.docx" in textes
    assert slide.notes_slide.notes_text_frame.text.strip()
