"""Reconstruit PPT-IGPDE-DSFR-base-intervenant.pptx depuis gabarits-ppt-igpde.pptx.

Utile quand le source IGPDE original a été déplacé. Le PPTX de gabarits porte
déjà les layouts rescalés (13,33" × 7,5"), il suffit de supprimer ses slides.
"""

from pathlib import Path
import shutil
from pptx import Presentation

SOURCES = Path(__file__).resolve().parent.parent / "_source" / "presentations-source"
SRC = SOURCES / "gabarits-ppt-igpde.pptx"
DST = SOURCES / "PPT-IGPDE-DSFR-base-intervenant.pptx"

NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def main():
    if DST.exists():
        DST.unlink()
    shutil.copy(SRC, DST)
    prs = Presentation(str(DST))

    sldIdLst = prs.slides._sldIdLst
    slide_ids = list(sldIdLst)
    for sldId in slide_ids:
        rId = sldId.get("{%s}id" % NS_R)
        try:
            prs.part.drop_rel(rId)
        except Exception:
            pass
        sldIdLst.remove(sldId)

    prs.save(str(DST))
    prs2 = Presentation(str(DST))
    print(f"[OK] Template reconstruit : {DST.name}")
    print(f"     Slides : {len(prs2.slides)} (doit être 0)")
    print(f"     Layouts : {len(prs2.slide_masters[0].slide_layouts)}")


if __name__ == "__main__":
    main()
