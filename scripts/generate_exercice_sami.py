"""Générateur des fichiers de l'exercice Sami.

Produit :
- _assets/graphique-inaccessible.png  (barres couleurs seules)
- _assets/graphique-accessible.png    (barres avec motifs + étiquettes)
- _assets/icone-enveloppe.png         (icône e-mail)
- _assets/organigramme.png            (organigramme du service)
- tp-doc-inaccessible.docx           (version fautive sans aide)
- tp-doc-aide-correction.docx        (version fautive annotée)
- tp-doc-accessible.docx             (version corrigée)
"""

# PDG-LARGE-FILE-JUSTIFICATION: générateur unique des trois variantes dont les
# helpers OOXML et l'ordre canonique doivent rester comparables dans un même flux.

import argparse
import re
import subprocess
from pathlib import Path
from xml.sax.saxutils import escape

import matplotlib

matplotlib.use("Agg")
import matplotlib.patches as mpatches
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from exercice_sami_matrice import load_sami_matrix
from lxml import etree

from config import load_formation_config

PROJECT = Path(__file__).resolve().parent.parent
ASSETS = PROJECT / "_assets"
ASSETS.mkdir(exist_ok=True)
CHECKLIST_BASENAME = "checklist-accessibilite-bureautique"
CHECKLIST_MARKDOWN = PROJECT / "_source" / f"{CHECKLIST_BASENAME}.md"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
ADEC_NS = "http://schemas.microsoft.com/office/drawing/2017/decorative"
DECORATIVE_EXT_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"


# ------------------------------------------------------------------
# 1. Graphiques PNG
# ------------------------------------------------------------------

INDICATEURS = ["Accès\ndirects", "Moteurs de\nrecherche", "Sites\nréférents"]
T4_2024 = [18200, 21000, 6000]
T1_2025 = [20400, 23400, 6800]
EVOL = ["+12 %", "+11 %", "+13 %"]
CHART_TITLE = "Évolution du trafic web"
CHART_ALT_INACCESSIBLE = "Graphique : évolution du trafic web entre T4 2024 et T1 2025."
CHART_ALT_ACCESSIBLE = (
    "T4 2024 puis T1 2025 : accès directs, 18 200 puis 20 400 (+12 %) ; "
    "moteurs de recherche, 21 000 puis 23 400 (+11 %) ; sites référents, "
    "6 000 puis 6 800 (+13 %)."
)
CONTRAST_SAMPLE_TEXT = "Information complémentaire : résultats provisoires."


def generate_chart_inaccessible():
    """Barres dont les deux séries ne se distinguent que par la couleur."""
    fig, ax = plt.subplots(figsize=(5, 3))
    x = np.arange(len(INDICATEURS))
    w = 0.35
    bars1 = ax.bar(
        x - w / 2,
        T4_2024,
        w,
        color="#D00000",
        label="T4 2024",
    )
    bars2 = ax.bar(
        x + w / 2,
        T1_2025,
        w,
        color="#18753C",
        label="T1 2025",
    )
    ax.bar_label(
        bars1,
        labels=[f"{value:,}".replace(",", " ") for value in T4_2024],
        padding=2,
        fontsize=6,
    )
    ax.bar_label(
        bars2,
        labels=[f"{value:,}".replace(",", " ") for value in T1_2025],
        padding=2,
        fontsize=6,
    )
    ax.set_xticks(x)
    ax.set_xticklabels(INDICATEURS, fontsize=8)
    ax.set_title(CHART_TITLE, fontsize=10)
    ax.set_ylim(0, max(T1_2025) * 1.25)
    ax.tick_params(axis="y", labelsize=7)
    ax.legend(fontsize=7, loc="upper left")
    fig.tight_layout()
    path = ASSETS / "graphique-inaccessible.png"
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def generate_chart_accessible():
    """Barres avec motifs distincts et étiquettes sur chaque barre."""
    fig, ax = plt.subplots(figsize=(5, 3))
    x = np.arange(len(INDICATEURS))
    w = 0.35
    bars1 = ax.bar(
        x - w / 2,
        T4_2024,
        w,
        color="#6C6C6C",
        edgecolor="black",
        linewidth=0.8,
        hatch="///",
        label="T4 2024",
    )
    bars2 = ax.bar(
        x + w / 2,
        T1_2025,
        w,
        color="#B0B0B0",
        edgecolor="black",
        linewidth=0.8,
        hatch="...",
        label="T1 2025",
    )

    for bar, val in zip(bars1, T4_2024):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height() + 500,
            f"{val:,}".replace(",", " "),
            ha="center",
            va="bottom",
            fontsize=6,
        )
    for bar, val, ev in zip(bars2, T1_2025, EVOL):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height() + 500,
            f"{val:,}".replace(",", " ") + f" ({ev})",
            ha="center",
            va="bottom",
            fontsize=6,
        )

    ax.set_xticks(x)
    ax.set_xticklabels(INDICATEURS, fontsize=8)
    ax.set_title(CHART_TITLE, fontsize=10)
    ax.set_ylim(0, max(T1_2025) * 1.25)
    ax.tick_params(axis="y", labelsize=7)
    ax.legend(fontsize=7, loc="upper left")
    fig.tight_layout()
    path = ASSETS / "graphique-accessible.png"
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


# ------------------------------------------------------------------
# 2. Icone enveloppe + organigramme
# ------------------------------------------------------------------


def generate_icon_enveloppe():
    """Petite icone d'enveloppe pour le paragraphe contact."""
    fig, ax = plt.subplots(figsize=(0.5, 0.5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 10)
    ax.set_aspect("equal")
    ax.axis("off")
    # Corps de l'enveloppe
    rect = mpatches.FancyBboxPatch(
        (1, 2), 8, 5, boxstyle="round,pad=0.3", facecolor="#000091", edgecolor="#000091"
    )
    ax.add_patch(rect)
    # Rabat triangulaire
    ax.plot([1, 5, 9], [7, 3.5, 7], color="white", linewidth=1.5)
    fig.subplots_adjust(left=0, right=1, top=1, bottom=0)
    path = ASSETS / "icone-enveloppe.png"
    fig.savefig(path, dpi=100, transparent=True, bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)
    return path


def generate_organigramme():
    """Organigramme simple de la Direction des affaires juridiques."""
    fig, ax = plt.subplots(figsize=(6, 3.5))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 7)
    ax.axis("off")

    box_style = dict(boxstyle="round,pad=0.4", facecolor="#000091", edgecolor="#000091")
    text_kw = dict(
        ha="center",
        va="center",
        fontsize=8,
        color="white",
        fontweight="bold",
        bbox=box_style,
    )

    ax.text(6, 6, "Direction des\naffaires juridiques", **text_kw)

    equipes = [
        (1.5, 2.5, "Bureau du\ndroit public"),
        (4.5, 2.5, "Bureau du\ndroit social"),
        (7.5, 2.5, "Bureau de la\ncommunication"),
        (10.5, 2.5, "Bureau des\naffaires internationales"),
    ]
    for x, y, label in equipes:
        ax.text(x, y, label, **text_kw)
        ax.plot([x, x], [y + 0.8, 5.2], color="#000091", linewidth=1.5)

    ax.plot([1.5, 10.5], [5.2, 5.2], color="#000091", linewidth=1.5)
    ax.plot([6, 6], [5.2, 5.5], color="#000091", linewidth=1.5)

    fig.tight_layout()
    path = ASSETS / "organigramme.png"
    fig.savefig(path, dpi=150, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    return path


def generate_texte_image():
    """Genere une image contenant du texte (avis important) pour l'erreur texte-image."""
    fig, ax = plt.subplots(figsize=(5, 1.5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 3)
    ax.axis("off")
    rect = mpatches.FancyBboxPatch(
        (0.2, 0.2),
        9.6,
        2.6,
        boxstyle="round,pad=0.3",
        facecolor="#FFF3CD",
        edgecolor="#856404",
        linewidth=1.5,
    )
    ax.add_patch(rect)
    ax.text(
        5,
        1.5,
        "Avis important : les indicateurs du T2 2025\n"
        "seront transmis avant le 15 septembre 2025.",
        ha="center",
        va="center",
        fontsize=11,
        fontweight="bold",
        color="#856404",
    )
    fig.subplots_adjust(left=0, right=1, top=1, bottom=0)
    path = ASSETS / "texte-image.png"
    fig.savefig(path, dpi=150, facecolor="white", bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)
    return path


# ------------------------------------------------------------------
# 3. Document inaccessible
# ------------------------------------------------------------------


def _set_image_alt(doc, alt_text="", title=""):
    """Pose alt text et titre sur la derniere image inseree."""
    inline_shape = doc.inline_shapes[-1]
    pic = inline_shape._inline
    docPr = pic.find(qn("wp:docPr"))
    if docPr is not None:
        docPr.set("descr", alt_text)
        if title:
            docPr.set("title", title)


def _mark_image_decorative(doc):
    """Marque la derniere image comme decorative via l'extension Office."""
    _set_image_alt(doc, alt_text="", title="")
    inline_shape = doc.inline_shapes[-1]
    docPr = inline_shape._inline.find(qn("wp:docPr"))
    if docPr is None:
        return

    extLst = docPr.find(f"{{{A_NS}}}extLst")
    if extLst is None:
        extLst = etree.SubElement(docPr, f"{{{A_NS}}}extLst")
    for ext in list(extLst.findall(f"{{{A_NS}}}ext")):
        if ext.get("uri") == DECORATIVE_EXT_URI:
            extLst.remove(ext)
    ext = parse_xml(
        f'<a:ext xmlns:a="{A_NS}" uri="{DECORATIVE_EXT_URI}">'
        f'<adec:decorative xmlns:adec="{ADEC_NS}" val="1"/>'
        f"</a:ext>"
    )
    extLst.append(ext)


def _set_style_language(style, lang):
    """Pose la langue sur un style Word."""
    rPr = style.element.get_or_add_rPr()
    lang_el = rPr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = parse_xml(
            f'<w:lang {nsdecls("w")} w:val="{lang}" '
            f'w:eastAsia="{lang}" w:bidi="{lang}"/>'
        )
        rPr.append(lang_el)
    else:
        lang_el.set(qn("w:val"), lang)
        lang_el.set(qn("w:eastAsia"), lang)
        lang_el.set(qn("w:bidi"), lang)


def _set_doc_defaults_language(doc, lang):
    """Pose la langue par défaut du document dans styles.xml."""
    styles = doc.styles.element
    doc_defaults = styles.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = parse_xml(f"<w:docDefaults {nsdecls('w')}/>")
        styles.insert(0, doc_defaults)

    rPr_default = doc_defaults.find(qn("w:rPrDefault"))
    if rPr_default is None:
        rPr_default = parse_xml(f"<w:rPrDefault {nsdecls('w')}/>")
        doc_defaults.insert(0, rPr_default)

    rPr = rPr_default.find(qn("w:rPr"))
    if rPr is None:
        rPr = parse_xml(f"<w:rPr {nsdecls('w')}/>")
        rPr_default.append(rPr)

    lang_el = rPr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = parse_xml(
            f'<w:lang {nsdecls("w")} w:val="{lang}" '
            f'w:eastAsia="{lang}" w:bidi="{lang}"/>'
        )
        rPr.append(lang_el)
    else:
        lang_el.set(qn("w:val"), lang)
        lang_el.set(qn("w:eastAsia"), lang)
        lang_el.set(qn("w:bidi"), lang)


def _set_content_styles_language(doc, lang):
    """Pose la langue sur les styles réellement utilisés dans le document."""
    for style_name in (
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Heading 4",
        "List Bullet",
        "List Number",
    ):
        _set_style_language(doc.styles[style_name], lang)


def _add_hyperlink(paragraph, text, url):
    """Ajoute un vrai lien hypertexte Word avec style visuel standard."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = parse_xml(
        f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}">'
        f"<w:r>"
        f'<w:rPr><w:color w:val="0000FF"/><w:u w:val="single"/></w:rPr>'
        f"<w:t>{escape(text)}</w:t>"
        f"</w:r>"
        f"</w:hyperlink>"
    )
    paragraph._p.append(hyperlink)


def _add_toc(doc):
    """Insere un champ Table des matieres automatique."""
    p = doc.add_paragraph()
    run = p.add_run()
    fldChar_begin = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run._r.append(fldChar_begin)
    run2 = p.add_run()
    instrText = parse_xml(
        f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>'
    )
    run2._r.append(instrText)
    run3 = p.add_run()
    fldChar_separate = parse_xml(
        f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'
    )
    run3._r.append(fldChar_separate)
    run4 = p.add_run("(Table des matières - mettre à jour avec F9)")
    run4.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    run4.font.size = Pt(10)
    run5 = p.add_run()
    fldChar_end = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run5._r.append(fldChar_end)


def _add_simple_field(paragraph, instr):
    """Ajoute un champ Word simple dans un paragraphe."""
    run_begin = paragraph.add_run()
    run_begin._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
    run_instr = paragraph.add_run()
    run_instr._r.append(
        parse_xml(
            f'<w:instrText {nsdecls("w")} xml:space="preserve"> {instr} </w:instrText>'
        )
    )
    run_sep = paragraph.add_run()
    run_sep._r.append(
        parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    )
    run_result = paragraph.add_run("1")
    run_end = paragraph.add_run()
    run_end._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))
    return run_result


GUIDE_TITLE = "Rendre un document Word accessible"
DOCUMENT_TITLE = GUIDE_TITLE
HEADER_TEXT = "Guide pratique - Rendre un document Word accessible"
AFFICHE_P06 = ASSETS / "affiche-sig-handicap.jpg"
STATUT_DOCUMENT = "Document confidentiel"
AFFICHE_ALT = (
    "Affiche des 20 ans de la loi handicap : Agents publics, changeons "
    "cette réalité grâce aux outils disponibles sur accessibilite.gouv.fr. "
    "Le handicap n'est pas un choix. L'accessibilité non plus."
)
INTRO_TEXT = (
    "Ce guide pratique présente les principales vérifications à effectuer pour "
    "rendre un document Word accessible. Il associe chaque règle à une "
    "manipulation et à une preuve de correction."
)


def _format_header_footer_run(run, font_size=9):
    """Applique le style explicite aux textes d'en-tête et pied de page."""
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def _format_header_paragraph(paragraph, font_size=9):
    for run in paragraph.runs:
        _format_header_footer_run(run, font_size)


def _add_header_status(header, font_size=9):
    """Porte le statut du document dans l'en-tête, support de P-11."""
    paragraph = header.add_paragraph(STATUT_DOCUMENT)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _format_header_paragraph(paragraph, font_size=font_size)


def _add_page_footer(doc, document_name=DOCUMENT_TITLE, font_size=9, separator=" / "):
    """Ajoute un pied de page Page X / Y avec champs Word natifs."""
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if document_name:
        p.add_run(document_name)
        p.add_run(" - ")
    p.add_run("Page ")
    _add_simple_field(p, "PAGE")
    p.add_run(separator)
    _add_simple_field(p, "NUMPAGES")
    for run in p.runs:
        _format_header_footer_run(run, font_size)


def _add_fake_list_item(doc, marker, text):
    """Ajoute un item tape manuellement mais aligne comme une vraie liste."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    run = p.add_run(f"{marker}\t{text}")
    run.font.name = "Arial"
    run.font.size = Pt(11)
    return p


# Identifiant en tête de piste et intitulés des rubriques, mis en gras.
GUIDANCE_LABELS = re.compile(
    r"^(?:Document \u2014 )?P-\d+"
    r"|(?<=\s)(?:Problème|Impact|Règle|Piste|Première action|Procédure Word) :"
)


def _guidance_segments(text):
    """Découpe une piste en segments (texte, gras) autour des intitulés."""
    segments, position = [], 0
    for match in GUIDANCE_LABELS.finditer(text):
        if match.start() > position:
            segments.append((text[position : match.start()], False))
        segments.append((match.group(), True))
        position = match.end()
    if position < len(text):
        segments.append((text[position:], False))
    return segments


def _add_guidance_comment(doc, runs, text):
    """Ajoute un commentaire Word de correction sur un ou plusieurs runs."""
    if not runs:
        return
    if not isinstance(runs, (list, tuple)):
        runs = [runs]
    runs = [run for run in runs if run is not None]
    if runs:
        comment = doc.add_comment(
            runs,
            text="",
            author="Formation IGPDE",
            initials="IGPDE",
        )
        for segment, bold in _guidance_segments(text):
            run = comment.paragraphs[0].add_run(segment)
            if bold:
                run.bold = True
        # Le commentaire est rédigé en français : sans langue propre, il hérite
        # du défaut de langue volontaire du document et Word le souligne.
        # Corps 14 pt et interligne double pour la lisibilité des pistes.
        for paragraph in comment.paragraphs:
            paragraph.paragraph_format.line_spacing = 2.0
            for run in paragraph.runs:
                run.font.size = Pt(14)
                run._element.get_or_add_rPr().append(
                    parse_xml(
                        f'<w:lang {nsdecls("w")} w:val="fr-FR" '
                        f'w:eastAsia="fr-FR" w:bidi="fr-FR"/>'
                    )
                )


def _station_controls(matrix, station_id):
    """Retourne les contrôles d'une station dans l'ordre canonique."""
    return [
        control for control in matrix["controles"] if control["station"] == station_id
    ]


def _station_title(matrix, station_id):
    """Retourne le titre canonique d'une station."""
    return next(
        block["titre"] for block in matrix["sequence"] if block["id"] == station_id
    )


def _guidance_text(control, *, document_level=False, precision=None):
    """Compose une piste à partir du contrôle canonique."""
    prefix = f"Document {chr(0x2014)} " if document_level else ""
    intitule = (
        f"{control['intitule']} ({precision})" if precision else control["intitule"]
    )
    return (
        f"{prefix}{control['id']} - {intitule}. "
        f"Problème : {control['defaut']} Impact : {control['impact']} "
        f"Règle : {control['regle']} Piste : {control['piste']} "
        f"Première action : {control['action_attendue']} Procédure Word : "
        f"{control['procedure_word']}"
    )


def _add_control_details(doc, control):
    """Ajoute le contenu éditorial commun d'un contrôle."""
    problem = control["defaut"] or (
        "Cette étape de finalisation ne peut pas être prouvée par le DOCX seul."
    )
    for label, value in (
        ("Problème", problem),
        ("Pourquoi", control["impact"]),
        ("Règle", control["regle"]),
        ("Dans Word", control["procedure_word"]),
        ("Dans LibreOffice Writer", control["procedure_writer"]),
        ("À faire", control["action_attendue"]),
        ("Preuve", control["preuve"]["attendu"]),
    ):
        # Intitulé en gras pour repérer chaque rubrique d'un coup d'œil.
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{label} :").bold = True
        paragraph.add_run(f" {value}")


def _create_heading_numbering(doc):
    """Crée une numérotation multiniveau native pour les titres de station."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(item.get(qn("w:abstractNumId")))
        for item in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    levels = []
    for level in range(4):
        level_text = ".".join(f"%{index}" for index in range(1, level + 2)) + "."
        levels.append(
            f'<w:lvl w:ilvl="{level}">'
            '<w:start w:val="1"/>'
            '<w:numFmt w:val="decimal"/>'
            f'<w:lvlText w:val="{level_text}"/>'
            '<w:suff w:val="space"/>'
            "</w:lvl>"
        )
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        '<w:multiLevelType w:val="multilevel"/>'
        f"{''.join(levels)}"
        "</w:abstractNum>"
    )
    num = parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}">'
        f'<w:abstractNumId w:val="{abstract_id}"/>'
        "</w:num>"
    )
    numbering.append(abstract)
    numbering.append(num)
    return num_id


def _apply_heading_numbering(paragraph, num_id, level):
    """Associe un titre à un niveau de la numérotation de station."""
    p_pr = paragraph._p.get_or_add_pPr()
    current = p_pr.find(qn("w:numPr"))
    if current is not None:
        p_pr.remove(current)
    p_pr.append(
        parse_xml(
            f"<w:numPr {nsdecls('w')}>"
            f'<w:ilvl w:val="{level}"/>'
            f'<w:numId w:val="{num_id}"/>'
            "</w:numPr>"
        )
    )


def _add_station_one_p01(
    doc,
    control,
    station_title,
    numbering_id,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute le titre principal et le premier point du guide."""
    if corrected:
        title = doc.add_paragraph(GUIDE_TITLE, style="Title")
        station_heading = doc.add_heading(station_title, level=1)
        control_heading = doc.add_heading(
            f"{control['id']} - {control['intitule']}", level=2
        )
    else:
        title = doc.add_paragraph()
        title_run = title.add_run(GUIDE_TITLE)
        title_run.bold = True
        title_run.font.name = "Arial"
        title_run.font.size = Pt(20)
        station_heading = doc.add_heading(station_title, level=1)
        control_heading = doc.add_heading(
            f"{control['id']} - {control['intitule']}", level=2
        )

    _apply_heading_numbering(station_heading, numbering_id, 0)
    _apply_heading_numbering(control_heading, numbering_id, 1)

    _add_control_details(doc, control)

    if with_guidance:
        _add_guidance_comment(doc, title.runs, _guidance_text(control))


def _add_station_one_p02(
    doc, control, numbering_id, *, corrected=False, with_guidance=False
):
    """Ajoute le contrôle de hiérarchie avec ou sans saut de niveau."""
    level = 2 if corrected else 4
    heading = doc.add_heading(f"{control['id']} - {control['intitule']}", level=level)
    _apply_heading_numbering(heading, numbering_id, 1 if corrected else 3)
    _add_control_details(doc, control)
    if with_guidance:
        _add_guidance_comment(doc, heading.runs, _guidance_text(control))


def _add_station_one_control(doc, control, numbering_id):
    """Ajoute un point de station avec son titre structurel numéroté."""
    text = f"{control['id']} - {control['intitule']}"
    heading = doc.add_heading(text, level=2)
    _apply_heading_numbering(heading, numbering_id, 1)
    _add_control_details(doc, control)
    return heading


def _set_section_columns(section, count):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = parse_xml(f"<w:cols {nsdecls('w')}/>")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(count))


def _add_station_one_p05(
    doc, control, numbering_id, *, corrected=False, with_guidance=False
):
    """Ajoute l'occurrence composite de mise en page de la station 1."""
    _add_station_one_control(doc, control, numbering_id)
    zone_heading = doc.add_heading("Mise en page robuste", level=3)
    if corrected:
        spacing = doc.add_paragraph("Espacement entre les paragraphes")
        spacing.paragraph_format.space_after = Pt(12)

        indent = doc.add_paragraph(
            "Retrait du paragraphe aligné avec le contenu précédent."
        )
        indent.paragraph_format.left_indent = Inches(0.5)

        doc.add_paragraph("Ligne principale")
        doc.add_paragraph("Suite sur un paragraphe distinct")

        next_page = doc.add_paragraph("Début de la page suivante")
        next_page.paragraph_format.page_break_before = True

        two_columns = doc.add_section(WD_SECTION.CONTINUOUS)
        _set_section_columns(two_columns, 2)
        doc.add_paragraph("Colonne gauche : structure")
        doc.add_paragraph("Colonne droite : navigation")
        one_column = doc.add_section(WD_SECTION.CONTINUOUS)
        _set_section_columns(one_column, 1)
    else:
        doc.add_paragraph("Espacement entre les paragraphes")
        for _ in range(4):
            doc.add_paragraph()

        doc.add_paragraph("Retrait du paragraphe    aligné avec le contenu précédent.")

        line = doc.add_paragraph()
        line_run = line.add_run("Ligne principale")
        line_run.add_break()
        line_run.add_text("Suite sur un paragraphe distinct")

        for _ in range(4):
            doc.add_paragraph()
        doc.add_paragraph("Début de la page suivante")

        columns = doc.add_paragraph()
        columns.add_run("Colonne gauche : structure")
        columns.add_run().add_tab()
        columns.add_run("Colonne droite : navigation")

    doc.add_paragraph("Fin de la zone Mise en page robuste.")
    if with_guidance:
        _add_guidance_comment(doc, zone_heading.runs, _guidance_text(control))


def _add_station_two_p06(
    doc,
    control,
    station_title,
    numbering_id,
    image_path,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute l'image informative simple de la station 2."""
    station_heading = doc.add_heading(station_title, level=1)
    _apply_heading_numbering(station_heading, numbering_id, 0)
    _add_station_one_control(doc, control, numbering_id)
    occurrence_paragraph = None

    # L'affiche porte son message dans l'image : il ne peut pas être remplacé
    # par du texte réel (cas de P-09), l'alternative doit le restituer.
    if image_path:
        doc.add_picture(str(image_path), width=Inches(2.2))
        occurrence_paragraph = doc.paragraphs[-1]
        occurrence_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if corrected:
            _set_image_alt(
                doc,
                alt_text=AFFICHE_ALT,
                title="Affiche des 20 ans de la loi handicap",
            )
        else:
            _set_image_alt(doc, alt_text="", title="")

    if with_guidance and occurrence_paragraph is not None:
        _add_guidance_comment(doc, occurrence_paragraph.runs, _guidance_text(control))


def _add_station_two_p07(
    doc,
    control,
    numbering_id,
    organigramme_path,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute l'image complexe et son équivalent textuel."""
    _add_station_one_control(doc, control, numbering_id)
    occurrence_paragraph = None

    if organigramme_path:
        doc.add_picture(str(organigramme_path), width=Inches(5.0))
        occurrence_paragraph = doc.paragraphs[-1]
        occurrence_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if corrected:
            _set_image_alt(
                doc,
                alt_text="Organigramme de la Direction des affaires juridiques "
                "(description ci-dessous).",
                title="Organigramme du service",
            )
            description = doc.add_paragraph(
                "La Direction des affaires juridiques comprend 4 bureaux : "
                "le Bureau du droit public, le Bureau du droit social, "
                "le Bureau de la communication et le Bureau des affaires "
                "internationales. Chaque bureau est rattaché directement "
                "à la direction."
            )
            for run in description.runs:
                run.font.name = "Arial"
                run.font.size = Pt(10)
        else:
            _set_image_alt(doc, alt_text="image.png")

    if with_guidance and occurrence_paragraph is not None:
        _add_guidance_comment(doc, occurrence_paragraph.runs, _guidance_text(control))


def _add_station_two_p08(
    doc,
    control,
    numbering_id,
    icon_path,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute le pictogramme redondant à traiter comme décoratif."""
    _add_station_one_control(doc, control, numbering_id)
    occurrence_run = None

    if icon_path:
        paragraph = doc.add_paragraph()
        paragraph.add_run("Pour toute question, contactez-nous par ")
        image_run = paragraph.add_run()
        image_run.add_picture(str(icon_path), width=Inches(0.18))
        occurrence_run = image_run
        if corrected:
            _mark_image_decorative(doc)
        else:
            _set_image_alt(doc, alt_text="E-mail")
        paragraph.add_run(" e-mail pour plus d'informations.")

    if with_guidance and occurrence_run is not None:
        _add_guidance_comment(doc, occurrence_run, _guidance_text(control))


def _add_station_two_p09(
    doc,
    control,
    numbering_id,
    texte_image_path,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute la transformation d'une image de texte en texte réel."""
    _add_station_one_control(doc, control, numbering_id)
    occurrence_paragraph = None
    if corrected:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(
            "Avis important : les indicateurs du T2 2025 "
            "seront transmis avant le 15 septembre 2025."
        )
        run.font.name = "Arial"
        run.font.size = Pt(11)
    elif texte_image_path:
        doc.add_picture(str(texte_image_path), width=Inches(4.5))
        occurrence_paragraph = doc.paragraphs[-1]
        occurrence_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if with_guidance and occurrence_paragraph is not None:
        _add_guidance_comment(doc, occurrence_paragraph.runs, _guidance_text(control))


def _add_station_two_p10(
    doc,
    control,
    numbering_id,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute le téléchargement avec un libellé adapté à la variante."""
    _add_station_one_control(doc, control, numbering_id)
    paragraph = doc.add_paragraph("Pour accéder aux annexes, ")
    label = (
        "Consulter les annexes du rapport T1 2025 (PDF, 1,2 Mo, français)"
        if corrected
        else "cliquez ici"
    )
    _add_hyperlink(
        paragraph,
        label,
        "https://example.org/annexes-rapport-t1-2025.pdf",
    )
    paragraph.add_run(".")

    if with_guidance:
        _add_guidance_comment(doc, paragraph.runs, _guidance_text(control))


def _add_station_two_p11(
    doc,
    control,
    numbering_id,
    *,
    corrected=False,
    with_guidance=False,
    guidance_anchor=None,
):
    """Reprend dans le corps le statut porté par l'en-tête."""
    _add_station_one_control(doc, control, numbering_id)
    if corrected:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(STATUT_DOCUMENT)
        run.font.name = "Arial"
        run.font.size = Pt(11)

    if with_guidance and guidance_anchor is not None:
        _add_guidance_comment(
            doc,
            guidance_anchor.runs,
            _guidance_text(control, document_level=True),
        )


def _mark_first_row_as_header(table):
    """Déclare la première ligne comme en-tête répétable."""
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    for current in list(tr_pr.findall(qn("w:tblHeader"))):
        tr_pr.remove(current)
    tr_pr.append(parse_xml(f'<w:tblHeader {nsdecls("w")} w:val="true"/>'))


def _set_table_alt_text(table, title, description):
    """Renseigne le titre et la description du tableau (texte de remplacement)."""
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblCaption", "w:tblDescription"):
        for current in list(tbl_pr.findall(qn(tag))):
            tbl_pr.remove(current)
    for tag, value in (("w:tblCaption", title), ("w:tblDescription", description)):
        value = escape(value, {'"': "&quot;"})
        tbl_pr.append(parse_xml(f'<{tag} {nsdecls("w")} w:val="{value}"/>'))


def _set_table_column_widths(table, widths):
    """Fixe la largeur de chaque colonne, dans la grille et dans les cellules."""
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
        for cell in column.cells:
            cell.width = width


def _set_style_font_family(style, font_name):
    """Remplace les polices de thème d'un style par une police explicite."""
    rpr = style.element.get_or_add_rPr()
    for current in list(rpr.findall(qn("w:rFonts"))):
        rpr.remove(current)
    rpr.insert(
        0,
        parse_xml(
            f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" '
            f'w:eastAsia="{font_name}" w:cs="{font_name}"/>'
        ),
    )


def _prevent_table_row_splitting(table):
    """Interdit le fractionnement des lignes sur deux pages."""
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        for current in list(tr_pr.findall(qn("w:cantSplit"))):
            tr_pr.remove(current)
        tr_pr.append(parse_xml(f'<w:cantSplit {nsdecls("w")} w:val="true"/>'))


def _add_station_three_p12(
    doc,
    control,
    station_title,
    numbering_id,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute l'exemple de contraste textuel à mesurer."""
    station_heading = doc.add_heading(station_title, level=1)
    _apply_heading_numbering(station_heading, numbering_id, 0)
    _add_station_one_control(doc, control, numbering_id)

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(CONTRAST_SAMPLE_TEXT)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = (
        RGBColor(0x59, 0x59, 0x59) if corrected else RGBColor(0x9A, 0x9A, 0x9A)
    )
    if with_guidance:
        _add_guidance_comment(doc, run, _guidance_text(control))


def _add_station_three_p13(
    doc,
    control,
    numbering_id,
    chart_path,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute le graphique fautif ou corrigé de la station 3."""
    _add_station_one_control(doc, control, numbering_id)
    doc.add_heading("Détail par canal", level=3)
    doc.add_picture(str(chart_path), width=Inches(4.5))
    chart_paragraph = doc.paragraphs[-1]
    chart_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    alt_text = CHART_ALT_ACCESSIBLE if corrected else CHART_ALT_INACCESSIBLE
    _set_image_alt(doc, alt_text=alt_text, title="Trafic web T1 2025")
    if with_guidance:
        _add_guidance_comment(doc, chart_paragraph.runs, _guidance_text(control))


def _add_station_three_p14(
    doc,
    control,
    numbering_id,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute le tableau de données fautif ou corrigé de la station 3."""
    _add_station_one_control(doc, control, numbering_id)
    doc.add_heading("Répartition par service", level=3)
    doc.add_paragraph("Périmètre : Direction des affaires juridiques.")
    table = doc.add_table(rows=3, cols=3, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    data = [
        ["Service", "Effectif", "Budget"],
        ["Communication", "12", "45 000"],
        ["Juridique", "28", "120 000"],
    ]
    for row_index, row_data in enumerate(data):
        for column_index, cell_text in enumerate(row_data):
            cell = table.cell(row_index, column_index)
            cell.text = cell_text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
                    if row_index == 0:
                        run.bold = True

    if corrected:
        _mark_first_row_as_header(table)
        _prevent_table_row_splitting(table)
    else:
        merged_header = table.cell(0, 1).merge(table.cell(0, 2))
        merged_header.text = "Effectif\nBudget"
        for paragraph in merged_header.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(10)
        if with_guidance:
            _add_guidance_comment(
                doc,
                merged_header.paragraphs[0].runs,
                _guidance_text(control),
            )


def _add_station_four_p15(
    doc,
    control,
    station_title,
    numbering_id,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute le passage anglais de la station 4 et son balisage de langue."""
    station_heading = doc.add_heading(station_title, level=1)
    _apply_heading_numbering(station_heading, numbering_id, 0)
    _add_station_one_control(doc, control, numbering_id)

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(
        "The quarterly report is available upon request. "
        "Please contact the communication department for further details."
    )
    run.font.name = "Arial"
    run.font.size = Pt(11)
    if corrected:
        run_properties = run._r.get_or_add_rPr()
        run_properties.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="en-US"/>'))
    if with_guidance:
        _add_guidance_comment(doc, run, _guidance_text(control))


def _add_station_four_p16(
    doc,
    control,
    numbering_id,
    *,
    with_guidance=False,
):
    """Ajoute un paragraphe qui hérite du style Normal à corriger."""
    _add_station_one_control(doc, control, numbering_id)
    paragraph = doc.add_paragraph(
        "La mise en forme du corps du document est pilotée par le style Normal."
    )
    if with_guidance:
        _add_guidance_comment(doc, paragraph.runs, _guidance_text(control))


def _add_station_four_p17(
    doc,
    control,
    numbering_id,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute l'intertitre saisi ou mis en forme en majuscules."""
    _add_station_one_control(doc, control, numbering_id)
    if corrected:
        heading = doc.add_heading("Annexes - accessibilité", level=3)
        for run in heading.runs:
            run.font.all_caps = True
    else:
        heading = doc.add_heading("ANNEXES - ACCESSIBILITE", level=3)
    if with_guidance:
        _add_guidance_comment(doc, heading.runs, _guidance_text(control))


def _add_station_four_p18(
    doc,
    control,
    numbering_id,
    *,
    corrected=False,
    with_guidance=False,
):
    """Ajoute la première occurrence d'un acronyme à développer."""
    _add_station_one_control(doc, control, numbering_id)
    if corrected:
        text = (
            "Le Référentiel général d’amélioration de l’accessibilité (RGAA) "
            "structure le contrôle des documents numériques."
        )
    else:
        text = "Le RGAA structure le contrôle des documents numériques."
    paragraph = doc.add_paragraph(text)
    if with_guidance:
        _add_guidance_comment(doc, paragraph.runs, _guidance_text(control))


def _add_station_five_p19(
    doc,
    control,
    station_title,
    numbering_id,
    *,
    with_guidance=False,
):
    """Ajoute le contrôle des propriétés et son ancre documentaire stable."""
    station_heading = doc.add_heading(station_title, level=1)
    _apply_heading_numbering(station_heading, numbering_id, 0)
    _add_station_one_control(doc, control, numbering_id)
    anchor = doc.add_paragraph(
        f"Document {chr(0x2014)} propriétés : vérifiez le titre, l’auteur, la "
        "langue et le nom de votre copie de travail."
    )
    if with_guidance:
        _add_guidance_comment(
            doc,
            anchor.runs,
            _guidance_text(control, document_level=True),
        )


def _add_station_five_followups(doc, controls, numbering_id):
    """Ajoute les contrôles de finalisation sans défaut injecté."""
    for control in controls:
        _add_station_one_control(doc, control, numbering_id)


def _paginate_corrected_guide(doc, matrix):
    """Découpe le guide par station et par contrôle sans réduire le texte."""
    paragraphs = {paragraph.text: paragraph for paragraph in doc.paragraphs}
    for station in (
        block for block in matrix["sequence"] if block["id"].startswith("station-")
    ):
        station_heading = paragraphs[station["titre"]]
        station_heading.paragraph_format.page_break_before = True
        station_heading.paragraph_format.keep_with_next = True
        controls = [
            control
            for control in matrix["controles"]
            if control["station"] == station["id"]
        ]
        for index, control in enumerate(controls):
            heading = paragraphs[f"{control['id']} - {control['intitule']}"]
            heading.paragraph_format.keep_with_next = True
            if index:
                heading.paragraph_format.page_break_before = True


def _ensure_minimum_body_font_size(doc, minimum_points):
    """Relève les tailles directes du corps sans réduire les titres."""
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None and run.font.size.pt < minimum_points:
                run.font.size = Pt(minimum_points)


def build_inaccessible(
    chart_path: Path,
    icon_path: Path = None,
    organigramme_path: Path = None,
    texte_image_path: Path = None,
    with_guidance: bool = False,
    output_name: str | None = None,
    matrix=None,
    output_dir: Path | None = None,
    affiche_path: Path | None = None,
):
    matrix = matrix or load_sami_matrix()
    station_1_controls = _station_controls(matrix, "station-1")
    station_1_title = _station_title(matrix, "station-1")
    p01, p02, p03, p04, p05 = station_1_controls
    p06, p07, p08, p09, p10, p11 = _station_controls(matrix, "station-2")
    station_2_title = _station_title(matrix, "station-2")
    p12, p13, p14 = _station_controls(matrix, "station-3")
    station_3_title = _station_title(matrix, "station-3")
    p15, p16, p17, p18 = _station_controls(matrix, "station-4")
    station_4_title = _station_title(matrix, "station-4")
    p19, *station_5_followups = _station_controls(matrix, "station-5")
    station_5_title = _station_title(matrix, "station-5")
    doc = Document()
    station_numbering_id = _create_heading_numbering(doc)
    doc.core_properties.title = ""
    doc.core_properties.author = ""
    doc.core_properties.subject = ""
    doc.core_properties.language = "de-DE"

    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Arial"
    style_normal.font.size = Pt(11)
    # Défaut P-16 : texte justifié, interligne simple et corps inférieur à 12 pt.
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_normal.paragraph_format.line_spacing = 1.0
    _set_doc_defaults_language(doc, "de-DE")
    _set_content_styles_language(doc, "de-DE")

    # En-tête fictif
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = HEADER_TEXT
    hp.style.font.name = "Arial"
    hp.style.font.size = Pt(9)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _format_header_paragraph(hp)
    _add_header_status(header)
    _add_page_footer(doc)

    _add_station_one_p01(
        doc,
        p01,
        station_1_title,
        station_numbering_id,
        with_guidance=with_guidance,
    )
    _add_station_one_p02(doc, p02, station_numbering_id, with_guidance=with_guidance)

    doc.add_heading("Introduction", level=1)

    intro_paragraph = doc.add_paragraph(INTRO_TEXT)

    # Défaut P-03 : faux sommaire tapé à la main avec points de suite manuels.
    p = doc.add_paragraph()
    run = p.add_run("Sommaire")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x91)
    if with_guidance:
        _add_guidance_comment(
            doc,
            run,
            _guidance_text(p03, document_level=True),
        )
    for titre_som, page in [
        ("Résultats du trimestre", "2"),
        (station_2_title, "3"),
        ("Détail par canal", "4"),
        ("Annexes", "5"),
    ]:
        doc.add_paragraph(f"{titre_som} .............. {page}")

    _add_station_one_control(doc, p03, station_numbering_id)

    doc.add_paragraph()

    # La mention explicite ne repose pas sur la couleur seule et reste contrastée.
    p = doc.add_paragraph()
    run = p.add_run("Urgent - Retour attendu avant le 30 juin 2025")
    run.bold = True
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    run.font.name = "Arial"
    run.font.size = Pt(11)

    doc.add_paragraph("La direction demande un retour rapide sur les indicateurs.")

    doc.add_heading("Résultats du trimestre", level=2)

    # Tableau historique déjà structuré : la station 3 utilise un autre tableau cible.
    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    data = [
        ["Indicateur", "T4 2024", "T1 2025", "Évolution"],
        ["Visiteurs uniques", "45 200", "50 600", "+12 %"],
        ["Pages vues", "128 000", "142 000", "+11 %"],
        ["Taux de rebond", "42 %", "38 %", "-4 pts"],
    ]
    for i, row_data in enumerate(data):
        for j, cell_text in enumerate(row_data):
            cell = table.cell(i, j)
            cell.text = cell_text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
    _mark_first_row_as_header(table)
    _prevent_table_row_splitting(table)

    _add_station_one_control(doc, p04, station_numbering_id)

    # Défaut P-04 : fausse liste à puces saisie et indentée manuellement.
    doc.add_paragraph("Objectifs du trimestre :")
    first_fake_bullet = None
    for item in [
        "Augmenter le trafic de 10 %",
        "Publier 3 articles par semaine",
        "Réduire le taux de rebond sous 40 %",
    ]:
        p = _add_fake_list_item(doc, "\u2022", item)
        first_fake_bullet = first_fake_bullet or p
    if with_guidance and first_fake_bullet:
        _add_guidance_comment(
            doc,
            first_fake_bullet.runs,
            _guidance_text(p04, precision="liste à puces"),
        )

    # Défaut P-04 : fausse liste numérotée saisie et indentée manuellement.
    doc.add_paragraph("Priorités pour le prochain trimestre :")
    first_fake_number = None
    for numero, item in enumerate(
        [
            "Refonte de la page d'accueil",
            "Mise en conformité accessibilité",
            "Déploiement de la newsletter",
        ],
        start=1,
    ):
        p = _add_fake_list_item(doc, f"{numero}.", item)
        first_fake_number = first_fake_number or p
    if with_guidance and first_fake_number:
        _add_guidance_comment(
            doc,
            first_fake_number.runs,
            _guidance_text(p04, precision="liste numérotée"),
        )

    _add_station_one_p05(doc, p05, station_numbering_id, with_guidance=with_guidance)

    _add_station_two_p06(
        doc,
        p06,
        station_2_title,
        station_numbering_id,
        affiche_path,
        with_guidance=with_guidance,
    )
    _add_station_two_p07(
        doc,
        p07,
        station_numbering_id,
        organigramme_path,
        with_guidance=with_guidance,
    )
    _add_station_two_p08(
        doc,
        p08,
        station_numbering_id,
        icon_path,
        with_guidance=with_guidance,
    )
    _add_station_two_p09(
        doc,
        p09,
        station_numbering_id,
        texte_image_path,
        with_guidance=with_guidance,
    )
    _add_station_two_p10(
        doc,
        p10,
        station_numbering_id,
        with_guidance=with_guidance,
    )
    _add_station_two_p11(
        doc,
        p11,
        station_numbering_id,
        with_guidance=with_guidance,
        guidance_anchor=intro_paragraph,
    )

    _add_station_three_p12(
        doc,
        p12,
        station_3_title,
        station_numbering_id,
        with_guidance=with_guidance,
    )
    _add_station_three_p13(
        doc,
        p13,
        station_numbering_id,
        chart_path,
        with_guidance=with_guidance,
    )
    _add_station_three_p14(
        doc,
        p14,
        station_numbering_id,
        with_guidance=with_guidance,
    )

    _add_station_four_p15(
        doc,
        p15,
        station_4_title,
        station_numbering_id,
        with_guidance=with_guidance,
    )
    _add_station_four_p16(
        doc,
        p16,
        station_numbering_id,
        with_guidance=with_guidance,
    )
    _add_station_four_p17(
        doc,
        p17,
        station_numbering_id,
        with_guidance=with_guidance,
    )
    _add_station_four_p18(
        doc,
        p18,
        station_numbering_id,
        with_guidance=with_guidance,
    )
    _add_station_five_p19(
        doc,
        p19,
        station_5_title,
        station_numbering_id,
        with_guidance=with_guidance,
    )
    _add_station_five_followups(doc, station_5_followups, station_numbering_id)

    version_key = "avec_pistes" if with_guidance else "inaccessible"
    output = (output_dir or PROJECT / "_source") / (
        output_name or matrix["identite_editoriale"]["versions"][version_key]
    )
    doc.save(str(output))
    _remove_quarantine(output)
    print(f"  -> {output.name}")
    return output


# ------------------------------------------------------------------
# 3. Document accessible
# ------------------------------------------------------------------


def build_accessible(
    chart_path: Path,
    icon_path: Path = None,
    organigramme_path: Path = None,
    texte_image_path: Path = None,
    matrix=None,
    output_dir: Path | None = None,
    output_name: str | None = None,
    affiche_path: Path | None = None,
):
    matrix = matrix or load_sami_matrix()
    station_1_controls = _station_controls(matrix, "station-1")
    station_1_title = _station_title(matrix, "station-1")
    p01, p02, p03, p04, p05 = station_1_controls
    p06, p07, p08, p09, p10, p11 = _station_controls(matrix, "station-2")
    station_2_title = _station_title(matrix, "station-2")
    p12, p13, p14 = _station_controls(matrix, "station-3")
    station_3_title = _station_title(matrix, "station-3")
    p15, p16, p17, p18 = _station_controls(matrix, "station-4")
    station_4_title = _station_title(matrix, "station-4")
    p19, *station_5_followups = _station_controls(matrix, "station-5")
    station_5_title = _station_title(matrix, "station-5")
    doc = Document()
    station_numbering_id = _create_heading_numbering(doc)

    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Arial"
    style_normal.font.size = Pt(12)
    # Alignement à gauche, sans justification.
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_normal.paragraph_format.line_spacing = 1.15

    # Langue du document : fr-FR par défaut, passage anglais balisé plus bas.
    _set_doc_defaults_language(doc, "fr-FR")
    _set_content_styles_language(doc, "fr-FR")

    # Configurer les styles de titre
    for level, size in [(1, 16), (2, 14), (3, 12)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x00, 0x00, 0x91)

    # En-tête
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = HEADER_TEXT
    hp.style.font.name = "Arial"
    hp.style.font.size = Pt(12)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _format_header_paragraph(hp, font_size=12)
    _add_header_status(header, font_size=12)
    _add_page_footer(doc, font_size=12)

    _add_station_one_p01(
        doc, p01, station_1_title, station_numbering_id, corrected=True
    )
    _add_station_one_p02(doc, p02, station_numbering_id, corrected=True)

    # Titre 1
    doc.add_heading("Introduction", level=1)

    doc.add_paragraph(INTRO_TEXT)

    # Sommaire automatique (table des matieres generee depuis les styles)
    doc.add_heading("Sommaire", level=2)
    _add_toc(doc)
    _add_station_one_control(doc, p03, station_numbering_id)

    # Mention urgente accessible (#C00000 ratio 6.5:1 sur blanc)
    p = doc.add_paragraph()
    run = p.add_run("Urgent - Retour attendu avant le 30 juin 2025")
    run.bold = True
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    run.font.name = "Arial"
    run.font.size = Pt(11)

    doc.add_paragraph("La direction demande un retour rapide sur les indicateurs.")

    # Titre 2
    doc.add_heading("Résultats du trimestre", level=2)

    # Tableau avec en-tete balisee
    table = doc.add_table(rows=4, cols=4, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    data = [
        ["Indicateur", "T4 2024", "T1 2025", "Évolution"],
        ["Visiteurs uniques", "45 200", "50 600", "+12 %"],
        ["Pages vues", "128 000", "142 000", "+11 %"],
        ["Taux de rebond", "42 %", "38 %", "-4 pts"],
    ]
    for i, row_data in enumerate(data):
        for j, cell_text in enumerate(row_data):
            cell = table.cell(i, j)
            cell.text = cell_text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
            if i == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

    _mark_first_row_as_header(table)
    _prevent_table_row_splitting(table)

    _add_station_one_control(doc, p04, station_numbering_id)

    # Vraie liste a puces native
    doc.add_paragraph("Objectifs du trimestre :")
    for item in [
        "Augmenter le trafic de 10 %",
        "Publier 3 articles par semaine",
        "Réduire le taux de rebond sous 40 %",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    # Vraie liste numerotee native
    doc.add_paragraph("Priorités pour le prochain trimestre :")
    for item in [
        "Refonte de la page d'accueil",
        "Mise en conformité accessibilité",
        "Déploiement de la newsletter",
    ]:
        doc.add_paragraph(item, style="List Number")

    _add_station_one_p05(doc, p05, station_numbering_id, corrected=True)

    _add_station_two_p06(
        doc,
        p06,
        station_2_title,
        station_numbering_id,
        affiche_path,
        corrected=True,
    )
    _add_station_two_p07(
        doc,
        p07,
        station_numbering_id,
        organigramme_path,
        corrected=True,
    )
    _add_station_two_p08(
        doc,
        p08,
        station_numbering_id,
        icon_path,
        corrected=True,
    )
    _add_station_two_p09(
        doc,
        p09,
        station_numbering_id,
        texte_image_path,
        corrected=True,
    )
    _add_station_two_p10(
        doc,
        p10,
        station_numbering_id,
        corrected=True,
    )
    _add_station_two_p11(
        doc,
        p11,
        station_numbering_id,
        corrected=True,
    )

    _add_station_three_p12(
        doc,
        p12,
        station_3_title,
        station_numbering_id,
        corrected=True,
    )
    _add_station_three_p13(
        doc,
        p13,
        station_numbering_id,
        chart_path,
        corrected=True,
    )
    _add_station_three_p14(
        doc,
        p14,
        station_numbering_id,
        corrected=True,
    )

    _add_station_four_p15(
        doc,
        p15,
        station_4_title,
        station_numbering_id,
        corrected=True,
    )
    _add_station_four_p16(doc, p16, station_numbering_id)
    _add_station_four_p17(
        doc,
        p17,
        station_numbering_id,
        corrected=True,
    )
    _add_station_four_p18(
        doc,
        p18,
        station_numbering_id,
        corrected=True,
    )
    _add_station_five_p19(
        doc,
        p19,
        station_5_title,
        station_numbering_id,
    )
    _add_station_five_followups(doc, station_5_followups, station_numbering_id)

    # Proprietes du document
    doc.core_properties.title = DOCUMENT_TITLE
    doc.core_properties.author = "Sami Dupont"
    doc.core_properties.language = "fr-FR"
    doc.core_properties.subject = (
        "Guide pratique de mise en accessibilité d’un document Word"
    )
    doc.core_properties.keywords = (
        "accessibilité, document bureautique, Word, guide pratique, "
        "communication accessible"
    )
    _paginate_corrected_guide(doc, matrix)
    _ensure_minimum_body_font_size(doc, 12)

    output = (output_dir or PROJECT / "_source") / (
        output_name or matrix["identite_editoriale"]["versions"]["corrigee"]
    )
    doc.save(str(output))
    _remove_quarantine(output)
    print(f"  -> {output.name}")
    return output


def _remove_quarantine(path: Path):
    """Retire le flag com.apple.quarantine sur macOS."""
    try:
        subprocess.run(
            ["xattr", "-d", "com.apple.quarantine", str(path)],
            capture_output=True,
        )
    except FileNotFoundError:
        pass


def _checklist_groups(matrix):
    """Retourne les contrôles dans l'ordre des stations, puis les signalements."""
    groups = []
    for block in matrix["sequence"]:
        if not block["id"].startswith("station-"):
            continue
        controls = [
            control
            for control in matrix["controles"]
            if control["station"] == block["id"]
        ]
        groups.append((block["titre"], controls))
    signalled = [control for control in matrix["controles"] if control["niveau"] == "S"]
    groups.append(("Contrôles signalés", signalled))
    return groups


def _checklist_markdown(matrix, formation_code):
    lines = [
        "---",
        'title: "Checklist accessibilité des documents bureautiques"',
        'subtitle: "Suivi progressif du TP Word accessible"',
        f'author: "IGPDE - Formation {formation_code}"',
        "lang: fr",
        "---",
        "",
        "<!-- Généré depuis `_source/exercice-sami-matrice.yml` par `make checklist`. Ne pas modifier directement. -->",
        "",
        # Style local : mêmes police et pagination que la checklist DOCX.
        "<style>",
        "li { break-inside: avoid; }",
        ":root { font-family: Arial, sans-serif; }",
        "body { hyphens: manual; }",
        "@page { font-family: Arial, sans-serif;",
        "  @top-center { font-family: Arial, sans-serif; }",
        '  @bottom-center { content: "Page " counter(page) " sur " counter(pages);',
        "    font-family: Arial, sans-serif; } }",
        "</style>",
        "",
        "# Checklist accessibilité des documents bureautiques",
        "",
        "Utilisez cette même checklist dès le début du TP, puis complétez-la après chaque station.",
        "",
        "- **P - pratiqué :** une action est réalisée et sa preuve est conservée.",
        "- **C - contrôlé :** un outil ou une vérification humaine est exécuté et son résultat est noté.",
        "- **S - signalé :** le point est vérifié dans la checklist, sans manipulation obligatoire pendant le TP.",
        "",
    ]
    for title, controls in _checklist_groups(matrix):
        lines.extend((f"## {title}", ""))
        if controls and controls[0]["niveau"] == "S":
            lines.extend(
                (
                    "Ces points sont à vérifier avant diffusion, mais ne constituent pas des manipulations obligatoires du TP.",
                    "",
                )
            )
        for control in controls:
            lines.extend(
                (
                    f"- **{control['id']} · {control['niveau']}** - {control['checklist']}",
                    "  - Suivi : ☐ À vérifier · ☐ Fait · ☐ À reprendre",
                    "  - Notes :",
                    "",
                )
            )
    return "\n".join(lines).rstrip() + "\n"


def _checklist_docx(matrix):
    doc = Document()
    doc.core_properties.title = "Checklist accessibilité des documents bureautiques"
    doc.core_properties.author = "IGPDE"
    doc.core_properties.subject = "Suivi du TP Word accessible"
    doc.core_properties.language = "fr-FR"

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_after = Pt(6)
    _set_doc_defaults_language(doc, "fr-FR")
    _set_content_styles_language(doc, "fr-FR")
    for style_name in ("Title", "Heading 1"):
        _set_style_font_family(doc.styles[style_name], "Arial")

    doc.add_heading("Checklist accessibilité des documents bureautiques", level=0)
    doc.add_paragraph(
        "Utilisez cette même checklist dès le début du TP, puis complétez-la après chaque station."
    )
    for definition in (
        "P - pratiqué : une action est réalisée et sa preuve est conservée.",
        "C - contrôlé : un outil ou une vérification humaine est exécuté et son résultat est noté.",
        "S - signalé : le point est vérifié sans manipulation obligatoire pendant le TP.",
    ):
        label, explanation = definition.split(" : ", 1)
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(f"{label} :").bold = True
        paragraph.add_run(f" {explanation}")

    for title, controls in _checklist_groups(matrix):
        doc.add_heading(title, level=1)
        if controls and controls[0]["niveau"] == "S":
            doc.add_paragraph(
                "Ces points sont à vérifier avant diffusion, mais ne constituent pas des manipulations obligatoires du TP."
            )
        table = doc.add_table(rows=1, cols=3, style="Table Grid")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.rows[0].cells[0].text = "Identifiant et niveau"
        table.rows[0].cells[1].text = "Point à vérifier"
        table.rows[0].cells[2].text = "Suivi"
        for cell in table.rows[0].cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True
        _mark_first_row_as_header(table)
        _set_table_alt_text(
            table,
            title,
            "Une ligne par point : identifiant et niveau, point à vérifier, "
            "cases de suivi et notes.",
        )
        for control in controls:
            cells = table.add_row().cells
            cells[0].text = f"{control['id']} · {control['niveau']}"
            cells[1].text = control["checklist"]
            cells[2].text = "☐ À vérifier\n☐ Fait\n☐ À reprendre\nNotes :"
        # Colonne des points à vérifier élargie : la largeur utile fait 6 pouces.
        _set_table_column_widths(table, (Inches(1.3), Inches(3.2), Inches(1.5)))
        _prevent_table_row_splitting(table)

    _add_page_footer(doc, document_name=None, font_size=12, separator=" sur ")
    return doc


def build_checklists(
    *,
    matrix=None,
    markdown_output: Path,
    docx_output: Path,
    formation_code: str | None = None,
):
    """Génère les deux checklists depuis la matrice canonique."""
    matrix = matrix or load_sami_matrix()
    formation_code = formation_code or load_formation_config()["code"]
    markdown_output = Path(markdown_output)
    docx_output = Path(docx_output)
    markdown_output.parent.mkdir(parents=True, exist_ok=True)
    docx_output.parent.mkdir(parents=True, exist_ok=True)
    markdown_output.write_text(
        _checklist_markdown(matrix, formation_code), encoding="utf-8"
    )
    _checklist_docx(matrix).save(docx_output)
    _remove_quarantine(docx_output)
    return markdown_output, docx_output


def build_default_checklists():
    """Génère les checklists dans leurs emplacements contractuels."""
    config = load_formation_config()
    docx_output = (
        PROJECT
        / config["livrables"]
        / "Livrables-Stagiaires"
        / "tp-word-igpde"
        / f"{CHECKLIST_BASENAME}.docx"
    )
    outputs = build_checklists(
        markdown_output=CHECKLIST_MARKDOWN,
        docx_output=docx_output,
    )
    for output in outputs:
        print(f"  -> {output.relative_to(PROJECT)}")
    return outputs


# ------------------------------------------------------------------
# Main
# ------------------------------------------------------------------


def generate_sami_documents():
    """Génère les trois DOCX et leurs ressources graphiques."""
    print("Génération des graphiques...")
    chart_bad = generate_chart_inaccessible()
    print(f"  -> {chart_bad.name}")
    chart_good = generate_chart_accessible()
    print(f"  -> {chart_good.name}")

    print("\nGénération des images supplémentaires...")
    icon = generate_icon_enveloppe()
    print(f"  -> {icon.name}")
    orga = generate_organigramme()
    print(f"  -> {orga.name}")
    txt_img = generate_texte_image()
    print(f"  -> {txt_img.name}")

    print("\nGénération des documents Word...")
    build_inaccessible(
        chart_bad,
        icon_path=icon,
        organigramme_path=orga,
        texte_image_path=txt_img,
        affiche_path=AFFICHE_P06,
    )
    build_inaccessible(
        chart_bad,
        icon_path=icon,
        organigramme_path=orga,
        texte_image_path=txt_img,
        with_guidance=True,
        affiche_path=AFFICHE_P06,
    )
    build_accessible(
        chart_good,
        icon_path=icon,
        organigramme_path=orga,
        texte_image_path=txt_img,
        affiche_path=AFFICHE_P06,
    )

    print("\nTerminé.")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--checklist",
        action="store_true",
        help="génère le DOCX et la source Markdown de la checklist",
    )
    args = parser.parse_args(argv)
    if args.checklist:
        print("Génération des checklists...")
        build_default_checklists()
        print("\nTerminé.")
    else:
        generate_sami_documents()


if __name__ == "__main__":
    main()
