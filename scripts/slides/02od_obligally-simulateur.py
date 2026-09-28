"""Slide 02od : activité Obligally - découvrir ses obligations légales."""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from igpde_dsfr_components import (
    CONTENT_W, MARGIN_L,
    Stack,
    add_alert, add_highlight, add_image, add_notes, add_texte_libre,
    estimate_alert_height, estimate_highlight_height,
    new_slide,
)

URL_OBLIGALLY = "https://obligations-legales-accessibilite-numerique.fr/fr/"
URL_SIMULATION = "https://obligations-legales-accessibilite-numerique.fr/fr/simulation/"
URL_COMPRENDRE = "https://obligations-legales-accessibilite-numerique.fr/fr/comprendre/"


def _add_link_line(slide, label, url, top):
    txbox = slide.shapes.add_textbox(
        Inches(MARGIN_L), Inches(top),
        Inches(CONTENT_W), Inches(0.325),
    )
    tf = txbox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]

    label_run = p.add_run()
    label_run.text = label
    label_run.font.size = Pt(14)
    label_run.font.bold = True
    label_run.font.color.rgb = RGBColor(0x16, 0x16, 0x16)

    link_run = p.add_run()
    link_run.text = url
    link_run.font.size = Pt(14)
    link_run.font.color.rgb = RGBColor(0x00, 0x00, 0x91)
    link_run.font.underline = True
    link_run.hyperlink.address = url


def build(prs, layouts, ctx):
    slide = new_slide(
        prs, layouts,
        layout_name="titre_contenu",
        titre="Quelles obligations pour votre structure ?",
        fil_ariane="1. Q3 - Cadre légal | Obligations",
        footer_text=f"{ctx.footer_base} / Module 1",
        date_text=ctx.date,
        page_num=ctx.page_num,
    )

    stack = Stack(top=2.30, gap=0.35)

    accroche = (
        "Vous connaissez le cadre. "
        "Mais concrètement, quelles obligations "
        "s'appliquent à votre poste et à votre structure ?"
    )
    add_highlight(
        slide, accroche,
        top=stack.push(estimate_highlight_height(accroche, CONTENT_W)),
    )

    alert_titre = "Activité - simulateur Obligally (10 min)"
    alert_bullets = [
        "Cliquez sur Simuler et répondez aux questions",
        "Notez le résultat : quelles normes s'appliquent à vous ?",
        "Visualiser (infographie) : .../fr/visualisation/",
        "Approfondir (article détaillé) : .../fr/comprendre/",
    ]
    alert_w = CONTENT_W - 2.0
    alert_h = estimate_alert_height(
        alert_titre, alert_bullets, alert_w, line_spacing=1.15,
    )
    add_alert(
        slide,
        alert_titre,
        alert_bullets,
        top=stack.push(alert_h),
        left=MARGIN_L,
        width=alert_w,
        alert_type="info",
    )

    qr_w = 1.8
    qr_left = MARGIN_L + alert_w + 0.15
    qr_top = stack.cursor - alert_h
    add_image(
        slide,
        "scripts/images/qr-obligally.png",
        top=qr_top,
        left=qr_left,
        width=qr_w,
        alt_text=(
            "QR code vers le simulateur Obligally : "
            + URL_OBLIGALLY
        ),
    )

    add_texte_libre(
        slide,
        "Scannez-moi !",
        top=qr_top + qr_w + 0.05,
        left=qr_left, width=qr_w, height=0.35,
        size=14, bold=True,
    )

    txbox = slide.shapes.add_textbox(
        Inches(MARGIN_L), Inches(stack.cursor + 0.05),
        Inches(CONTENT_W), Inches(0.35),
    )
    tf = txbox.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = URL_OBLIGALLY
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x91)
    run.font.underline = True
    run.hyperlink.address = URL_OBLIGALLY

    _add_link_line(
        slide,
        "Simuler : ",
        URL_SIMULATION,
        top=stack.cursor + 0.40,
    )
    _add_link_line(
        slide,
        "Comprendre : ",
        URL_COMPRENDRE,
        top=stack.cursor + 0.73,
    )

    add_notes(
        slide,
        "Activité individuelle 10 min + débrief collectif 5 min. "
        "Faire ressortir : la plupart des structures publiques sont "
        "soumises au RGAA ; les communicants produisent des contenus "
        "concernés (PDF, vidéos, réseaux sociaux). "
        "Fallback si pas de réseau : distribuer le PDF de l'arbre "
        "de décision Idéance (dans _source/références/).",
    )
    return slide
