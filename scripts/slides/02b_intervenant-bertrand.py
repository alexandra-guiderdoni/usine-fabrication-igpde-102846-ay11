"""Slide 01b : Intervenant - Bertrand Matge."""

from pptx.util import Pt
from pptx.oxml.ns import qn

from igpde_dsfr_components import (
    add_callout, add_image, add_notes, new_slide,
    MARGIN_L, CONTENT_W, BLEU_FRANCE,
)

PHOTO_W = 3.5
BIO_LEFT = round(MARGIN_L + PHOTO_W + 0.28, 2)
BIO_W = round(CONTENT_W - PHOTO_W - 0.28, 2)
EMAIL = "bertrand.matge@finances.gouv.fr"


def _add_mailto(slide, email):
    """Ajoute un hyperlien mailto sur le dernier bullet du callout-body."""
    for shape in reversed(list(slide.shapes)):
        if shape.name == "DSFR-callout-body":
            for para in shape.text_frame.paragraphs:
                if email in para.text:
                    for run in para.runs:
                        if email in run.text:
                            rPr = run._r.get_or_add_rPr()
                            hlinkClick = rPr.makeelement(qn("a:hlinkClick"), {})
                            rel = slide.part.relate_to(
                                f"mailto:{email}", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                is_external=True,
                            )
                            hlinkClick.set(qn("r:id"), rel)
                            rPr.append(hlinkClick)
                            run.font.color.rgb = BLEU_FRANCE
                            run.font.underline = True
                    return
            return


def build(prs, layouts, ctx):
    slide = new_slide(
        prs, layouts,
        layout_name="titre_contenu",
        titre="Bertrand Matge",
        fil_ariane="1. Introduction | Intervenants",
        footer_text=f"{ctx.footer_base} / Intervenants",
        date_text=ctx.date,
        page_num=ctx.page_num,
    )

    add_image(
        slide,
        "_assets/photo-bertrand.jpg",
        top=2.40, left=0.71, width=2.39,
        alt_text="Photo Bertrand Matge",
    )

    add_callout(
        slide,
        "Responsable pôle support web - Mission Ingénierie du Web, SG-SNUM",
        [
            "Formateur Opquast certifié, référent Assurance Qualité Web",
            "Forme et sensibilise à l'accessibilité depuis plusieurs années",
            "Passionné par la conception inclusive et l'expérience utilisateur",
            "Formé à l'audit d'accessibilité numérique",
            f"Email : {EMAIL}",
        ],
        top=2.4,
        left=BIO_LEFT,
        width=BIO_W,
        line_spacing=1.5,
    )

    _add_mailto(slide, EMAIL)

    add_notes(slide, "Présentation de Bertrand : parcours SIRCOM, Opquast, sensibilisation accessibilité.")
    return slide
