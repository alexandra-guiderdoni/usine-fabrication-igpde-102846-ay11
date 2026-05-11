"""Slide 02t : synthèse Module 1 - 3 points clés + plan d'action."""

from copy import deepcopy
from pptx.util import Pt

from igpde_dsfr_components import (
    COL_R, COL_W, MARGIN_L, Stack, FONT, NOIR,
    add_alert, add_encadre, add_notes, new_slide,
)


def _bold_substring(slide, shape_name, text, substring):
    """Met en gras un sous-texte dans un run existant."""
    for shape in slide.shapes:
        if shape.name != shape_name or not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if substring in run.text:
                    before, _, after = run.text.partition(substring)
                    run.text = before
                    bold_run = deepcopy(run._r)
                    bold_run.text = substring
                    bold_run.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}rPr').set('b', '1')
                    run._r.addnext(bold_run)
                    if after:
                        after_run = deepcopy(run._r)
                        after_run.text = after
                        bold_run.addnext(after_run)
                    return


def build(prs, layouts, ctx):
    slide = new_slide(
        prs, layouts,
        layout_name="titre_contenu",
        titre="Module 1 - ce que vous retenez",
        fil_ariane="1. Introduction | Points clés",
        footer_text=f"{ctx.footer_base} / Module 1",
        date_text=ctx.date,
        page_num=ctx.page_num,
    )

    stack = Stack(top=2.30, gap=0.35)

    add_alert(
        slide,
        "3 points à retenir",
        [
            "L'accessibilité relève de l'équité et du droit, pas de la bonne volonté",
            "Les WCAG et le RGAA donnent un cadre pour vérifier ce qui est conforme",
            "Les premiers réflexes : texte lisible, contraste testé, alternative disponible",
        ],
        top=stack.push(1.7),
        alert_type="info",
    )

    _bold_substring(slide, "DSFR-alert-info-body",
                    "• L'accessibilité relève de l'équité", "équité")

    encadre_top = stack.cursor
    encadre_h = 6.75 - encadre_top

    add_encadre(
        slide,
        top=encadre_top,
        left=MARGIN_L,
        width=COL_W,
        height=encadre_h,
        titre="Dès demain matin",
        bullets=[
            "Relire une communication récente",
            "Vérifier texte et contraste",
            "Chercher la déclaration d'accessibilité du site",
        ],
    )

    add_encadre(
        slide,
        top=encadre_top,
        left=COL_R,
        width=COL_W,
        height=encadre_h,
        titre="Cette semaine",
        bullets=[
            "Partager les constats avec l'équipe",
            "Choisir 3 règles transversales à appliquer à chaque publication",
        ],
    )

    add_notes(
        slide,
        "Récupération active : demander à 2-3 stagiaires de citer un point retenu. "
        "Faire le lien avec la suite : on part de ce cadrage général, puis on descend dans les gestes concrets "
        "sur les documents bureautiques, le web et les réseaux sociaux.",
    )
    return slide
