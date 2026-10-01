"""Slide 02qcc : exemple concret - rapport DIA en PDF accessible et version FALC."""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from igpde_dsfr_components import (
    CONTENT_W, COL_W, COL_R, MARGIN_L,
    Stack,
    add_callout, add_highlight, add_image, add_notes,
    estimate_callout_height, estimate_highlight_height,
    new_slide,
)

URL_PDF = "https://handicap.gouv.fr/sites/handicap/files/2025-03/delegation-interministerielle-accessibilite-rapport-activite-2024.pdf"
URL_FALC = "https://handicap.gouv.fr/sites/handicap/files/2025-03/delegation-interministerielle-accessibilite-synthese-FLAC-rapport-activite-2024.pdf"

IMG_W = 1.2


def _inject_links_in_callouts(slide, urls):
    """Injecte des hyperliens sur le dernier bullet de chaque callout body."""
    idx = 0
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.name != "DSFR-callout-body":
            continue
        tf = shape.text_frame
        last_para = tf.paragraphs[-1]
        if "handicap.gouv.fr" in last_para.text and idx < len(urls):
            for run in last_para.runs:
                if "handicap.gouv.fr" in run.text:
                    run.font.color.rgb = RGBColor(0x00, 0x00, 0x91)
                    run.font.underline = True
                    run.hyperlink.address = urls[idx]
                    break
            idx += 1


def build(prs, layouts, ctx):
    slide = new_slide(
        prs, layouts,
        layout_name="titre_contenu",
        titre="Exemple : le rapport DIA en version accessible et FALC",
        fil_ariane="1. Q5 - Comment | FALC exemple",
        footer_text=f"{ctx.footer_base} / Module 1",
        date_text=ctx.date,
        page_num=ctx.page_num,
    )

    stack = Stack(top=1.90, gap=0.15)

    message = (
        "Le rapport d'activité 2024 de la Délégation interministérielle "
        "à l'accessibilité existe en PDF accessible et en version FALC."
    )
    add_highlight(
        slide, message,
        top=stack.push(estimate_highlight_height(message, CONTENT_W)),
    )

    img_top = stack.push(IMG_W + 0.05)

    col_center_l = MARGIN_L + (COL_W - IMG_W) / 2
    col_center_r = COL_R + (COL_W - IMG_W) / 2

    add_image(
        slide,
        "scripts/images/logo-falc-europe.jpg",
        top=img_top, left=col_center_l, width=IMG_W, height=IMG_W,
        alt_text="Logo europeen FALC - personnage lisant un document avec pouce leve",
    )
    add_image(
        slide,
        "scripts/images/rapport-dia-2024.jpg",
        top=img_top, left=col_center_r, width=IMG_W, height=IMG_W,
        alt_text="Couverture du rapport d'activité 2024 de la DIA en version FALC",
    )

    titre_pdf = "PDF accessible"
    bullets_pdf = [
        "Structure balisée (titres, listes, tableaux)",
        "Alt text sur les images, ordre de lecture logique",
        "Télécharger (handicap.gouv.fr)",
    ]
    titre_falc = "Version FALC"
    bullets_falc = [
        "Phrases courtes, vocabulaire simple, pictogrammes",
        "Validé par des personnes concernées",
        "Télécharger (handicap.gouv.fr)",
    ]
    col_h = max(
        estimate_callout_height(titre_pdf, bullets_pdf, COL_W, line_spacing=1.15, compact=True),
        estimate_callout_height(titre_falc, bullets_falc, COL_W, line_spacing=1.15, compact=True),
    )
    top_cols = stack.push(col_h)
    add_callout(
        slide, titre_pdf, bullets_pdf,
        top=top_cols, left=MARGIN_L, width=COL_W,
        line_spacing=1.15,
        compact=True,
    )
    add_callout(
        slide, titre_falc, bullets_falc,
        top=top_cols, left=COL_R, width=COL_W,
        line_spacing=1.15,
        compact=True,
    )

    _inject_links_in_callouts(slide, [URL_PDF, URL_FALC])

    add_notes(
        slide,
        "Montrer les deux documents si possible. Le rapport de la DIA est "
        "un bon exemple car il vient d'une institution publique et montre "
        "que le FALC n'est pas reserve aux associations. Le PDF accessible "
        "et la version FALC sont complementaires : le premier est le document "
        "officiel rendu navigable, le second est une version simplifiee pour "
        "un public plus large.\n"
        "PDF accessible : " + URL_PDF + "\n"
        "Version FALC : " + URL_FALC,
    )
    return slide
