"""Slide 02nab : comprendre pour mieux agir - simulations Atalan."""

from igpde_dsfr_components import (
    CONTENT_W, GAP, MARGIN_L,
    Stack,
    add_callout, add_highlight, add_image, add_notes, add_texte_libre,
    estimate_highlight_height, estimate_callout_height,
    new_slide, _set_run_hyperlink,
)
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN


URL_ATALAN = "https://atalan.fr/agissons/fr/index.html"
IMG_PATH = "scripts/images/atalan-agissons-home.png"

COL_LEFT_W = (CONTENT_W - GAP) / 2
COL_RIGHT_LEFT = MARGIN_L + COL_LEFT_W + GAP
COL_RIGHT_W = COL_LEFT_W


def build(prs, layouts, ctx):
    slide = new_slide(
        prs, layouts,
        layout_name="titre_contenu",
        titre="Comprendre pour mieux agir",
        fil_ariane="1. Q2 - Pour qui | Comprendre",
        footer_text=f"{ctx.footer_base} / Module 1",
        date_text=ctx.date,
        page_num=ctx.page_num,
    )

    stack = Stack(top=2.30, gap=0.25)

    message = (
        "Il ne s’agit pas de se mettre à la place "
        "d’une personne handicapée, mais d’observer "
        "ce que nos choix numériques peuvent faciliter ou bloquer."
    )
    h_highlight = estimate_highlight_height(message, COL_LEFT_W)
    add_highlight(
        slide, message,
        top=stack.push(h_highlight),
        left=MARGIN_L, width=COL_LEFT_W,
    )

    titre_sim = "Testez les simulations suivantes"
    simulations = [
        "Daltonisme",
        "Malvoyance",
        "Cécité",
        "Surdité",
        "Handicap moteur",
    ]
    h_callout = estimate_callout_height(titre_sim, simulations, COL_LEFT_W, compact=True)
    add_callout(
        slide,
        titre_sim,
        simulations,
        top=stack.push(h_callout),
        left=MARGIN_L, width=COL_LEFT_W,
        compact=True,
    )

    img_top = 2.30
    img_w = COL_RIGHT_W
    img_h = 3.30
    add_image(
        slide, IMG_PATH,
        top=img_top, left=COL_RIGHT_LEFT, width=img_w, height=img_h,
        alt_text="Page d'accueil du site Atalan - L'accessibilité numérique, et si nous agissions ?",
    )

    url_top = img_top + img_h + 0.20
    box = slide.shapes.add_textbox(
        Inches(COL_RIGHT_LEFT), Inches(url_top),
        Inches(COL_RIGHT_W), Inches(0.50),
    )
    box.name = "DSFR-texte-libre"
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "L'accessibilité numérique, et si nous agissions ?"
    run.font.size = Pt(14)
    run.font.bold = True
    _set_run_hyperlink(slide, run, URL_ATALAN)

    add_notes(
        slide,
        "Daltonisme : environ 8 % des hommes et 0,4 % des femmes. "
        "Laisser les stagiaires naviguer librement sur le site Atalan "
        "pendant 5 minutes. Chaque simulation illustre un type de "
        "handicap et montre ce que nos choix de mise en forme "
        "peuvent faciliter ou bloquer.",
    )
    return slide
