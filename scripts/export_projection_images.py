#!/usr/bin/env python3
"""Exporte une section du deck canonique en PPTX de projection image."""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from igpde_dsfr_components import finalize_pptx


def _arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Rend une plage de slides en images pleine page et conserve "
            "la transcription ainsi que les notes du présentateur."
        )
    )
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--first-slide", required=True, type=int)
    parser.add_argument("--last-slide", required=True, type=int)
    parser.add_argument(
        "--images-dir",
        type=Path,
        help="Dossier de pages PNG déjà rendues, nommé slide-NNN.png",
    )
    parser.add_argument("--width", type=int, default=1672)
    parser.add_argument("--height", type=int, default=941)
    return parser.parse_args()


def _shape_text(shape) -> list[str]:
    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        return [text] if text else []
    if getattr(shape, "has_table", False):
        values = [
            cell.text.strip()
            for row in shape.table.rows
            for cell in row.cells
            if cell.text.strip()
        ]
        return values
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        values: list[str] = []
        for child in shape.shapes:
            values.extend(_shape_text(child))
        return values
    return []


def _slide_transcription(slide) -> str:
    values: list[str] = []
    for shape in slide.shapes:
        values.extend(_shape_text(shape))
    return "\n\n".join(values)


def _slide_title(slide, transcription: str, number: int) -> str:
    title_shape = slide.shapes.title
    if title_shape is not None and title_shape.text.strip():
        return title_shape.text.strip()
    for line in transcription.splitlines():
        if line.strip():
            return line.strip()
    return f"Slide {number}"


def _render_slides(
    source: Path,
    destination: Path,
    first_slide: int,
    last_slide: int,
    width: int,
    height: int,
) -> list[Path]:
    soffice = shutil.which("soffice")
    pdftoppm = shutil.which("pdftoppm")
    if not soffice or not pdftoppm:
        raise RuntimeError("soffice et pdftoppm sont requis")

    subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(destination),
            str(source),
        ],
        check=True,
    )
    pdf = destination / f"{source.stem}.pdf"
    if not pdf.exists():
        raise RuntimeError(f"PDF intermédiaire absent : {pdf}")

    prefix = destination / "slide"
    subprocess.run(
        [
            pdftoppm,
            "-png",
            "-f",
            str(first_slide),
            "-l",
            str(last_slide),
            "-scale-to-x",
            str(width),
            "-scale-to-y",
            str(height),
            str(pdf),
            str(prefix),
        ],
        check=True,
    )

    def page_number(path: Path) -> int:
        match = re.search(r"-(\d+)$", path.stem)
        if not match:
            raise RuntimeError(f"Numéro de page introuvable : {path}")
        return int(match.group(1))

    images = sorted(destination.glob("slide-*.png"), key=page_number)
    expected = last_slide - first_slide + 1
    if len(images) != expected:
        raise RuntimeError(f"{expected} images attendues, {len(images)} produites")
    return images


def _existing_images(
    directory: Path, first_slide: int, last_slide: int
) -> list[Path]:
    images = [
        directory / f"slide-{number:03d}.png"
        for number in range(first_slide, last_slide + 1)
    ]
    missing = [str(path) for path in images if not path.exists()]
    if missing:
        raise FileNotFoundError("Images absentes : " + ", ".join(missing))
    return images


def export_projection(args: argparse.Namespace) -> Path:
    source_path = args.source.resolve()
    template_path = args.template.resolve()
    output_path = args.output.resolve()
    if not source_path.exists():
        raise FileNotFoundError(source_path)
    if not template_path.exists():
        raise FileNotFoundError(template_path)
    if args.first_slide < 1 or args.last_slide < args.first_slide:
        raise ValueError("Plage de slides invalide")

    source = Presentation(source_path)
    target = Presentation(template_path)
    expected = args.last_slide - args.first_slide + 1
    if len(source.slides) < args.last_slide:
        raise ValueError("Le deck source ne contient pas la dernière slide demandée")
    if len(target.slides) != expected:
        raise ValueError(
            f"Le gabarit doit contenir {expected} slides, pas {len(target.slides)}"
        )
    width_gap = abs(source.slide_width - target.slide_width)
    height_gap = abs(source.slide_height - target.slide_height)
    if width_gap > 1_000 or height_gap > 1_000:
        raise ValueError("Le format du gabarit diffère de celui du deck source")
    target.slide_width = source.slide_width
    target.slide_height = source.slide_height

    with tempfile.TemporaryDirectory(prefix="projection-images-") as temporary:
        temporary_path = Path(temporary)
        if args.images_dir:
            images = _existing_images(
                args.images_dir.resolve(), args.first_slide, args.last_slide
            )
        else:
            images = _render_slides(
                source_path,
                temporary_path,
                args.first_slide,
                args.last_slide,
                args.width,
                args.height,
            )

        for offset, image in enumerate(images):
            source_number = args.first_slide + offset
            source_slide = source.slides[source_number - 1]
            target_slide = target.slides[offset]
            transcription = _slide_transcription(source_slide)
            title = _slide_title(source_slide, transcription, source_number)
            oral_notes = source_slide.notes_slide.notes_text_frame.text.strip()
            notes = (
                f"Lire la transcription\n\n{transcription}\n\n"
                f"Lire le discours oral\n\n{oral_notes}"
            )

            for shape in list(target_slide.shapes):
                element = shape._element
                element.getparent().remove(element)
            picture = target_slide.shapes.add_picture(
                str(image),
                0,
                0,
                width=target.slide_width,
                height=target.slide_height,
            )
            properties = picture._element.nvPicPr.cNvPr
            properties.set("name", f"{offset + 1:02d} - {title}")
            properties.set("title", title)
            properties.set(
                "descr",
                "Reproduction visuelle de la slide. La transcription complète "
                "et les notes du présentateur figurent dans les notes.",
            )
            target_slide.notes_slide.notes_text_frame.text = notes

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}-",
            suffix=output_path.suffix,
            dir=output_path.parent,
            delete=False,
        ) as temporary_output:
            candidate = Path(temporary_output.name)
        try:
            finalize_pptx(
                target,
                candidate,
                title="Partie II - Documents bureautiques accessibles - TP",
                subject="Support de projection image avec notes du présentateur",
            )
            candidate.replace(output_path)
        finally:
            candidate.unlink(missing_ok=True)
    return output_path


def main() -> None:
    output = export_projection(_arguments())
    print(output)


if __name__ == "__main__":
    main()
