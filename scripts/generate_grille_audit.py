"""Génère le carnet de diagnostic guidé des 13 points de contrôle rapides.

Sortie : 03-easy-checks/grille-audit-easy-checks.xlsx

Le classeur contient 15 onglets :
- Démarrer : consignes, informations communes et accès aux 13 exercices ;
- 13 fiches verticales : une page du site et un point de contrôle par onglet ;
- Synthèse : consolidation automatique des constats, corrections et preuves.

Usage :
    python3 scripts/generate_grille_audit.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.page import PageMargins

from config import load_formation_config


FORMATION = load_formation_config()

BLEU_FRANCE = "000091"
BLEU_CLAIR = "E3E3FD"
GRIS_MOYEN = "E5E5E5"
VERT_SUCCES = "18753C"
JAUNE_EDITABLE = "FFF4CC"
BLANC = "FFFFFF"

SIZE_TITLE = 20
SIZE_SECTION = 16
SIZE_HEADER = 14
SIZE_CELL = 14

BASE_URL = FORMATION["site_url"].rstrip("/")
WORKBOOK_TITLE = "Carnet de diagnostic guidé - 13 points de contrôle rapides"

CHECKS = [
    {
        "id": 1,
        "onglet": "#01 - Images",
        "titre": "Texte alternatif des images",
        "wcag": "1.1.1 - Contenu non textuel",
        "rgaa": "1.1, 1.2, 1.3, 1.6, 1.7, 1.8, 1.9, 6.1",
        "question": "Chaque image reçoit-elle le traitement adapté à son rôle : informative, décorative, fonctionnelle ou complexe ?",
        "methode": (
            "Afficher les attributs alt avec Web Developer "
            "(Images > Display Alt Attributes) ou utiliser le bookmarklet « Check images », "
            "puis vérifier leur pertinence selon le rôle de l'image et le contexte du lien."
        ),
        "conformite": "Alternative utile pour une image informative ou fonctionnelle. Alternative vide pour une image décorative. Description détaillée pour un contenu complexe.",
    },
    {
        "id": 2,
        "onglet": "#02 - Titre de page",
        "titre": "Vérifier le titre de la page",
        "wcag": "2.4.2 - Titre de page",
        "rgaa": "8.5, 8.6",
        "question": "Le titre permet-il d'identifier immédiatement la page et son état, avec l'information spécifique en premier ?",
        "methode": (
            "Lire le titre dans l'onglet du navigateur. Pour afficher sa valeur complète, "
            "utiliser le bookmarklet « 02 Titre de page ». En alternative, afficher le code "
            "source, rechercher <title> et lire son contenu."
        ),
        "conformite": (
            "Titre présent, pertinent et distinct. La requête et la pagination utiles sont "
            "indiquées. L'information propre à la page apparaît avant le nom du site."
        ),
    },
    {
        "id": 3,
        "onglet": "#03 - Titres",
        "titre": "Titres et hiérarchie",
        "wcag": "1.3.1 - Information et relations ; 2.4.6 - En-têtes et étiquettes",
        "rgaa": "9.1",
        "question": "Les titres visibles sont-ils réellement balisés et organisés sans saut de niveau incohérent ?",
        "methode": (
            "Comparer le plan visuel avec les titres identifiés par HeadingsMap, "
            "le bookmarklet « 9.1 - Lister les titres » ou WAVE."
        ),
        "conformite": "Titres balisés de h1 à h6. Un titre principal pertinent. Hiérarchie logique qui décrit la structure du contenu.",
    },
    {
        "id": 4,
        "onglet": "#04 - Contrastes",
        "titre": "Vérifier le contraste des couleurs",
        "wcag": "1.4.3 - Contraste minimum ; 1.4.11 - Contraste du contenu non textuel",
        "rgaa": "3.2, 3.3",
        "question": "Les textes, composants et informations graphiques atteignent-ils les contrastes minimaux requis ?",
        "methode": (
            "Repérer les contrastes à vérifier avec WCAG Contrast Checker pour Firefox "
            "ou WCAG Color contrast checker pour Chrome. Mesurer les couples de couleurs "
            "avec Colour Contrast Analyser (CCA) de Vispero ou Tanaguru Contrast-Finder. "
            "Attention : les extensions automatiques peuvent produire des faux positifs ; confirmer chaque "
            "résultat à partir des couleurs calculées, de la taille du texte et du rôle de l'élément."
        ),
        "conformite": "Texte courant : 4,5:1. Grand texte : 3:1. Composants et informations graphiques nécessaires : 3:1.",
    },
    {
        "id": 5,
        "onglet": "#05 - Évitement",
        "titre": "Lien d'évitement",
        "wcag": "2.4.1 - Contourner des blocs",
        "rgaa": "12.7",
        "question": "Au premier appui sur Tab, un lien visible permet-il d'atteindre directement le contenu principal ?",
        "methode": "Recharger la page puis appuyer une fois sur Tab.",
        "conformite": "Lien d'évitement présent, visible au focus, placé au début du parcours clavier et relié à une cible valide.",
    },
    {
        "id": 6,
        "onglet": "#06 - Clavier",
        "titre": "Focus et navigation clavier",
        "wcag": "2.1.1 - Clavier ; 2.1.2 - Pas de piège au clavier ; 2.4.3 - Parcours du focus ; 2.4.7 - Visibilité du focus",
        "rgaa": "7.3, 10.7, 12.8, 12.9",
        "question": "Le parcours complet fonctionne-t-il au clavier, avec un focus visible, un ordre logique et sans piège ?",
        "methode": "Cacher la souris puis utiliser Tab, Maj+Tab, Entrée, Espace et les flèches.",
        "conformite": "Tous les éléments interactifs sont atteignables et activables. Focus visible. Ordre logique. Aucun blocage ni perte de focus.",
    },
    {
        "id": 7,
        "onglet": "#07 - Langue",
        "titre": "Déclarer les langues utilisées dans la page",
        "wcag": "3.1.1 - Langue de la page ; 3.1.2 - Langue d'un passage",
        "rgaa": "8.3, 8.4, 8.7, 8.8",
        "question": "La langue principale et les changements de langue sont-ils correctement déclarés ?",
        "methode": (
            "Afficher les attributs lang et dir avec le bookmarklet « 08 Langue et direction ». "
            "En complément, ouvrir le validateur HTML du W3C depuis Web Developer "
            "(Outils > Valider le HTML), puis utiliser Tanaguru. "
            "Vérifier manuellement que chaque code de langue correspond au contenu."
        ),
        "conformite": "Langue principale déclarée avec un code valide. Changements de langue balisés lorsqu'ils modifient la prononciation.",
    },
    {
        "id": 8,
        "onglet": "#08 - Zoom 200 %",
        "titre": "Zoom à 200 %",
        "wcag": "1.4.4 - Redimensionnement du texte ; 1.4.10 - Redistribution",
        "rgaa": "10.4, 10.11",
        "question": "À 200 % et dans une fenêtre étroite, les contenus restent-ils lisibles et utilisables sans perte d'information ?",
        "methode": "Zoomer à 200 %, puis tester séparément une zone d'affichage de 320 CSS px et parcourir tout le contenu.",
        "conformite": "Aucun texte coupé ou superposé. Contenus et commandes utilisables. Pas de défilement dans deux directions pour lire le contenu courant.",
    },
    {
        "id": 9,
        "onglet": "#09 - Sous-titres",
        "titre": "Sous-titres vidéo",
        "wcag": "1.2.2 - Sous-titres pour un média pré-enregistré",
        "rgaa": "4.3, 4.4",
        "question": "La vidéo dispose-t-elle de sous-titres synchronisés et relus pour les paroles et les sons utiles ?",
        "methode": "Lancer la vidéo, couper le son et activer les sous-titres.",
        "conformite": "Sous-titres disponibles, synchronisés et relus. Paroles, identification utile des locuteurs et sons nécessaires à la compréhension.",
    },
    {
        "id": 10,
        "onglet": "#10 - Transcription",
        "titre": "Transcriptions audio et vidéo",
        "wcag": "1.2.1 - Contenu seulement audio ou vidéo pré-enregistré",
        "rgaa": "4.1, 4.2",
        "question": "Une transcription complète, structurée et proche du média permet-elle d'accéder à son contenu sans écouter l'audio ?",
        "methode": "Chercher une transcription affichée à proximité du lecteur ou un lien explicite qui permet de l'ouvrir.",
        "conformite": "Transcription complète et structurée, affichée près du média ou accessible par un lien explicite placé à proximité.",
    },
    {
        "id": 11,
        "onglet": "#11 - Audiodescription",
        "titre": "Audiodescription",
        "wcag": "1.2.3 - Audiodescription ou version de remplacement ; 1.2.5 - Audiodescription",
        "rgaa": "4.5, 4.6",
        "question": "Au niveau A, l'information visuelle essentielle est-elle disponible par l'audio principal, une audiodescription ou une alternative ; au niveau AA, une audiodescription synchronisée est-elle fournie lorsqu'elle est nécessaire ?",
        "methode": "Écouter d'abord sans regarder l'image, puis comparer avec la vidéo et ses alternatives.",
        "conformite": "Niveau A : audio principal, audiodescription ou alternative complète. Niveau AA : audiodescription synchronisée lorsqu'elle est nécessaire.",
    },
    {
        "id": 12,
        "onglet": "#12 - Étiquettes",
        "titre": "Étiquettes de formulaire",
        "wcag": "1.3.1 - Information et relations ; 2.5.3 - Étiquette dans le nom ; 3.3.2 - Étiquettes ou instructions",
        "rgaa": "11.1, 11.2, 11.3, 11.5, 11.6, 11.7",
        "question": "Chaque champ possède-t-il une étiquette visible, persistante et correctement associée, y compris pour les groupes ?",
        "methode": "Cliquer sur chaque libellé, vérifier le nom accessible et contrôler les légendes des groupes.",
        "conformite": "Étiquette visible et associée au champ. Nom accessible cohérent avec le texte visible. Légende pertinente pour chaque groupe.",
    },
    {
        "id": 13,
        "onglet": "#13 - Erreurs",
        "titre": "Champs obligatoires et erreurs",
        "wcag": "3.3.1 - Identification des erreurs ; 3.3.2 - Étiquettes ou instructions ; 3.3.3 - Suggestion après une erreur",
        "rgaa": "11.10, 11.11",
        "question": "Avant l'envoi, les champs obligatoires sont-ils indiqués ; après la soumission, les erreurs expliquent-elles comment corriger ?",
        "methode": "Observer l'état initial, puis soumettre un formulaire incomplet au clavier.",
        "conformite": "Obligations annoncées avant l'envoi. Après soumission, erreurs précises, reliées aux champs et accompagnées d'une aide à la correction.",
    },
]

EXERCICE_PAGES = [
    {"id": "ec01-images", "title": "Contrôler les images avant publication", "slides": "84-86"},
    {"id": "ec02-page-title", "title": "Résultats de recherche RGAA", "slides": "87"},
    {"id": "ec03-headings", "title": "Structurer la page avec des titres", "slides": "88-89"},
    {"id": "ec04-contrast", "title": "Vérifier le contraste des couleurs", "slides": "90-91"},
    {"id": "ec05-skiplinks", "title": "Accès rapide aux contenus", "slides": "92"},
    {"id": "ec06-keyboard-focus", "title": "Parcours clavier", "slides": "93-96"},
    {"id": "ec07-language", "title": "Déclarer les langues utilisées dans la page", "slides": "97"},
    {"id": "ec08-zoom", "title": "Ressources à zoomer", "slides": "98"},
    {"id": "ec09-captions", "title": "Vidéo de sensibilisation", "slides": "99-100"},
    {"id": "ec10-transcript", "title": "Écouter un podcast", "slides": "101"},
    {"id": "ec11-audio-description", "title": "Démonstration vidéo", "slides": "102"},
    {"id": "ec12-form-labels", "title": "Inscription à un webinaire", "slides": "105-107"},
    {"id": "ec13-required-errors", "title": "Formulaire de contact", "slides": "108"},
]

THIN = Side(border_style="thin", color=GRIS_MOYEN)
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT_TITLE = Font(name="Calibri", size=SIZE_TITLE, bold=True, color=BLEU_FRANCE)
FONT_SECTION = Font(name="Calibri", size=SIZE_SECTION, bold=True, color=BLEU_FRANCE)
FONT_HEADER = Font(name="Calibri", size=SIZE_HEADER, bold=True, color=BLANC)
FONT_CELL = Font(name="Calibri", size=SIZE_CELL)
FONT_CELL_BOLD = Font(name="Calibri", size=SIZE_CELL, bold=True)
FONT_LINK = Font(name="Calibri", size=SIZE_CELL, color=BLEU_FRANCE, underline="single")
FILL_HEADER = PatternFill(fill_type="solid", start_color=BLEU_FRANCE, end_color=BLEU_FRANCE)
FILL_SECTION = PatternFill(fill_type="solid", start_color=BLEU_CLAIR, end_color=BLEU_CLAIR)
FILL_EDITABLE = PatternFill(fill_type="solid", start_color=JAUNE_EDITABLE, end_color=JAUNE_EDITABLE)
ALIGN_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(ws, widths):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def style_header_row(ws, row, columns):
    for column in range(1, columns + 1):
        cell = ws.cell(row=row, column=column)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_ALL


def configure_print(ws, orientation="landscape", fit_height=0):
    ws.page_setup.orientation = orientation
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_height
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(
        left=0.4,
        right=0.4,
        top=0.5,
        bottom=0.5,
        header=0.3,
        footer=0.3,
    )
    ws.oddFooter.center.text = "Page &P / &N"
    ws.oddFooter.center.size = 11


def external_url(page_id, version="site-inaccessible"):
    return f"{BASE_URL}/{version}/{page_id}.html"


def add_link(cell, label, target):
    cell.value = label
    if target.startswith("#"):
        cell.hyperlink = Hyperlink(
            ref=cell.coordinate,
            location=target.removeprefix("#"),
            display=label,
        )
    else:
        cell.hyperlink = target
    cell.font = FONT_LINK
    cell.alignment = ALIGN_WRAP


def add_title(ws, number, title):
    ws["A1"] = number
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = ALIGN_CENTER
    ws["B1"] = title
    ws["B1"].font = FONT_TITLE
    ws["B1"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 42


def build_demarrer(wb):
    ws = wb.create_sheet("Démarrer")
    ws.sheet_properties.tabColor = BLEU_FRANCE
    set_widths(ws, [14, 30, 30, 22, 23, 23])

    ws.merge_cells("A1:F1")
    ws["A1"] = WORKBOOK_TITLE
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34
    ws["A3"] = "Objectif"
    ws["A3"].font = FONT_SECTION
    ws.merge_cells("B3:F3")
    ws["B3"] = (
        "Diagnostiquer un écart, proposer sa mise en conformité et conserver une preuve de l'écart. "
        "Ce carnet de sensibilisation ne remplace pas un audit RGAA formel."
    )
    ws["B3"].font = FONT_CELL
    ws["B3"].alignment = ALIGN_WRAP
    ws.row_dimensions[3].height = 46

    ws["A5"] = "Auditeur"
    ws["C5"] = "Date"
    ws["A6"] = "Navigateur"
    ws["C6"] = "Outils"
    for coordinate in ("A5", "C5", "A6", "C6"):
        ws[coordinate].font = FONT_CELL_BOLD
        ws[coordinate].alignment = ALIGN_WRAP
    for coordinate in ("B5", "D5", "B6", "D6"):
        ws[coordinate].fill = FILL_EDITABLE
        ws[coordinate].border = BORDER_ALL
        ws[coordinate].font = FONT_CELL
    ws["F5"] = "Cellules jaunes : à renseigner"
    ws["F5"].font = FONT_CELL
    ws["F5"].alignment = ALIGN_WRAP

    ws["A8"] = "Parcours"
    ws["A8"].font = FONT_SECTION
    ws.merge_cells("B8:F8")
    ws["B8"] = (
        "Ouvrir la page à tester, réaliser le contrôle, puis renseigner le constat, la mise en conformité à réaliser et la preuve. "
        "Consulter l'aide ou la correction seulement après la recherche autonome."
    )
    ws["B8"].font = FONT_CELL
    ws["B8"].alignment = ALIGN_WRAP
    ws.row_dimensions[8].height = 46

    headers = ["#", "Point testé", "Page d'exercice", "Slides", "Ouvrir la fiche", "Ouvrir la page"]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=10, column=column, value=header)
    style_header_row(ws, 10, len(headers))
    ws.row_dimensions[10].height = 34

    for row, (check, page) in enumerate(zip(CHECKS, EXERCICE_PAGES, strict=True), start=11):
        values = [f"#{check['id']:02d}", check["titre"], page["title"], page["slides"]]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.font = FONT_CELL
            cell.alignment = ALIGN_WRAP
            cell.border = BORDER_ALL
        add_link(ws.cell(row=row, column=5), "Ouvrir la fiche", f"#'{check['onglet']}'!A1")
        add_link(ws.cell(row=row, column=6), "Ouvrir la page", external_url(page["id"]))
        for column in (5, 6):
            ws.cell(row=row, column=column).border = BORDER_ALL
        ws.row_dimensions[row].height = 46

    section_row = 26
    ws.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=6)
    ws.cell(row=section_row, column=1, value="Après la recherche autonome")
    ws.cell(row=section_row, column=1).font = FONT_SECTION
    ws.cell(row=section_row + 1, column=1, value="#")
    ws.cell(row=section_row + 1, column=2, value="Page d'exercice")
    ws.cell(row=section_row + 1, column=3, value="Aide à la correction")
    ws.cell(row=section_row + 1, column=4, value="Version corrigée")
    style_header_row(ws, section_row + 1, 4)
    ws.row_dimensions[section_row + 1].height = 42

    for row, (check, page) in enumerate(zip(CHECKS, EXERCICE_PAGES, strict=True), start=section_row + 2):
        ws.cell(row=row, column=1, value=f"#{check['id']:02d}")
        ws.cell(row=row, column=2, value=page["title"])
        add_link(ws.cell(row=row, column=3), "Voir l'aide", external_url(page["id"], "site-aide-correction"))
        add_link(ws.cell(row=row, column=4), "Voir la correction", external_url(page["id"], "site-accessible"))
        for column in range(1, 5):
            cell = ws.cell(row=row, column=column)
            if column in (1, 2):
                cell.font = FONT_CELL
            cell.alignment = ALIGN_WRAP
            cell.border = BORDER_ALL
        ws.row_dimensions[row].height = 40

    sources_row = section_row + 16
    ws.merge_cells(start_row=sources_row, start_column=1, end_row=sources_row, end_column=6)
    ws.cell(row=sources_row, column=1, value="Références")
    ws.cell(row=sources_row, column=1).font = FONT_SECTION
    sources = [
        ("Points de contrôle rapides W3C", "https://www.w3.org/WAI/test-evaluate/easy-checks/"),
        ("WCAG 2.2", "https://www.w3.org/TR/WCAG22/"),
        ("RGAA 4.1.2", "https://accessibilite.numerique.gouv.fr/"),
    ]
    for row, (label, url) in enumerate(sources, start=sources_row + 1):
        add_link(ws.cell(row=row, column=2), label, url)

    ws.freeze_panes = "A11"
    ws.auto_filter.ref = "A10:F23"
    ws.print_area = f"A1:F{sources_row + len(sources)}"
    configure_print(ws, orientation="landscape")


def build_check_sheet(wb, check, page, previous_sheet, next_sheet):
    ws = wb.create_sheet(check["onglet"])
    ws.sheet_properties.tabColor = BLEU_CLAIR
    set_widths(ws, [34, 100])
    add_title(ws, f"#{check['id']:02d}", check["titre"])

    rows = [
        (3, "Page à auditer", page["title"]),
        (4, "Slides projetées", page["slides"]),
        (5, "Question à tester", check["question"]),
        (6, "Référence WCAG 2.2", check["wcag"]),
        (7, "Référence RGAA 4.1.2", check["rgaa"]),
        (8, "Comment tester", check["methode"]),
        (9, "Ce qu'il faut rendre conforme", check["conformite"]),
    ]
    heights = {3: 40, 4: 32, 5: 58, 6: 48, 7: 40, 8: 55, 9: 70}
    if check["id"] == 4:
        heights[8] = 78
    for row, label, value in rows:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=1).font = FONT_CELL_BOLD
        ws.cell(row=row, column=1).fill = FILL_SECTION
        ws.cell(row=row, column=2, value=value)
        ws.cell(row=row, column=2).font = FONT_CELL
        for column in (1, 2):
            ws.cell(row=row, column=column).alignment = ALIGN_WRAP
            ws.cell(row=row, column=column).border = BORDER_ALL
        ws.row_dimensions[row].height = heights[row]
    add_link(ws["B3"], page["title"], external_url(page["id"]))
    ws["B3"].border = BORDER_ALL

    for row, label, height in (
        (11, "Constat", 80),
        (12, "Mise en conformité à réaliser", 90),
        (13, "Preuve de l'écart", 80),
    ):
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=1).font = FONT_CELL_BOLD
        ws.cell(row=row, column=1).fill = FILL_SECTION
        ws.cell(row=row, column=2, value="")
        ws.cell(row=row, column=2).fill = FILL_EDITABLE
        ws.cell(row=row, column=2).font = FONT_CELL
        for column in (1, 2):
            ws.cell(row=row, column=column).alignment = ALIGN_WRAP
            ws.cell(row=row, column=column).border = BORDER_ALL
        ws.row_dimensions[row].height = height

    add_link(ws["A15"], "Retour à Démarrer", "#'Démarrer'!A1")
    add_link(ws["B15"], "Point précédent", f"#'{previous_sheet}'!A1")
    add_link(ws["A16"], "Point suivant", f"#'{next_sheet}'!A1")
    add_link(ws["B16"], "Voir la synthèse", "#'Synthèse'!A1")
    ws.row_dimensions[15].height = 30
    ws.row_dimensions[16].height = 30

    ws.freeze_panes = "A3"
    ws.print_area = "A1:B16"
    configure_print(ws, orientation="portrait", fit_height=1)


def build_synthese(wb):
    ws = wb.create_sheet("Synthèse")
    ws.sheet_properties.tabColor = VERT_SUCCES
    set_widths(ws, [8, 28, 34, 42, 46, 38])

    ws.merge_cells("A1:F1")
    ws["A1"] = "Synthèse des 13 diagnostics"
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:F2")
    ws["A2"] = (
        "Les constats, mises en conformité et preuves sont repris automatiquement depuis les 13 fiches. "
        "« À renseigner » signale un contrôle dont le constat est encore vide."
    )
    ws["A2"].font = FONT_CELL
    ws["A2"].alignment = ALIGN_WRAP
    ws.row_dimensions[2].height = 48

    headers = ["#", "Page", "Point testé", "Constat", "Mise en conformité à réaliser", "Preuve de l'écart"]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=4, column=column, value=header)
    style_header_row(ws, 4, len(headers))
    ws.row_dimensions[4].height = 42

    for row, (check, page) in enumerate(zip(CHECKS, EXERCICE_PAGES, strict=True), start=5):
        sheet_ref = check["onglet"].replace("'", "''")
        ws.cell(row=row, column=1, value=f"#{check['id']:02d}")
        ws.cell(row=row, column=2, value=page["title"])
        add_link(ws.cell(row=row, column=3), check["titre"], f"#'{check['onglet']}'!A1")
        ws.cell(row=row, column=4, value=f'=IF(\'{sheet_ref}\'!B11="","À renseigner",\'{sheet_ref}\'!B11)')
        ws.cell(
            row=row,
            column=5,
            value=f'=IF(\'{sheet_ref}\'!B12="","",\'{sheet_ref}\'!B12)',
        )
        ws.cell(
            row=row,
            column=6,
            value=f'=IF(\'{sheet_ref}\'!B13="","",\'{sheet_ref}\'!B13)',
        )
        for column in range(1, 7):
            cell = ws.cell(row=row, column=column)
            if column != 3:
                cell.font = FONT_CELL
            cell.alignment = ALIGN_WRAP
            cell.border = BORDER_ALL
        ws.row_dimensions[row].height = 72

    add_link(ws["B19"], "Retour à Démarrer", "#'Démarrer'!A1")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:F17"
    ws.print_title_rows = "1:4"
    ws.print_area = "A1:F19"
    configure_print(ws, orientation="landscape")


def main():
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "IGPDE"
    wb.properties.title = f"IGPDE - {FORMATION['footer']} - {WORKBOOK_TITLE}"
    wb.properties.language = "fr-FR"
    wb.properties.subject = "Carnet pédagogique de diagnostic des 13 points de contrôle rapides du W3C"
    wb.properties.description = "Un onglet par point de contrôle : constat, mise en conformité et preuve de l'écart."
    wb.properties.keywords = "accessibilité, W3C, WCAG 2.2, RGAA 4.1.2, diagnostic"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    build_demarrer(wb)
    for index, (check, page) in enumerate(zip(CHECKS, EXERCICE_PAGES, strict=True)):
        previous_sheet = "Démarrer" if index == 0 else CHECKS[index - 1]["onglet"]
        next_sheet = "Synthèse" if index == len(CHECKS) - 1 else CHECKS[index + 1]["onglet"]
        build_check_sheet(wb, check, page, previous_sheet, next_sheet)
    build_synthese(wb)

    wb.active = 0
    output = Path(__file__).parent.parent / "03-easy-checks" / "grille-audit-easy-checks.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"Classeur généré : {output}")
    print(f"Onglets : {len(wb.sheetnames)}")


if __name__ == "__main__":
    main()
